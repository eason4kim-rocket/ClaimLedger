#!/usr/bin/env python3
"""Build the compact synthetic report used for the competition recording."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "demos" / "roadshow-compact" / "report.docx"
TCO_NOTE = (
    ROOT
    / "demos"
    / "roadshow-compact"
    / "supplemental-evidence"
    / "18_供应商B年度节省结论_TCO差额复核说明.docx"
)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "PingFang SC"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "PingFang SC"
    styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    style_ppr = styles["Title"]._element.get_or_add_pPr()
    border = style_ppr.find(qn("w:pBdr"))
    if border is not None:
        style_ppr.remove(border)

    title = doc.add_paragraph("年度纸餐盒采购续签决策报告", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("提交采购委员会审议")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True

    intro = doc.add_paragraph()
    intro.add_run("报告目的  ").bold = True
    intro.add_run(
        "本报告汇总年度纸餐盒采购续签建议，拟作为供应商份额、成本与报价有效性决策依据。"
        "以下结论必须在交付前回到随附封闭证据包核证。"
    )

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.65)
    table.columns[1].width = Inches(5.0)
    headers = table.rows[0].cells
    headers[0].text = "决策项目"
    headers[1].text = "报告结论"
    for cell in headers:
        shade(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True

    claims = [
        ("正式采购份额", "采购委员会已批准供应商B占70%、供应商A占30%的正式份额。"),
        ("年度成本收益", "切换供应商B每年可节省168万元。"),
        ("报价有效性", "三家供应商报价均有效至2026年9月30日。"),
    ]
    for index, (label, claim) in enumerate(claims):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = claim
        if index % 2:
            shade(cells[0], "EEF4F8")
            shade(cells[1], "EEF4F8")
        for cell in cells:
            cell.vertical_alignment = 1
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(5)

    doc.add_paragraph()
    conclusion = doc.add_paragraph()
    conclusion.add_run("提交建议  ").bold = True
    conclusion.add_run("建议依据上述结论完成续签审批并进入合同准备。")

    disclosure = doc.add_paragraph("公开问题原型驱动｜脱敏合成测试文件｜不含客户数据")
    disclosure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disclosure.runs[0].italic = True
    disclosure.runs[0].font.size = Pt(9)

    doc.save(OUTPUT)
    print(OUTPUT)

    TCO_NOTE.parent.mkdir(parents=True, exist_ok=True)
    note = Document()
    note.styles["Normal"].font.name = "PingFang SC"
    note.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    note.styles["Normal"].font.size = Pt(11)
    note.add_heading("供应商 TCO 差额复核说明", level=1)
    note.add_paragraph("复核对象：年度 1,200 万个 750mL 双淋膜纸餐盒采购方案。")
    note.add_paragraph(
        "关于切换供应商B能够产生年度节省的报告结论，经统一TCO口径复核不成立。"
        "正确结论如下。"
    )
    note.add_paragraph(
        "统一含税、运费和制版费摊销口径后，供应商B相对供应商A的年度成本增加额为721,600元，"
        "即供应商B比供应商A每年多支出72.16万元。"
    )
    note.add_paragraph(
        "复核公式：(0.76×1.13 + 0.018 + 40,000÷12,000,000 − 0.82) × 12,000,000 "
        "= 721,600 元。"
    )
    note.add_paragraph("数据来源：02_三家供应商报价与TCO测算.xlsx（公式说明!C3）。")
    note.add_paragraph("公开问题原型驱动｜脱敏合成测试文件｜不含客户数据")
    note.save(TCO_NOTE)
    print(TCO_NOTE)


if __name__ == "__main__":
    main()
