from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
import json

from docx import Document
from openpyxl import load_workbook

from claimledger.demo import create_demo
from claimledger.engine import new_job, run_audit
from claimledger.exporters import _replace_span, export_final, export_ledger, export_standard_artifacts
from claimledger.models import Decision, JobRequest
from claimledger.parsers import sha256_file


def test_demo_audit_exports_all_artifacts(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    demo = create_demo(tmp_path / "demo")
    job = new_job(
        JobRequest(
            report_path=demo["report"],
            sources_path=demo["sources"],
            case_name="test-demo",
            rule_pack="lithium-demo",
        )
    )
    job, chunks = run_audit(job)
    assert job.status == "completed"
    assert chunks
    assert any(item.status.value == "conflict" for item in job.findings)
    outputs = export_standard_artifacts(job)
    assert all(Path(path).is_file() for path in outputs.values())
    annotated = Document(outputs["annotated_report"])
    assert annotated.paragraphs
    with ZipFile(outputs["annotated_report"]) as archive:
        assert "word/comments.xml" in archive.namelist()
    workbook = load_workbook(outputs["claim_ledger"], read_only=True)
    assert {
        "审计摘要",
        "结论证据台账",
        "候选证据",
        "人工决定日志",
        "文件清单",
    } <= set(workbook.sheetnames)
    workbook.close()

    for finding in job.findings:
        if finding.severity == "high":
            job.decisions.append(Decision(finding_id=finding.id, action="waive", reason="test fixture"))
    final = export_final(job)
    assert final.is_file()
    assert final.resolve() != Path(job.request.report_path).resolve()
    manifest = json.loads(Path(job.artifacts["manifest"]).read_text(encoding="utf-8"))
    assert {"delivery_report", "audit"} <= set(manifest["outputs"])
    assert manifest["outputs"]["delivery_report"] == sha256_file(final)


def test_final_export_is_blocked_for_unresolved_high_risk(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    demo = create_demo(tmp_path / "demo")
    job = new_job(JobRequest(report_path=demo["report"], sources_path=demo["sources"], rule_pack="lithium-demo"))
    job, _ = run_audit(job)
    if job.summary()["unresolved_high"]:
        try:
            export_final(job)
        except ValueError as exc:
            assert "blocked" in str(exc)
        else:
            raise AssertionError("final export should have been blocked")


def test_all_demo_scenarios_are_self_contained(tmp_path):
    for scenario in ("lithium", "operations", "consulting"):
        demo = create_demo(tmp_path / scenario, scenario)
        assert Path(demo["report"]).is_file()
        assert Path(demo["sources"]).is_dir()


def test_final_export_blocks_original_input_changed_after_audit(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    demo = create_demo(tmp_path / "demo")
    job = new_job(JobRequest(report_path=demo["report"], sources_path=demo["sources"], rule_pack="lithium-demo"))
    job, _ = run_audit(job)
    for finding in job.findings:
        if finding.severity == "high":
            job.decisions.append(Decision(finding_id=finding.id, action="waive", reason="test fixture"))
    report = Path(demo["report"])
    report.write_bytes(report.read_bytes() + b"changed-after-audit")
    try:
        export_final(job)
    except ValueError as exc:
        assert "original input changed" in str(exc)
    else:
        raise AssertionError("final export should be blocked when an original input hash changes")


def test_excel_ledger_neutralizes_formula_like_claim_text(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    demo = create_demo(tmp_path / "demo")
    job = new_job(JobRequest(report_path=demo["report"], sources_path=demo["sources"], rule_pack="lithium-demo"))
    job, _ = run_audit(job)
    job.claims[0].text = '=HYPERLINK("https://example.invalid","click")'
    path = export_ledger(job, tmp_path / "ledger.xlsx")
    workbook = load_workbook(path, data_only=False)
    try:
        value = workbook["结论证据台账"]["F2"].value
        assert value.startswith("'=")
    finally:
        workbook.close()


def test_precise_replacement_preserves_surrounding_runs():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("前缀：")
    paragraph.add_run("错误结论").bold = True
    paragraph.add_run("；后缀保留。")
    start = paragraph.text.index("错误结论")
    _replace_span(paragraph, start, start + len("错误结论"), "错误结论", "已核实结论")
    assert paragraph.text == "前缀：已核实结论；后缀保留。"
    assert paragraph.runs[-1].text == "；后缀保留。"
