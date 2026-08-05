from __future__ import annotations

from openpyxl import load_workbook

from claimledger.models import (
    AuditJob,
    Claim,
    Finding,
    FindingStatus,
    IssueCode,
    JobRequest,
    Locator,
)
from claimledger.reviewer import export_blind_workbook, import_blind_workbook


def _job() -> AuditJob:
    return AuditJob(
        id="job-blind",
        request=JobRequest(
            report_path="/tmp/report.docx",
            sources_path="/tmp/evidence",
            case_name="blind-case",
        ),
        claims=[
            Claim(
                id="claim-1",
                text="=公式开头的风险结论",
                claim_type="conclusion",
                locator=Locator(kind="paragraph", paragraph=0),
            ),
            Claim(
                id="claim-2",
                text="已有充分来源的结论",
                claim_type="conclusion",
                locator=Locator(kind="paragraph", paragraph=1),
            ),
        ],
        findings=[
            Finding(
                id="finding-1",
                claim_id="claim-1",
                status=FindingStatus.UNSUPPORTED,
                severity="high",
                confidence=0.9,
                explanation="无来源",
                issue_codes=[IssueCode.MISSING_EVIDENCE],
            ),
            Finding(
                id="finding-2",
                claim_id="claim-2",
                status=FindingStatus.SUPPORTED,
                severity="low",
                confidence=0.9,
                explanation="支持",
                issue_codes=[IssueCode.DIRECT_SUPPORT],
            ),
        ],
    )


def test_blind_export_omits_system_prediction_and_neutralizes_formula(tmp_path):
    path = export_blind_workbook(_job(), tmp_path / "blind.xlsx")
    workbook = load_workbook(path)
    try:
        sheet = workbook["Blind review"]
        headers = [cell.value for cell in sheet[1]]
        assert "status" not in headers
        assert "severity" not in headers
        assert sheet["D2"].value.startswith("'=")
    finally:
        workbook.close()


def test_blind_import_calculates_metrics_and_disagreements(tmp_path):
    job = _job()
    path = export_blind_workbook(job, tmp_path / "blind.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Blind review"]
    sheet["I2"] = "risk"
    sheet["J2"] = "missing_evidence"
    sheet["I3"] = "risk"
    workbook.save(path)
    workbook.close()

    metrics = import_blind_workbook(job, path)
    assert metrics["precision"] == 1.0
    assert metrics["recall"] == 0.5
    assert metrics["agreement"] == 0.5
    assert metrics["issue_code_f1"] == 1.0
    assert metrics["disagreements"] == [
        {"finding_id": "finding-2", "system": "clean", "reviewer": "risk"}
    ]
