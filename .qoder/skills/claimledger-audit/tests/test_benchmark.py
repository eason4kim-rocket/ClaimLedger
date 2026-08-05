from __future__ import annotations

import json

from docx import Document

from claimledger.benchmark import run_benchmark
from claimledger.demo import create_demo


def test_benchmark_computes_auditable_metrics(tmp_path):
    demo = create_demo(tmp_path / "demo")
    dataset = {
        "name": "tiny",
        "cases": [
            {
                "name": "lithium",
                "report": demo["report"],
                "sources": demo["sources"],
                "rule_pack": "lithium-demo",
                "claims": [
                    {
                        "text": "SC6 锂辉石与低品位锂云母可以直接按照吨价比较采购成本。",
                        "should_flag": True,
                        "accepted_statuses": ["conflict"],
                        "evidence_file": "public-evidence.docx",
                    },
                    {
                        "text": "2026年7月，电池级碳酸锂现货参考价为每吨 120000 元。",
                        "should_flag": False,
                        "accepted_statuses": ["supported"],
                        "evidence_file": "public-evidence.docx",
                    },
                ],
            }
        ],
    }
    path = tmp_path / "labels.json"
    path.write_text(json.dumps(dataset, ensure_ascii=False), encoding="utf-8")
    result = run_benchmark(path)
    assert result["counts"]["labeled_claims"] == 2
    assert result["metrics"]["claim_extraction_recall"] == 1.0
    assert result["metrics"]["evidence_top5_recall"] == 1.0


def test_missing_risk_claim_is_a_false_negative(tmp_path):
    report = Document()
    report.add_heading("空白审阅报告", level=0)
    report.add_paragraph("本次登记材料共2份。")
    report.save(tmp_path / "report.docx")
    sources = tmp_path / "sources"
    sources.mkdir()
    evidence = Document()
    evidence.add_paragraph("来源材料仅记录一般背景。")
    evidence.add_paragraph("本次登记材料共2份。")
    evidence.save(sources / "evidence.docx")
    dataset = {
        "name": "missing-risk",
        "cases": [
            {
                "name": "missing-risk",
                "report": "report.docx",
                "sources": "sources",
                "rule_pack": "generic-zh",
                "claims": [
                    {
                        "text": "报告声称不存在的高风险结论为100万元。",
                        "should_flag": True,
                        "expected_severity": "high",
                        "accepted_statuses": ["conflict"],
                    }
                ],
            }
        ],
    }
    path = tmp_path / "labels.json"
    path.write_text(json.dumps(dataset, ensure_ascii=False), encoding="utf-8")

    result = run_benchmark(path)

    assert result["counts"]["true_positives"] == 0
    assert result["counts"]["false_negatives"] == 1
    assert result["metrics"]["error_recall"] == 0.0


def test_unlabeled_supported_claim_is_not_error_false_positive(tmp_path):
    report = Document()
    report.add_paragraph("报告统计期间为2026年7月1日至2026年7月31日。")
    report.add_paragraph("本次登记材料共2份。")
    report.save(tmp_path / "report.docx")
    sources = tmp_path / "sources"
    sources.mkdir()
    evidence = Document()
    evidence.add_heading("原始证据", level=1)
    evidence.add_paragraph("报告统计期间为2026年7月1日至2026年7月31日。")
    evidence.add_paragraph("本次登记材料共2份。")
    evidence.save(sources / "evidence.docx")
    dataset = {
        "name": "unlabeled-supported",
        "cases": [
            {
                "name": "unlabeled-supported",
                "report": "report.docx",
                "sources": "sources",
                "rule_pack": "generic-zh",
                "claims": [
                    {
                        "text": "报告统计期间为2026年7月1日至2026年7月31日。",
                        "should_flag": False,
                        "accepted_statuses": ["supported"],
                        "evidence_file": "evidence.docx",
                    }
                ],
            }
        ],
    }
    path = tmp_path / "labels.json"
    path.write_text(json.dumps(dataset, ensure_ascii=False), encoding="utf-8")

    result = run_benchmark(path)

    assert result["counts"]["unlabeled_claims"] == 1
    assert result["counts"]["unlabeled_risk_predictions"] == 0
    assert result["counts"]["false_positives"] == 0
