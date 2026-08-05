from __future__ import annotations

import json
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook

from claimledger.cli import build_parser
from claimledger.demo import DISCLOSURE, create_demo
from claimledger.parsers import OcrEngine, parse_docx, parse_pdf, parse_xlsx


def _claims(payload: dict) -> list[dict]:
    return [claim for case in payload["cases"] for claim in case["claims"]]


def test_golden_v2_contains_100_cross_industry_claims(tmp_path):
    generated = create_demo(tmp_path / "golden-v2", "golden-v2")
    labels = json.loads(Path(generated["labels"]).read_text(encoding="utf-8"))
    manifest = json.loads(Path(generated["manifest"]).read_text(encoding="utf-8"))

    assert labels["disclosure"] == DISCLOSURE
    assert [case["name"] for case in labels["cases"]] == [
        "procurement-renewal",
        "operations-fulfillment-monthly",
        "consulting-market-entry",
    ]
    assert [len(case["claims"]) for case in labels["cases"]] == [44, 28, 28]
    assert len(_claims(labels)) == 100
    assert sum(claim["should_flag"] for claim in _claims(labels)) == 49
    assert manifest["claims"] == 100
    assert manifest["planted_risks"] == 49
    assert manifest["evidence_groups"] == 36
    assert manifest["case_ids"] == ["CL-PROC-001", "CL-OPS-001", "CL-CONS-001"]


def test_golden_labels_are_complete_relative_and_self_contained(tmp_path):
    generated = create_demo(tmp_path / "golden-v2", "golden-v2")
    labels_path = Path(generated["labels"])
    labels = json.loads(labels_path.read_text(encoding="utf-8"))
    serialized = labels_path.read_text(encoding="utf-8")
    assert str(Path.home() / "Documents" / "code") not in serialized
    suite_manifest = json.loads(Path(generated["manifest"]).read_text(encoding="utf-8"))
    assert suite_manifest["labels"] == "labels.json"

    for case in labels["cases"]:
        report_path = labels_path.parent / case["report"]
        source_dir = labels_path.parent / case["sources"]
        case_manifest_path = report_path.parent / "case_manifest.json"
        case_manifest = json.loads(case_manifest_path.read_text(encoding="utf-8"))
        assert case_manifest["report"] == "report.docx"
        assert case_manifest["sources"] == "evidence"
        assert case_manifest["labels"] == "labels.json"
        assert report_path.is_file()
        assert source_dir.is_dir()
        report_text = {paragraph.text for paragraph in Document(report_path).paragraphs}
        for claim in case["claims"]:
            assert claim["text"] in report_text
            assert claim["expected_status"] in claim["accepted_statuses"]
            assert claim["expected_issue_codes"]
            assert claim["expected_severity"] in {"low", "medium", "high"}
            assert claim["evidence_file"]
            assert claim["evidence_quote"]
            assert claim["evidence_locator"]
            assert (source_dir / claim["evidence_file"]).is_file()


def test_operations_and_consulting_gold_quotes_resolve_exactly(tmp_path):
    for scenario in ("operations", "consulting"):
        generated = create_demo(tmp_path / scenario, scenario)
        labels = json.loads(Path(generated["labels"]).read_text(encoding="utf-8"))
        source_dir = Path(generated["sources"])
        cache_dir = tmp_path / "parse-cache" / scenario
        cache_dir.mkdir(parents=True, exist_ok=True)
        chunks_by_file = {}
        for path in source_dir.iterdir():
            if path.suffix.lower() == ".docx":
                chunks = parse_docx(path)
            elif path.suffix.lower() == ".xlsx":
                chunks = parse_xlsx(path)
            elif path.suffix.lower() == ".pdf":
                chunks, warnings = parse_pdf(path, OcrEngine(enabled=False), cache_dir)
                assert not warnings
            else:
                continue
            chunks_by_file[path.name] = chunks

        for claim in labels["cases"][0]["claims"]:
            matches = [
                chunk
                for chunk in chunks_by_file[claim["evidence_file"]]
                if claim["evidence_quote"] in chunk.quote
                and chunk.locator.label() == claim["evidence_locator"]
            ]
            assert matches, (
                scenario,
                claim["text"],
                claim["evidence_file"],
                claim["evidence_quote"],
                claim["evidence_locator"],
            )


def test_every_generated_material_carries_reconstruction_disclosure(tmp_path):
    for scenario in ("operations", "consulting"):
        generated = create_demo(tmp_path / scenario, scenario)
        report = Document(generated["report"])
        assert DISCLOSURE in " ".join(paragraph.text for section in report.sections for paragraph in section.footer.paragraphs)
        for path in Path(generated["sources"]).iterdir():
            if path.suffix.lower() == ".docx":
                document = Document(path)
                footer_text = " ".join(
                    paragraph.text
                    for section in document.sections
                    for paragraph in section.footer.paragraphs
                )
                assert DISCLOSURE in footer_text
            elif path.suffix.lower() == ".xlsx":
                workbook = load_workbook(path, read_only=False)
                try:
                    assert workbook.properties.subject == DISCLOSURE
                    assert all(sheet.oddFooter.center.text == DISCLOSURE for sheet in workbook.worksheets)
                finally:
                    workbook.close()
            elif path.suffix.lower() == ".pdf":
                document = fitz.open(path)
                try:
                    assert document.metadata["subject"] == DISCLOSURE
                finally:
                    document.close()


def test_cli_accepts_golden_v2_scenario(tmp_path):
    parser = build_parser()
    args = parser.parse_args(
        ["demo", "--output", str(tmp_path / "suite"), "--scenario", "golden-v2", "--json"]
    )
    assert args.command == "demo"
    assert args.scenario == "golden-v2"
