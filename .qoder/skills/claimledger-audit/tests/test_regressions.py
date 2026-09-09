from __future__ import annotations

import fitz
import os
import stat
from docx import Document
from openpyxl import Workbook
from pathlib import Path

from claimledger.engine import extract_claims, extract_facts, lexical_score, new_job, run_audit, verify
from claimledger.models import IssueCode, JobRequest, Locator, ParsedChunk
from claimledger.parsers import parse_xlsx, sha256_text


def make_chunk(
    text: str,
    *,
    chunk_id: str,
    file_name: str,
    file_hash: str,
    source_type: str = "docx",
    context_text: str | None = None,
    locator: Locator | None = None,
    metadata: dict | None = None,
    confidence: float = 1.0,
) -> ParsedChunk:
    return ParsedChunk(
        id=chunk_id,
        file_path=f"/tmp/{file_name}",
        file_name=file_name,
        file_hash=file_hash,
        source_type=source_type,
        text=text,
        raw_text=text,
        context_text=context_text or text,
        locator=locator or Locator(kind="paragraph", paragraph=0),
        metadata=metadata or {},
        confidence=confidence,
    )


def test_adjacent_measurements_keep_their_scale_units():
    facts = extract_facts("年度需求为1200万个750mL双淋膜纸餐盒。")
    values = {(item.base_unit, item.base_value) for item in facts if item.kind == "number"}
    assert ("COUNT", "12000000") in values
    assert ("ML", "750") in values


def test_average_does_not_trigger_the_absolute_qualifier_jun():
    report = make_chunk(
        "供应商B平均交付周期为10个自然日。",
        chunk_id="report-1",
        file_name="report.docx",
        file_hash="a" * 64,
    )
    claim = extract_claims([report])[0]
    assert "均" not in claim.qualifiers


def test_entity_allocation_reversal_is_not_hidden_by_unordered_percentages():
    rules = {
        "freshness_terms": [],
        "exclusive_categories": [
            {
                "id": "supplier-entity",
                "issue_code": "entity_mismatch",
                "values": {
                    "supplier_a": ["供应商A"],
                    "supplier_b": ["供应商B"],
                },
            }
        ],
    }
    report = make_chunk(
        "供应商B占70%、供应商A占30%。",
        chunk_id="report-2",
        file_name="report.docx",
        file_hash="b" * 64,
    )
    claim = extract_claims([report])[0]
    evidence = make_chunk(
        "委员会仅批准供应商B占30%、供应商A占70%。",
        chunk_id="evidence-2",
        file_name="minutes.docx",
        file_hash="c" * 64,
    )
    finding = verify(claim, [(evidence, 0.8)], rules)
    assert finding.status.value == "conflict"
    assert IssueCode.NUMERIC_MISMATCH in finding.issue_codes


def test_quote_price_is_not_compared_with_a_separate_freight_cell():
    report = make_chunk(
        "供应商A报价0.82元/个，含税含运。",
        chunk_id="report-3",
        file_name="report.docx",
        file_hash="d" * 64,
    )
    claim = extract_claims([report])[0]
    quote = make_chunk(
        "报价：0.82元/个，含税含运",
        chunk_id="quote-3",
        file_name="supplier-a-quote.pdf",
        file_hash="e" * 64,
        source_type="pdf",
        locator=Locator(kind="pdf_page", page=1),
    )
    freight = make_chunk(
        "0 元/个",
        chunk_id="freight-3",
        file_name="tco.xlsx",
        file_hash="f" * 64,
        source_type="xlsx",
        context_text="供应商A | 运费(元/个): 0",
        locator=Locator(kind="sheet_cell", sheet="TCO", cell="F2"),
        metadata={"header": "运费(元/个)", "row_key": "供应商A"},
    )
    finding = verify(claim, [(quote, 0.8), (freight, 0.6)], {"freshness_terms": []})
    assert finding.status.value == "supported"


def test_distinct_receipts_reconcile_an_aggregate():
    report = make_chunk(
        "历史三批采购累计付款58.4万元。",
        chunk_id="report-4",
        file_name="report.docx",
        file_hash="1" * 64,
    )
    claim = extract_claims([report])[0]
    candidates = []
    for index, amount in enumerate(("18万元", "19.6万元", "20.8万元"), start=1):
        receipt = make_chunk(
            f"付款金额：{amount}",
            chunk_id=f"receipt-{index}",
            file_name=f"16{index}_银行电子回单.png",
            file_hash=str(index + 1) * 64,
            source_type="png",
            locator=Locator(kind="image", bbox=[0, 0, 10, 10]),
        )
        candidates.append((receipt, 0.7 - index / 100))
    finding = verify(claim, candidates, {"freshness_terms": [], "ocr": {"review_below": 0.72}})
    assert finding.status.value == "supported"
    assert IssueCode.DERIVED_CALCULATION in finding.issue_codes
    assert finding.checks[0].evidence_value == "584000"


def test_top_five_evidence_is_source_diverse():
    report = make_chunk(
        "项目金额达到100万元。",
        chunk_id="report-5",
        file_name="report.docx",
        file_hash="9" * 64,
    )
    claim = extract_claims([report])[0]
    candidates = [
        (
            make_chunk(
                "项目金额达到100万元。",
                chunk_id=f"same-{index}",
                file_name="many-cells.xlsx",
                file_hash="7" * 64,
            ),
            0.9 - index / 100,
        )
        for index in range(6)
    ]
    candidates.append(
        (
            make_chunk(
                "项目金额达到100万元。",
                chunk_id="other-source",
                file_name="signed-contract.pdf",
                file_hash="8" * 64,
                source_type="pdf",
                locator=Locator(kind="pdf_page", page=1),
            ),
            0.7,
        )
    )
    finding = verify(claim, candidates, {"freshness_terms": []})
    assert len({item.file_hash for item in finding.evidence[:5]}) == 2
    assert finding.evidence[1].file_name == "signed-contract.pdf"


def test_evidence_quote_hash_matches_the_stored_truncated_quote():
    report = make_chunk(
        "项目金额达到100万元。",
        chunk_id="report-quote-hash",
        file_name="report.docx",
        file_hash="4" * 64,
    )
    claim = extract_claims([report])[0]
    long_quote = "项目金额达到100万元。" + "补充说明" * 500
    source = make_chunk(
        long_quote,
        chunk_id="long-evidence",
        file_name="long-contract.pdf",
        file_hash="5" * 64,
        source_type="pdf",
        locator=Locator(kind="pdf_page", page=1),
    )
    finding = verify(claim, [(source, 0.9)], {"freshness_terms": []})
    stored = finding.evidence[0]
    assert len(stored.quote) == 1600
    assert stored.quote_hash == sha256_text(stored.quote)
    assert stored.quote_hash != sha256_text(long_quote)


def test_absolute_threshold_counterexample_is_promoted_into_top_five():
    report = make_chunk(
        "所有仓库库存准确率均不低于99%。",
        chunk_id="report-threshold",
        file_name="report.docx",
        file_hash="6" * 64,
    )
    claim = extract_claims([report])[0]
    candidates = []
    for index, value in enumerate(("99.2%", "99.5%", "99.4%", "99.1%", "99.3%", "98.7%"), start=2):
        candidates.append(
            (
                make_chunk(
                    f"库存准确率: {value}",
                    chunk_id=f"warehouse-{index}",
                    file_name="inventory.xlsx",
                    file_hash="7" * 64,
                    source_type="xlsx",
                    locator=Locator(kind="sheet_cell", sheet="库存准确率", cell=f"B{index}"),
                    metadata={"header": "库存准确率", "row_key": f"仓库{index}"},
                ),
                0.95 - index / 100,
            )
        )
    candidates.append(
        (
            make_chunk(
                "库存统计口径按月末盘点结果计算。",
                chunk_id="inventory-method",
                file_name="inventory-method.docx",
                file_hash="8" * 64,
            ),
            0.7,
        )
    )
    finding = verify(claim, candidates, {"freshness_terms": []})
    first_five = finding.evidence[:5]
    assert any(item.quote == "库存准确率: 98.7%" for item in first_five)
    assert len({item.file_hash for item in first_five}) == 2


def test_review_evidence_prefers_local_category_counterexample_over_page_context():
    report = make_chunk(
        "检测报告样品与750mL双淋膜采购品一致。",
        chunk_id="report-sample-scope",
        file_name="report.docx",
        file_hash="9" * 64,
    )
    claim = extract_claims([report])[0]
    rules = {
        "freshness_terms": [],
        "exclusive_categories": [
            {
                "id": "package-volume",
                "issue_code": "scope_mismatch",
                "values": {"ml650": ["650mL"], "ml750": ["750mL"]},
            },
            {
                "id": "coating-spec",
                "issue_code": "scope_mismatch",
                "values": {"single": ["单淋膜"], "double": ["双淋膜"]},
            },
        ],
    }
    candidates = [
        (
            make_chunk(
                "淋膜: 双淋膜",
                chunk_id="scope-top",
                file_name="comparison.xlsx",
                file_hash="a" * 64,
                source_type="xlsx",
                locator=Locator(kind="sheet_cell", sheet="比价", cell="C2"),
                metadata={"header": "淋膜", "row_key": "供应商A"},
            ),
            0.9,
        ),
        (
            make_chunk(
                "产品检测报告",
                chunk_id="scope-page-title",
                file_name="产品检测报告.pdf",
                file_hash="b" * 64,
                source_type="pdf",
                context_text="产品检测报告 | 样品：650mL单淋膜纸餐盒",
                locator=Locator(kind="pdf_page", page=1),
            ),
            0.8,
        ),
        (
            make_chunk(
                "产品：750mL双淋膜纸餐盒",
                chunk_id="scope-support-1",
                file_name="comparison.xlsx",
                file_hash="a" * 64,
            ),
            0.78,
        ),
        (
            make_chunk(
                "产品规格: 750mL",
                chunk_id="scope-support-2",
                file_name="comparison.xlsx",
                file_hash="a" * 64,
            ),
            0.76,
        ),
        (
            make_chunk(
                "样品信息见检测报告正文。",
                chunk_id="scope-support-3",
                file_name="comparison.xlsx",
                file_hash="a" * 64,
            ),
            0.74,
        ),
        (
            make_chunk(
                "样品：650mL单淋膜纸餐盒",
                chunk_id="scope-local-conflict",
                file_name="产品检测报告.pdf",
                file_hash="b" * 64,
                source_type="pdf",
                context_text="产品检测报告 | 样品：650mL单淋膜纸餐盒",
                locator=Locator(kind="pdf_page", page=1),
            ),
            0.7,
        ),
    ]
    finding = verify(claim, candidates, rules)
    assert any(item.quote == "样品：650mL单淋膜纸餐盒" for item in finding.evidence[:5])


def test_xlsx_rate_cells_keep_exact_cell_locator_and_percent_context(tmp_path):
    path = tmp_path / "kpi.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "履约KPI"
    sheet.append(["供应商", "准时交付率"])
    sheet.append(["供应商A", 0.932])
    workbook.save(path)

    chunks = parse_xlsx(path)
    rate = next(item for item in chunks if item.locator.cell == "B2")
    assert rate.locator.label() == "履约KPI!B2"
    assert "93.2%" in rate.context_text
    assert rate.quote == "准时交付率: 0.932"


def test_derived_cost_and_rank_cells_receive_narrow_semantic_bonus():
    saving_claim = extract_claims([
        make_chunk(
            "切换供应商B每年可节省168万元。",
            chunk_id="report-saving",
            file_name="report.docx",
            file_hash="a" * 64,
        )
    ])[0]
    result = make_chunk(
        "721600",
        chunk_id="cost-result",
        file_name="02_三家供应商报价与TCO测算.xlsx",
        file_hash="b" * 64,
        source_type="xlsx",
        context_text="公式说明 | 结果 | B相对A年度成本增加额 | 721600",
        locator=Locator(kind="sheet_cell", sheet="公式说明", cell="C3"),
        metadata={"header": "结果", "row_key": "B相对A年度成本增加额"},
    )
    supplier = make_chunk(
        "供应商B",
        chunk_id="supplier-cell",
        file_name="02_三家供应商报价与TCO测算.xlsx",
        file_hash="b" * 64,
        source_type="xlsx",
        context_text="TCO复核 | 供应商 | 供应商B",
        locator=Locator(kind="sheet_cell", sheet="TCO复核", cell="A3"),
        metadata={"header": "供应商", "row_key": "供应商B"},
    )
    assert lexical_score(saving_claim, result) > lexical_score(saving_claim, supplier)

    rank_claim = extract_claims([
        make_chunk(
            "供应商B的综合到岸成本最低。",
            chunk_id="report-rank",
            file_name="report.docx",
            file_hash="c" * 64,
        )
    ])[0]
    rank = make_chunk(
        "3",
        chunk_id="rank-cell",
        file_name="02_三家供应商报价与TCO测算.xlsx",
        file_hash="b" * 64,
        source_type="xlsx",
        context_text="TCO复核 | 排名 | 供应商B | 排名: 3",
        locator=Locator(kind="sheet_cell", sheet="TCO复核", cell="K3"),
        metadata={"header": "排名", "row_key": "供应商B"},
    )
    assert lexical_score(rank_claim, rank) > lexical_score(rank_claim, supplier)


def test_partial_pdf_page_failure_marks_evidence_coverage_incomplete(tmp_path):
    report_path = tmp_path / "report.docx"
    report = Document()
    report.add_paragraph("项目金额为100万元。")
    report.save(report_path)

    sources = tmp_path / "evidence"
    sources.mkdir()
    pdf_path = sources / "contract.pdf"
    pdf = fitz.open()
    readable = pdf.new_page()
    readable.insert_text((72, 72), "Contract amount is CNY 1,000,000 and is approved.")
    pdf.new_page()  # No text layer; deterministic profile cannot OCR this page.
    pdf.save(pdf_path)
    pdf.close()

    job = new_job(
        JobRequest(
            report_path=str(report_path),
            sources_path=str(sources),
            profile="deterministic",
            rule_pack="generic-zh",
        )
    )
    completed, _chunks = run_audit(job)

    assert completed.status == "completed"
    assert completed.coverage_status == "incomplete"
    assert any("contract.pdf page 2" in warning for warning in completed.warnings)
    cache_dir = Path(completed.report_snapshot).parents[2] / "ocr-cache"
    assert cache_dir.is_dir()
    if os.name != "nt":
        assert stat.S_IMODE(cache_dir.stat().st_mode) == 0o700
        assert all(stat.S_IMODE(path.stat().st_mode) == 0o600 for path in cache_dir.iterdir())
