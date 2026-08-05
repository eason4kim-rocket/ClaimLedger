from __future__ import annotations

import ipaddress
import socket

from docx import Document

from claimledger.engine import new_job, run_audit
from claimledger.models import JobRequest


def test_lite_text_audit_attempts_no_non_loopback_network(tmp_path, monkeypatch):
    report = tmp_path / "report.docx"
    report_document = Document()
    report_document.add_paragraph("2026年6月订单履约率为95%。")
    report_document.save(report)
    sources = tmp_path / "evidence"
    sources.mkdir()
    evidence = sources / "operations-source.docx"
    evidence_document = Document()
    evidence_document.add_paragraph("运营系统来源记录：2026年6月订单履约率为95%。")
    evidence_document.save(evidence)
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))

    blocked: list[str] = []
    original_connect = socket.socket.connect
    original_create_connection = socket.create_connection

    def is_loopback(address) -> bool:
        if not isinstance(address, tuple) or not address:
            return True
        host = str(address[0])
        if host == "localhost":
            return True
        try:
            return ipaddress.ip_address(host).is_loopback
        except ValueError:
            return False

    def guarded_connect(sock, address):
        if not is_loopback(address):
            blocked.append(str(address))
            raise AssertionError(f"non-loopback network attempt: {address}")
        return original_connect(sock, address)

    def guarded_create_connection(address, *args, **kwargs):
        if not is_loopback(address):
            blocked.append(str(address))
            raise AssertionError(f"non-loopback network attempt: {address}")
        return original_create_connection(address, *args, **kwargs)

    monkeypatch.setattr(socket.socket, "connect", guarded_connect)
    monkeypatch.setattr(socket, "create_connection", guarded_create_connection)

    job = new_job(
        JobRequest(
            report_path=str(report),
            sources_path=str(sources),
            case_name="offline-lite",
            profile="lite",
            rule_pack="operations-zh",
            as_of_date="2026-07-19",
        )
    )
    job, chunks = run_audit(job)

    assert job.status == "completed"
    assert job.coverage_status == "complete"
    assert chunks
    assert blocked == []
