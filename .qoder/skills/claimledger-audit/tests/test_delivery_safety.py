from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
import json
import os
from pathlib import Path
from pathlib import PureWindowsPath
import shutil
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from claimledger.cli import decide as cli_decide
from claimledger.exporters import export_final, export_standard_artifacts
from claimledger.ingestion import snapshot_inputs
from claimledger.models import (
    AuditJob,
    Claim,
    Decision,
    DecisionRequest,
    Evidence,
    Finding,
    FindingStatus,
    JobRequest,
    Locator,
    SourceRecord,
)
from claimledger.parsers import sha256_file, sha256_text
from claimledger.service import create_app
from claimledger.storage import ConcurrentUpdateError, JobStore


def _high_risk_job() -> AuditJob:
    claim = Claim(
        id="claim-0001",
        text="供应商已经满足全部交付要求。",
        claim_type="conclusion",
        locator=Locator(
            kind="paragraph",
            paragraph=0,
            char_start=0,
            char_end=14,
            anchor_precision="span",
        ),
    )
    finding = Finding(
        id="finding-0001",
        claim_id=claim.id,
        status=FindingStatus.UNSUPPORTED,
        severity="high",
        confidence=0.9,
        explanation="No source supports the absolute statement.",
    )
    return AuditJob(
        id="job-safety",
        request=JobRequest(report_path="report.docx", sources_path="evidence"),
        status="completed",
        coverage_status="complete",
        claims=[claim],
        findings=[finding],
    )


def _exportable_job(tmp_path: Path) -> AuditJob:
    original_root = tmp_path / "private-client-inputs"
    original_root.mkdir(parents=True)
    report = original_root / "board-report.docx"
    report_document = Document()
    report_document.add_paragraph("供应商已经满足全部交付要求。")
    report_document.save(report)
    evidence_root = original_root / "evidence"
    evidence_root.mkdir()
    evidence_file = evidence_root / "acceptance-record.docx"
    evidence_document = Document()
    evidence_document.add_paragraph("验收记录仍有两项未完成。")
    evidence_document.save(evidence_file)
    report_hash = sha256_file(report)
    evidence_hash = sha256_file(evidence_file)
    quote = "验收记录仍有两项未完成。"
    job = _high_risk_job()
    job.request = JobRequest(
        report_path=str(report),
        sources_path=str(evidence_root),
        case_name="privacy-delivery",
    )
    job.report_snapshot = str(report)
    job.sources_snapshot = str(evidence_root)
    job.source_inventory = [
        SourceRecord(
            id=f"source-{report_hash[:12]}",
            role="report",
            original_path=str(report),
            snapshot_path=str(report),
            file_name=report.name,
            file_hash=report_hash,
            size_bytes=report.stat().st_size,
            status="parsed",
            parsed_chunks=1,
        ),
        SourceRecord(
            id=f"source-{evidence_hash[:12]}",
            role="evidence",
            original_path=str(evidence_file),
            snapshot_path=str(evidence_file),
            file_name=evidence_file.name,
            file_hash=evidence_hash,
            size_bytes=evidence_file.stat().st_size,
            status="parsed",
            parsed_chunks=1,
        ),
    ]
    job.findings[0].evidence = [
        Evidence(
            id="evidence-0001",
            chunk_id="chunk-0001",
            file_path=str(evidence_file),
            file_name=evidence_file.name,
            file_hash=evidence_hash,
            quote=quote,
            quote_hash=sha256_text(quote),
            locator=Locator(kind="paragraph", paragraph=0, anchor_precision="paragraph"),
            score=0.91,
        )
    ]
    return job


def _all_strings(value):
    if isinstance(value, dict):
        for item in value.values():
            yield from _all_strings(item)
    elif isinstance(value, list):
        for item in value:
            yield from _all_strings(item)
    elif isinstance(value, str):
        yield value


def _is_absolute(value: str) -> bool:
    return Path(value).is_absolute() or PureWindowsPath(value).is_absolute()


def test_review_uses_safe_data_bindings_and_csp(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    job = _high_risk_job()
    job.claims[0].text = '结论"><img src=x onerror="window.pwned=1">'
    store = JobStore(tmp_path / "jobs.sqlite3")
    store.save(job)
    app = create_app(store)
    client = TestClient(app)

    response = client.get(f"/review/{job.id}?token={app.state.token}")

    assert response.status_code == 200
    assert "onclick=" not in response.text
    assert 'data-decision-action="replace"' in response.text
    assert "<img src=x" not in response.text
    assert "&lt;img src=x" in response.text
    assert "script-src 'nonce-" in response.headers["content-security-policy"]
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["referrer-policy"] == "no-referrer"


def test_new_decision_invalidates_current_delivery_pointer(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    job = _high_risk_job()
    historical = tmp_path / "revision-000000-final" / "delivery_report.docx"
    historical.parent.mkdir()
    historical.write_bytes(b"historical")
    job.artifacts = {"delivery_report": str(historical)}
    job.artifacts_revision = 0
    job.artifacts_kind = "final"
    job.artifact_bundle_id = "revision-000000-final"
    store = JobStore(tmp_path / "jobs.sqlite3")
    store.save(job)
    app = create_app(store)
    client = TestClient(app)

    response = client.post(
        f"/api/v1/jobs/{job.id}/decisions",
        headers={"Authorization": f"Bearer {app.state.token}"},
        json={"finding_id": "finding-0001", "action": "reject", "reason": "confirmed false positive"},
    )

    assert response.status_code == 200
    current = store.get(job.id)
    assert current is not None
    assert current.decision_revision == 1
    assert current.artifacts == {}
    assert current.artifacts_revision is None
    assert current.summary()["delivery_artifact_current"] is False
    assert current.artifact_history[-1]["artifacts"]["delivery_report"] == str(historical)
    assert historical.is_file()


def test_concurrent_decisions_are_not_lost(tmp_path):
    database = tmp_path / "jobs.sqlite3"
    initial_store = JobStore(database)
    job = _high_risk_job()
    initial_store.save(job)
    stores = [JobStore(database) for _ in range(16)]

    def append(index: int) -> None:
        decision = Decision(
            finding_id="finding-0001",
            action="accept",
            actor=f"reviewer-{index}",
        )

        def apply(current: AuditJob) -> None:
            current.record_decision(decision)

        assert stores[index].mutate(job.id, apply) is not None

    with ThreadPoolExecutor(max_workers=8) as pool:
        list(pool.map(append, range(len(stores))))

    persisted = initial_store.get(job.id)
    assert persisted is not None
    assert len(persisted.decisions) == len(stores)
    assert persisted.decision_revision == len(stores)
    assert len({item.event_id for item in persisted.decisions}) == len(stores)


def test_stale_save_is_rejected_instead_of_overwriting_newer_state(tmp_path):
    store = JobStore(tmp_path / "jobs.sqlite3")
    job = _high_risk_job()
    store.save(job)
    first = store.get(job.id)
    stale = store.get(job.id)
    assert first is not None and stale is not None
    first.record_decision(Decision(finding_id="finding-0001", action="accept"))
    store.save(first)
    stale.record_decision(Decision(finding_id="finding-0001", action="reject", reason="stale"))
    with pytest.raises(ConcurrentUpdateError):
        store.save(stale)


def test_final_bundle_is_revisioned_and_updates_delivery_disclosure(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    report = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("供应商续签决策报告")
    document.add_paragraph("存在高风险结论，未经人工闭环不得交付")
    document.add_paragraph("供应商已经满足全部交付要求。")
    document.save(report)
    digest = sha256_file(report)
    evidence = tmp_path / "evidence.docx"
    evidence_document = Document()
    evidence_document.add_paragraph("证据留存。")
    evidence_document.save(evidence)
    evidence_digest = sha256_file(evidence)
    job = _high_risk_job()
    job.request = JobRequest(report_path=str(report), sources_path=str(tmp_path))
    job.report_snapshot = str(report)
    job.sources_snapshot = str(tmp_path)
    job.source_inventory = [
        SourceRecord(
            id=f"source-{digest[:12]}",
            role="report",
            original_path=str(report),
            snapshot_path=str(report),
            file_name=report.name,
            file_hash=digest,
            size_bytes=report.stat().st_size,
            status="parsed",
            parsed_chunks=3,
        ),
        SourceRecord(
            id=f"source-{evidence_digest[:12]}",
            role="evidence",
            original_path=str(evidence),
            snapshot_path=str(evidence),
            file_name=evidence.name,
            file_hash=evidence_digest,
            size_bytes=evidence.stat().st_size,
            status="parsed",
            parsed_chunks=1,
        ),
    ]
    job.claims[0].locator = Locator(
        kind="paragraph",
        paragraph=2,
        char_start=0,
        char_end=len(job.claims[0].text),
        anchor_precision="span",
    )
    job.record_decision(
        Decision(
            finding_id="finding-0001",
            action="waive",
            reason="approved continuity risk",
            risk_owner="procurement-director",
        )
    )

    final = export_final(job)

    assert f"revision-{job.decision_revision:06d}-final-" in str(final)
    assert job.artifacts_revision == job.decision_revision
    assert job.artifacts_kind == "final"
    assert job.summary()["delivery_artifact_current"] is True
    manifest = json.loads(Path(job.artifacts["manifest"]).read_text(encoding="utf-8"))
    assert manifest["decision_revision"] == job.decision_revision
    assert manifest["artifact_kind"] == "final"
    assert manifest["artifact_bundle_id"] == job.artifact_bundle_id
    workbook = load_workbook(job.artifacts["claim_ledger"], read_only=True, data_only=True)
    summary_values = {
        row[0]: row[1]
        for row in workbook["审计摘要"].iter_rows(min_row=2, values_only=True)
    }
    workbook.close()
    assert summary_values["产物是否为当前版本"] is True
    assert summary_values["最终交付物是否为当前版本"] is True
    rendered = "\n".join(paragraph.text for paragraph in Document(final).paragraphs)
    assert "已完成人工闭环，可以交付（含正式风险豁免）" in rendered
    assert "自动审计识别" in rendered
    assert "书面豁免" in rendered
    assert "未经人工闭环不得交付" not in rendered


def test_snapshot_creation_rejects_copy_time_content_change(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(data_dir))
    report = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("项目金额为100万元。")
    document.save(report)
    evidence_dir = tmp_path / "evidence"
    evidence_dir.mkdir()
    evidence = evidence_dir / "source.docx"
    source = Document()
    source.add_paragraph("项目金额为100万元。")
    source.save(evidence)
    job = AuditJob(
        id="job-copy-race",
        request=JobRequest(report_path=str(report), sources_path=str(evidence_dir)),
    )
    real_copy2 = shutil.copy2

    def changed_copy(source_path, destination_path):
        result = real_copy2(source_path, destination_path)
        if Path(source_path).resolve() == report.resolve():
            Path(destination_path).write_bytes(Path(destination_path).read_bytes() + b"changed")
        return result

    monkeypatch.setattr("claimledger.ingestion.shutil.copy2", changed_copy)

    with pytest.raises(ValueError, match="changed while its audit snapshot"):
        snapshot_inputs(job)


def test_standard_export_blocks_changed_snapshot_before_manifest(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(data_dir))
    original = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("供应商已经满足全部交付要求。")
    document.save(original)
    snapshot = tmp_path / "report-snapshot.docx"
    shutil.copy2(original, snapshot)
    digest = sha256_file(original)
    job = _high_risk_job()
    job.request = JobRequest(report_path=str(original), sources_path=str(tmp_path))
    job.report_snapshot = str(snapshot)
    job.sources_snapshot = str(tmp_path)
    job.source_inventory = [
        SourceRecord(
            id=f"source-{digest[:12]}",
            role="report",
            original_path=str(original),
            snapshot_path=str(snapshot),
            file_name=original.name,
            file_hash=digest,
            size_bytes=original.stat().st_size,
            status="parsed",
            parsed_chunks=1,
        )
    ]
    changed = Document()
    changed.add_paragraph("快照在审计后被替换，但仍是可打开的Word文件。")
    changed.save(snapshot)

    with pytest.raises(ValueError, match="snapshot hash changed"):
        export_standard_artifacts(job)

    artifact_root = data_dir / "jobs" / job.id / "artifacts"
    if artifact_root.exists():
        assert not list(artifact_root.rglob("artifact_manifest.json"))


def test_decision_api_rejects_client_supplied_audit_metadata(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    job = _high_risk_job()
    store = JobStore(tmp_path / "jobs.sqlite3")
    store.save(job)
    app = create_app(store)
    client = TestClient(app)
    headers = {"Authorization": f"Bearer {app.state.token}"}

    forged = client.post(
        f"/api/v1/jobs/{job.id}/decisions",
        headers=headers,
        json={
            "finding_id": "finding-0001",
            "action": "accept",
            "event_id": "decision-forged",
            "decided_at": "2000-01-01T00:00:00Z",
            "recheck_status": "passed",
            "recheck_message": "forged",
        },
    )

    assert forged.status_code == 422
    accepted = client.post(
        f"/api/v1/jobs/{job.id}/decisions",
        headers=headers,
        json={"finding_id": "finding-0001", "action": "accept"},
    )
    assert accepted.status_code == 200
    decision = accepted.json()["decision"]
    assert decision["event_id"].startswith("decision-")
    assert decision["event_id"] != "decision-forged"
    assert decision["decided_at"] != "2000-01-01T00:00:00Z"
    assert decision["recheck_status"] == "not_required"


def test_cli_and_api_share_strict_waiver_date_validation(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    with pytest.raises(ValueError):
        DecisionRequest(
            finding_id="finding-0001",
            action="waive",
            reason="temporary exception",
            waiver_expires_at="2026/08/31",
        )
    args = SimpleNamespace(
        job_id="missing",
        finding_id="finding-0001",
        action="waive",
        replacement=None,
        reason="temporary exception",
        evidence_id=None,
        actor="local-user",
        risk_owner="risk-owner",
        waiver_expires_at="2026/08/31",
        json=True,
    )
    assert cli_decide(args) == 2
    assert "invalid decision" in capsys.readouterr().out


def test_expired_waiver_relocks_delivery_and_no_expiry_is_explicitly_permanent():
    job = _high_risk_job()
    expired = (date.today() - timedelta(days=1)).isoformat()
    job.record_decision(
        Decision(
            finding_id="finding-0001",
            action="waive",
            reason="temporary exception",
            risk_owner="risk-owner",
            waiver_expires_at=expired,
        )
    )
    assert job.decisions[-1].waiver_scope == "temporary"
    assert job.summary()["expired_waivers"] == 1
    assert job.summary()["unresolved_high"] == 1
    assert job.summary()["delivery_ready"] is False

    job.record_decision(
        Decision(
            finding_id="finding-0001",
            action="waive",
            reason="board-approved permanent exception",
            risk_owner="risk-owner",
        )
    )
    assert job.decisions[-1].waiver_scope == "permanent"
    assert job.decisions[-1].waiver_expires_at is None
    assert job.summary()["expired_waivers"] == 0
    assert job.summary()["unresolved_high"] == 0
    assert job.summary()["delivery_ready"] is True


def test_delivery_artifacts_redact_internal_paths_without_mutating_runtime_job(tmp_path, monkeypatch):
    data_dir = tmp_path / "claimledger-private-data"
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(data_dir))
    job = _exportable_job(tmp_path)
    internal_report = job.request.report_path
    internal_sources = job.request.sources_path
    internal_evidence = job.findings[0].evidence[0].file_path

    outputs = export_standard_artifacts(job)

    audit = json.loads(Path(outputs["audit"]).read_text(encoding="utf-8"))
    manifest = json.loads(Path(outputs["manifest"]).read_text(encoding="utf-8"))
    for payload in (audit, manifest):
        absolute_values = [value for value in _all_strings(payload) if _is_absolute(value)]
        assert absolute_values == []
        serialized = json.dumps(payload, ensure_ascii=False)
        assert str(tmp_path) not in serialized
        assert str(data_dir) not in serialized
    assert audit["request"]["report_path"].startswith("input://report/")
    assert audit["request"]["sources_path"] == "input://evidence"
    assert audit["artifacts"]["audit"].startswith("artifact://")
    assert manifest["inputs"][0]["original_path"].startswith("input://")
    assert manifest["inputs"][0]["snapshot_path"].startswith("snapshot://")

    workbook = load_workbook(outputs["claim_ledger"], read_only=True, data_only=True)
    file_manifest = workbook["文件清单"]
    headers = [cell.value for cell in next(file_manifest.iter_rows())]
    assert "原始路径" not in headers
    assert "快照路径" not in headers
    assert "逻辑引用" in headers
    for row in file_manifest.iter_rows(min_row=2, values_only=True):
        for value in row:
            if isinstance(value, str):
                assert not _is_absolute(value)
                assert str(tmp_path) not in value
    workbook.close()

    # Redaction is a delivery-only transformation; previews still need the
    # internal snapshot paths held in the persisted job.
    assert job.request.report_path == internal_report
    assert job.request.sources_path == internal_sources
    assert job.findings[0].evidence[0].file_path == internal_evidence


@pytest.mark.parametrize(
    ("tamper", "message"),
    [
        ("quote", "evidence quote hash changed"),
        ("quote_hash", "evidence quote hash is missing"),
        ("file_hash", "evidence source hash is not in inventory"),
    ],
)
def test_export_blocks_citation_record_tampering(tmp_path, monkeypatch, tamper, message):
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(tmp_path / "data"))
    job = _exportable_job(tmp_path)
    evidence = job.findings[0].evidence[0]
    if tamper == "quote":
        evidence.quote = "被篡改的引文"
    elif tamper == "quote_hash":
        evidence.quote_hash = None
    else:
        evidence.file_hash = "f" * 64

    with pytest.raises(ValueError, match=message):
        export_standard_artifacts(job)

    artifact_root = tmp_path / "data" / "jobs" / job.id / "artifacts"
    assert not artifact_root.exists()


@pytest.mark.skipif(os.name != "posix", reason="POSIX permission bits are not portable")
def test_private_storage_snapshots_database_and_exports_use_restrictive_modes(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    monkeypatch.setenv("CLAIMLEDGER_DATA_DIR", str(data_dir))
    original_root = tmp_path / "incoming"
    original_root.mkdir()
    report = original_root / "report.docx"
    report_document = Document()
    report_document.add_paragraph("项目金额为100万元。")
    report_document.save(report)
    evidence_dir = original_root / "evidence"
    nested_evidence_dir = evidence_dir / "nested"
    nested_evidence_dir.mkdir(parents=True)
    evidence = nested_evidence_dir / "source.docx"
    evidence_document = Document()
    evidence_document.add_paragraph("证据确认项目金额为100万元，记录已归档。")
    evidence_document.save(evidence)
    snapshot_job = AuditJob(
        id="job-private-snapshot",
        request=JobRequest(report_path=str(report), sources_path=str(evidence_dir)),
    )

    snapshot_inputs(snapshot_job)
    store = JobStore(data_dir / "claimledger.sqlite3")
    store.save(snapshot_job)
    export_job = _exportable_job(tmp_path / "export-case")
    outputs = export_standard_artifacts(export_job)

    controlled_directories = [
        data_dir,
        data_dir / "jobs",
        data_dir / "jobs" / snapshot_job.id,
        data_dir / "jobs" / snapshot_job.id / "inputs",
        data_dir / "jobs" / snapshot_job.id / "inputs" / "report",
        data_dir / "jobs" / snapshot_job.id / "inputs" / "evidence",
        data_dir / "jobs" / snapshot_job.id / "inputs" / "evidence" / "nested",
        Path(outputs["audit"]).parent,
    ]
    for directory in controlled_directories:
        assert directory.stat().st_mode & 0o777 == 0o700
    sensitive_files = [
        data_dir / "claimledger.sqlite3",
        Path(snapshot_job.report_snapshot),
        Path(snapshot_job.sources_snapshot) / "nested" / "source.docx",
        *[Path(path) for path in outputs.values()],
    ]
    for file_path in sensitive_files:
        assert file_path.stat().st_mode & 0o777 == 0o600
