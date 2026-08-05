from __future__ import annotations

import json
import os
import platform
import shutil
import sys
from pathlib import Path, PureWindowsPath
from uuid import uuid4

import yaml
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import __version__
from .config import ensure_private_dir, ensure_private_file, jobs_dir, skill_root
from .engine import rule_path
from .ingestion import (
    verify_citation_integrity,
    verify_input_integrity,
    verify_snapshot_integrity,
)
from .localization import (
    ACTION_ZH,
    COVERAGE_ZH,
    ISSUE_ZH,
    OUTCOME_ZH,
    RECHECK_ZH,
    RELATION_ZH,
    ROLE_ZH,
    SEVERITY_ZH,
    SOURCE_STATUS_ZH,
    STATUS_ZH,
    WAIVER_SCOPE_ZH,
    locator_zh,
    message_zh,
    zh,
)
from .models import ArtifactManifest, AuditJob, Decision, Finding
from .parsers import sha256_file


STATUS_COLORS = {
    "supported": "E2F0D9",
    "partial": "FFF2CC",
    "unsupported": "F4CCCC",
    "conflict": "F8CBAD",
    "stale": "FCE4D6",
    "needs_review": "E4DFEC",
}
NAVY = "1F3864"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="D9D9D9")


def _portable_name(value: str) -> str:
    return PureWindowsPath(value).name if "\\" in value else Path(value).name


def _is_absolute_path(value: str) -> bool:
    try:
        return Path(value).is_absolute() or PureWindowsPath(value).is_absolute()
    except (OSError, ValueError):
        return False


def _logical_source_reference(item, *, snapshot: bool = False) -> str:
    scheme = "snapshot" if snapshot else "input"
    role = "report" if item.role == "report" else "evidence"
    return f"{scheme}://{role}/{item.file_name}"


def _delivery_path_mapping(job: AuditJob) -> dict[str, str]:
    """Map internal paths to stable logical references without mutating a job."""

    mapping: dict[str, str] = {}

    def remember(value: str | None, replacement: str) -> None:
        if value and _is_absolute_path(value):
            mapping[value] = replacement

    report_name = _portable_name(job.request.report_path)
    remember(job.request.report_path, f"input://report/{report_name}")
    remember(job.request.sources_path, "input://evidence")
    remember(job.report_snapshot, f"snapshot://report/{report_name}")
    remember(job.sources_snapshot, "snapshot://evidence")
    for item in job.source_inventory:
        remember(item.original_path, _logical_source_reference(item))
        remember(item.snapshot_path, _logical_source_reference(item, snapshot=True))
    for finding in job.findings:
        for evidence in finding.evidence:
            remember(evidence.file_path, f"evidence://{evidence.file_name}")

    def remember_artifacts(artifacts: dict[str, str], bundle_id: str | None) -> None:
        bundle = bundle_id or "bundle"
        for name, value in artifacts.items():
            file_name = _portable_name(value) or name
            remember(value, f"artifact://{bundle}/{file_name}")

    remember_artifacts(job.artifacts, job.artifact_bundle_id)
    for history in job.artifact_history:
        artifacts = history.get("artifacts")
        if isinstance(artifacts, dict):
            remember_artifacts(artifacts, history.get("bundle_id"))
    return mapping


def _sanitize_delivery_value(value, replacements: list[tuple[str, str]]):
    if isinstance(value, dict):
        return {
            key: _sanitize_delivery_value(item, replacements)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [_sanitize_delivery_value(item, replacements) for item in value]
    if not isinstance(value, str):
        return value
    sanitized = value
    for internal, logical in replacements:
        sanitized = sanitized.replace(internal, logical)
    if sanitized == value and _is_absolute_path(sanitized):
        return f"redacted://{_portable_name(sanitized) or 'path'}"
    return sanitized


def _delivery_payload(job: AuditJob) -> dict:
    mapping = _delivery_path_mapping(job)
    replacements = sorted(mapping.items(), key=lambda item: len(item[0]), reverse=True)
    payload = _sanitize_delivery_value(job.model_dump(mode="json"), replacements)
    report_name = _portable_name(job.request.report_path)
    payload["request"]["report_path"] = f"input://report/{report_name}"
    payload["request"]["sources_path"] = "input://evidence"
    if job.report_snapshot:
        payload["report_snapshot"] = f"snapshot://report/{report_name}"
    if job.sources_snapshot:
        payload["sources_snapshot"] = "snapshot://evidence"
    for public, item in zip(payload["source_inventory"], job.source_inventory, strict=True):
        public["original_path"] = _logical_source_reference(item)
        public["snapshot_path"] = _logical_source_reference(item, snapshot=True)
    for public_finding, finding in zip(payload["findings"], job.findings, strict=True):
        for public_evidence, evidence in zip(
            public_finding["evidence"],
            finding.evidence,
            strict=True,
        ):
            if evidence.file_path:
                public_evidence["file_path"] = f"evidence://{evidence.file_name}"
    return payload


def _safe_cell(value):
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _latest(job: AuditJob, finding_id: str) -> Decision | None:
    return job.latest_decisions().get(finding_id)


def _comment_text(finding: Finding, decision: Decision | None = None) -> str:
    lines = [
        f"[{finding.id}] {zh(STATUS_ZH, finding.status.value)} / {zh(SEVERITY_ZH, finding.severity)}",
        f"问题类型：{', '.join(zh(ISSUE_ZH, item.value) for item in finding.issue_codes) or '无'}",
        message_zh(finding.explanation),
    ]
    for check in finding.checks:
        lines.append(
            f"核对项“{zh(ISSUE_ZH, check.code.value)}”："
            f"{zh(OUTCOME_ZH, check.outcome)}——{message_zh(check.message)}"
        )
        if check.claim_value or check.evidence_value:
            lines.append(f"报告值：{check.claim_value or '-'}｜证据值：{check.evidence_value or '-'}")
    for item in finding.evidence[:3]:
        lines.append(
            f"证据 {item.rank}：{item.file_name}——{locator_zh(item.locator)}——"
            f"{zh(RELATION_ZH, item.relation.value)}——匹配分 {item.score:.2f}"
        )
        lines.append(item.quote[:400])
    if decision:
        lines.append(
            f"最新决定：{zh(ACTION_ZH, decision.action)}；"
            f"操作人：{decision.actor}；时间：{decision.decided_at}"
        )
        if decision.reason:
            lines.append(f"理由：{decision.reason}")
        if decision.replacement_text:
            lines.append(f"修正文本：{decision.replacement_text}")
        if decision.action == "waive":
            expiry = decision.waiver_expires_at or "无到期日（永久豁免）"
            lines.append(f"豁免范围：{zh(WAIVER_SCOPE_ZH, decision.waiver_scope)}；到期日：{expiry}")
    return "\n".join(lines)


def _overlapping_runs(paragraph, start: int | None, end: int | None):
    if start is None or end is None:
        return paragraph.runs or [paragraph.add_run("")]
    runs = []
    offset = 0
    for run in paragraph.runs:
        run_end = offset + len(run.text)
        if offset < end and run_end > start:
            runs.append(run)
        offset = run_end
    return runs or paragraph.runs or [paragraph.add_run("")]


def _target_runs(document: Document, finding: Finding, claim_lookup: dict):
    claim = claim_lookup[finding.claim_id]
    locator = claim.locator
    if locator.kind == "paragraph" and locator.paragraph is not None:
        paragraph = document.paragraphs[locator.paragraph]
        return _overlapping_runs(paragraph, locator.char_start, locator.char_end), locator.anchor_precision
    if locator.kind == "table_cell" and None not in (locator.table, locator.row, locator.column):
        cell = document.tables[locator.table].rows[locator.row].cells[locator.column]
        paragraph = next((item for item in cell.paragraphs if claim.text in item.text), cell.paragraphs[0])
        start = paragraph.text.find(claim.text)
        return _overlapping_runs(paragraph, start if start >= 0 else None, start + len(claim.text) if start >= 0 else None), "cell"
    return [], "paragraph"


def export_annotated(job: AuditJob, destination: Path) -> Path:
    ensure_private_dir(destination.parent)
    source = Path(job.report_snapshot or job.request.report_path)
    document = Document(source)
    claim_lookup = {claim.id: claim for claim in job.claims}
    latest = job.latest_decisions()
    for finding in job.findings:
        runs, _precision = _target_runs(document, finding, claim_lookup)
        if not runs:
            continue
        if hasattr(document, "add_comment"):
            document.add_comment(
                runs=runs,
                text=_comment_text(finding, latest.get(finding.id)),
                author="ClaimLedger",
                initials="CL",
            )
        else:
            runs[-1].add_text(
                f" [ClaimLedger {zh(STATUS_ZH, finding.status.value)}：{message_zh(finding.explanation)}]"
            )
    document.save(destination)
    return ensure_private_file(destination)


def _append(sheet, values) -> None:
    sheet.append([_safe_cell(value) for value in values])


def _style_workbook(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet.max_row > 1 else None
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else None
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
        sheet.row_dimensions[1].height = 30
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column[:100]]
            width = min(56, max(10, (max(values) if values else 10) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        sheet.sheet_view.showGridLines = False


def export_ledger(job: AuditJob, destination: Path) -> Path:
    ensure_private_dir(destination.parent)
    delivery_replacements = sorted(
        _delivery_path_mapping(job).items(),
        key=lambda item: len(item[0]),
        reverse=True,
    )
    workbook = Workbook()
    summary = workbook.active
    summary.title = "审计摘要"
    _append(summary, ["指标", "数值"])
    _append(summary, ["案例名称", job.request.case_name])
    _append(summary, ["任务编号", job.id])
    _append(summary, ["审计基准日", job.request.as_of_date])
    _append(summary, ["规则包", job.request.rule_pack])
    _append(summary, ["运行模式", job.request.profile])
    summary_labels = {
        "supported": "证据充分",
        "partial": "部分支持",
        "unsupported": "缺少证据",
        "conflict": "证据冲突",
        "stale": "证据过期",
        "needs_review": "需要复核",
        "total": "结论总数",
        "decision_events": "人工决定事件数",
        "decision_revision": "决定版本",
        "unresolved_high": "未处理高风险",
        "expired_waivers": "已过期豁免",
        "coverage_status": "证据覆盖状态",
        "artifacts_current": "产物是否为当前版本",
        "delivery_ready": "是否满足交付条件",
        "delivery_artifact_current": "最终交付物是否为当前版本",
    }
    for key, value in job.summary().items():
        if key == "coverage_status":
            value = zh(COVERAGE_ZH, str(value))
        _append(summary, [summary_labels.get(key, key), value])
    _append(summary, ["证据覆盖情况", zh(COVERAGE_ZH, job.coverage_status)])
    _append(summary, ["解析警告", "\n".join(job.warnings)])

    ledger = workbook.create_sheet("结论证据台账")
    _append(
        ledger,
        [
            "发现编号", "结论编号", "审计状态", "风险级别", "问题类型", "报告结论", "报告位置",
            "审计说明", "首要证据文件", "证据位置", "证据关系", "置信度", "处理结果",
        ],
    )
    claims = {claim.id: claim for claim in job.claims}
    latest = job.latest_decisions()
    for finding in job.findings:
        claim = claims[finding.claim_id]
        top = finding.evidence[0] if finding.evidence else None
        decision = latest.get(finding.id)
        _append(
            ledger,
            [
                finding.id,
                claim.id,
                zh(STATUS_ZH, finding.status.value),
                zh(SEVERITY_ZH, finding.severity),
                "、".join(zh(ISSUE_ZH, item.value) for item in finding.issue_codes),
                claim.text,
                locator_zh(claim.locator),
                message_zh(finding.explanation),
                top.file_name if top else "",
                locator_zh(top.locator) if top else "",
                zh(RELATION_ZH, top.relation.value) if top else "",
                finding.confidence,
                zh(ACTION_ZH, decision.action if decision else "unreviewed"),
            ],
        )
        row = ledger.max_row
        ledger.cell(row, 3).fill = PatternFill("solid", fgColor=STATUS_COLORS[finding.status.value])

    actions = workbook.create_sheet("风险处理台账")
    _append(actions, ["发现编号", "风险说明", "当前处理", "修正文本", "处理理由", "责任人", "豁免范围", "到期日", "复验结果", "更新时间"])
    for finding in job.findings:
        if finding.status.value == "supported":
            continue
        decision = latest.get(finding.id)
        _append(
            actions,
            [
                finding.id,
                message_zh(finding.explanation),
                zh(ACTION_ZH, decision.action if decision else "unresolved"),
                decision.replacement_text if decision else "",
                decision.reason if decision else "",
                decision.risk_owner if decision else "",
                zh(WAIVER_SCOPE_ZH, decision.waiver_scope) if decision else "",
                decision.waiver_expires_at if decision else "",
                zh(RECHECK_ZH, decision.recheck_status) if decision else "",
                decision.decided_at if decision else "",
            ],
        )

    checks_sheet = workbook.create_sheet("逐项核对")
    _append(checks_sheet, ["发现编号", "问题类型", "核对结果", "报告值", "证据值", "说明", "证据编号"])
    for finding in job.findings:
        for check in finding.checks:
            _append(
                checks_sheet,
                [
                    finding.id,
                    zh(ISSUE_ZH, check.code.value),
                    zh(OUTCOME_ZH, check.outcome),
                    check.claim_value,
                    check.evidence_value,
                    message_zh(check.message),
                    ", ".join(check.evidence_ids),
                ],
            )

    evidence_sheet = workbook.create_sheet("候选证据")
    _append(
        evidence_sheet,
        [
            "证据编号", "发现编号", "排序", "证据关系", "来源权威级别", "文件", "SHA-256", "来源位置",
            "匹配分", "OCR 置信度", "来源日期", "证据原文", "行或页面上下文",
        ],
    )
    for finding in job.findings:
        for item in finding.evidence:
            _append(
                evidence_sheet,
                [
                    item.id,
                    finding.id,
                    item.rank,
                    zh(RELATION_ZH, item.relation.value),
                    item.authority,
                    item.file_name,
                    item.file_hash,
                    locator_zh(item.locator),
                    item.score,
                    item.confidence,
                    item.source_date,
                    item.quote,
                    item.context_text or "",
                ],
            )

    decisions_sheet = workbook.create_sheet("人工决定日志")
    _append(
        decisions_sheet,
        ["事件编号", "发现编号", "处理动作", "修正文本", "处理理由", "所选证据", "操作人", "风险责任人", "豁免范围", "豁免到期日", "复验结果", "复验说明", "时间"],
    )
    for decision in job.decisions:
        _append(
            decisions_sheet,
            [
                decision.event_id,
                decision.finding_id,
                zh(ACTION_ZH, decision.action),
                decision.replacement_text,
                decision.reason,
                ", ".join(decision.selected_evidence_ids),
                decision.actor,
                decision.risk_owner,
                zh(WAIVER_SCOPE_ZH, decision.waiver_scope),
                decision.waiver_expires_at,
                zh(RECHECK_ZH, decision.recheck_status),
                message_zh(decision.recheck_message),
                decision.decided_at,
            ],
        )

    files_sheet = workbook.create_sheet("文件清单")
    _append(files_sheet, ["来源编号", "文件角色", "逻辑引用", "文件名", "SHA-256", "字节数", "解析状态", "片段数", "OCR 片段数", "低置信度 OCR", "错误信息"])
    for item in job.source_inventory:
        _append(
            files_sheet,
            [
                item.id,
                zh(ROLE_ZH, item.role),
                _logical_source_reference(item),
                item.file_name,
                item.file_hash,
                item.size_bytes,
                zh(SOURCE_STATUS_ZH, item.status),
                item.parsed_chunks,
                item.ocr_chunks,
                item.low_confidence_chunks,
                _sanitize_delivery_value(item.error, delivery_replacements),
            ],
        )

    config_sheet = workbook.create_sheet("审计配置")
    _append(config_sheet, ["配置项", "取值"])
    _append(config_sheet, ["程序版本", __version__])
    _append(config_sheet, ["运行模式", job.request.profile])
    _append(config_sheet, ["规则包", job.request.rule_pack])
    _append(config_sheet, ["审计基准日", job.request.as_of_date])
    _append(config_sheet, ["审阅配置", job.policy_snapshot.profile_name or "未启用"])
    _append(config_sheet, ["岗位模板", job.policy_snapshot.role_template])
    _append(config_sheet, ["策略哈希", job.policy_snapshot.policy_hash or ""])
    _append(config_sheet, ["记忆快照哈希", job.policy_snapshot.memory_snapshot_hash or ""])
    for key, value in job.stage_timings_ms.items():
        _append(config_sheet, [f"阶段耗时 {key}（毫秒）", value])

    _style_workbook(workbook)
    workbook.save(destination)
    return ensure_private_file(destination)


def export_json(job: AuditJob, destination: Path) -> Path:
    ensure_private_dir(destination.parent)
    destination.write_text(
        json.dumps(_delivery_payload(job), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return ensure_private_file(destination)


def _rule_hash(rule_pack: str) -> str:
    paths = [rule_path("generic-zh")]
    if rule_pack != "generic-zh":
        paths.append(rule_path(rule_pack))
    digest = __import__("hashlib").sha256()
    for path in paths:
        digest.update(path.read_bytes())
    return digest.hexdigest()


def export_manifest(job: AuditJob, destination: Path, outputs: dict[str, str]) -> Path:
    ensure_private_dir(destination.parent)
    output_hashes = {name: sha256_file(Path(path)) for name, path in outputs.items() if Path(path).is_file()}
    models_path = skill_root() / "assets" / "models.yaml"
    models = yaml.safe_load(models_path.read_text(encoding="utf-8"))["profiles"].get(job.request.profile, {})
    public_inputs = [
        item.model_copy(
            update={
                "original_path": _logical_source_reference(item),
                "snapshot_path": _logical_source_reference(item, snapshot=True),
            }
        )
        for item in job.source_inventory
    ]
    manifest = ArtifactManifest(
        job_id=job.id,
        decision_revision=job.decision_revision,
        artifact_bundle_id=job.artifact_bundle_id,
        artifact_kind=job.artifacts_kind,
        inputs=public_inputs,
        outputs=output_hashes,
        profile=job.request.profile,
        rule_pack=job.request.rule_pack,
        rule_hash=_rule_hash(job.request.rule_pack),
        models=models,
        runtime={
            "python": sys.version.split()[0],
            "platform": platform.platform(),
            "machine": platform.machine(),
            "claimledger": __version__,
            "stage_timings_ms": job.stage_timings_ms,
            "runtime_metrics": job.runtime_metrics,
        },
        run_parameters={
            "as_of_date": job.request.as_of_date,
            "case_name": job.request.case_name,
            "review_profile": job.request.review_profile,
        },
        review_profile_id=job.policy_snapshot.profile_id,
        role_template_version=job.policy_snapshot.role_version,
        policy_version_hash=job.policy_snapshot.policy_hash,
        memory_snapshot_hash=job.policy_snapshot.memory_snapshot_hash,
        package_version=__version__,
    )
    replacements = sorted(
        _delivery_path_mapping(job).items(),
        key=lambda item: len(item[0]),
        reverse=True,
    )
    public_manifest = _sanitize_delivery_value(
        manifest.model_dump(mode="json"),
        replacements,
    )
    destination.write_text(
        json.dumps(public_manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return ensure_private_file(destination)


def _replace_span(paragraph, start: int, end: int, expected: str, replacement: str) -> None:
    text = paragraph.text
    if start < 0 or end > len(text) or text[start:end] != expected:
        located = text.find(expected)
        if located < 0:
            raise ValueError("precise report span no longer matches the audited claim")
        start, end = located, located + len(expected)
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(text[:start] + replacement + text[end:])
        return
    positions = []
    offset = 0
    for run in runs:
        positions.append((run, offset, offset + len(run.text)))
        offset += len(run.text)
    affected = [(run, left, right) for run, left, right in positions if left < end and right > start]
    if not affected:
        raise ValueError("no Word run overlaps the audited claim span")
    first_run, first_left, _first_right = affected[0]
    last_run, last_left, _last_right = affected[-1]
    first_original = first_run.text
    last_original = last_run.text
    prefix = first_original[: max(0, start - first_left)]
    suffix = last_original[max(0, end - last_left):]
    if first_run is last_run:
        first_run.text = prefix + replacement + suffix
        return
    first_run.text = prefix + replacement
    for run, _left, _right in affected[1:-1]:
        run.text = ""
    last_run.text = suffix


def _apply_replacements(job: AuditJob, destination: Path) -> Path:
    ensure_private_dir(destination.parent)
    source = Path(job.report_snapshot or job.request.report_path)
    document = Document(source)
    claims = {claim.id: claim for claim in job.claims}
    findings = {finding.id: finding for finding in job.findings}
    replacements = []
    for decision in job.latest_decisions().values():
        if decision.action != "replace" or decision.recheck_status != "passed" or not decision.replacement_text:
            continue
        finding = findings.get(decision.finding_id)
        if not finding:
            continue
        claim = claims[finding.claim_id]
        replacements.append((claim, decision.replacement_text))
    replacements.sort(key=lambda item: (item[0].locator.paragraph or -1, item[0].locator.char_start or -1), reverse=True)
    for claim, replacement in replacements:
        locator = claim.locator
        if locator.kind == "paragraph" and locator.paragraph is not None and locator.char_start is not None and locator.char_end is not None:
            paragraph = document.paragraphs[locator.paragraph]
            _replace_span(paragraph, locator.char_start, locator.char_end, claim.text, replacement)
        elif locator.kind == "table_cell" and None not in (locator.table, locator.row, locator.column):
            cell = document.tables[locator.table].rows[locator.row].cells[locator.column]
            paragraph = next((item for item in cell.paragraphs if claim.text in item.text), None)
            if paragraph is None:
                raise ValueError(f"cannot locate claim text in audited table cell: {claim.id}")
            start = paragraph.text.find(claim.text)
            _replace_span(paragraph, start, start + len(claim.text), claim.text, replacement)
        else:
            raise ValueError(f"claim does not have a replacement-safe Word span: {claim.id}")
    _apply_delivery_status(document, job)
    document.save(destination)
    return ensure_private_file(destination)


def _delivery_status_text(job: AuditJob) -> tuple[str, bool]:
    findings = {finding.id: finding for finding in job.findings}
    latest = job.latest_decisions()
    high_decisions = [
        decision
        for finding_id, decision in latest.items()
        if findings.get(finding_id) and findings[finding_id].severity == "high"
    ]
    waived = sum(1 for decision in high_decisions if decision.action == "waive")
    replaced = sum(
        1
        for decision in high_decisions
        if decision.action == "replace" and decision.recheck_status == "passed"
    )
    rejected = sum(1 for decision in high_decisions if decision.action == "reject")
    total_high = sum(1 for finding in job.findings if finding.severity == "high")
    if waived:
        text = (
            "ClaimLedger 交付状态：已完成人工闭环，可以交付（含正式风险豁免）。"
            f"自动审计识别的 {total_high} 项高风险中，{waived} 项由责任人书面豁免，"
            f"{replaced} 项修正并复验通过，{rejected} 项确认为误报；"
            "风险披露、责任人和处理理由保留在随附证据台账与决策日志中。"
        )
        return text, True
    text = (
        "ClaimLedger 交付状态：已完成人工闭环，可以交付（无正式风险豁免）。"
        f"自动审计识别的 {total_high} 项高风险均已闭环："
        f"{replaced} 项修正并复验通过，{rejected} 项确认为误报；"
        "原始风险发现与处理轨迹保留在随附证据台账与决策日志中。"
    )
    return text, False


def _style_delivery_status(paragraph, *, has_waiver: bool) -> None:
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.space_before = Pt(0)
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x7A, 0x3E, 0x00) if has_waiver else RGBColor(0x0B, 0x5D, 0x3B)
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), "FFF2CC" if has_waiver else "E2F0D9")
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), "C88719" if has_waiver else "177245")
    left.set(qn("w:space"), "6")


def _apply_delivery_status(document: Document, job: AuditJob) -> None:
    status_text, has_waiver = _delivery_status_text(job)
    disclosure = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if "未经人工闭环不得交付" in paragraph.text
        ),
        None,
    )
    if disclosure is not None:
        if disclosure.runs:
            disclosure.runs[0].text = status_text
            for run in disclosure.runs[1:]:
                run.text = ""
        else:
            disclosure.add_run(status_text)
        paragraph = disclosure
    else:
        paragraph = document.add_paragraph(status_text)
        if len(document.paragraphs) > 1:
            first = document.paragraphs[0]
            first._p.addprevious(paragraph._p)
    _style_delivery_status(paragraph, has_waiver=has_waiver)


def _export_bundle(job: AuditJob, *, include_final: bool) -> dict[str, str]:
    job.decision_revision = max(job.decision_revision, len(job.decisions))
    job_root = ensure_private_dir(jobs_dir() / job.id)
    output_root = ensure_private_dir(job_root / "artifacts")
    bundle_kind = "final" if include_final else "standard"
    bundle_id = f"revision-{job.decision_revision:06d}-{bundle_kind}-{uuid4().hex[:8]}"
    staging = output_root / f".staging-{bundle_id}"
    published = output_root / bundle_id
    artifact_names = {
        "annotated_report": "annotated_report.docx",
        "claim_ledger": "claim_ledger.xlsx",
        "audit": "audit.json",
        "manifest": "artifact_manifest.json",
    }
    if include_final:
        artifact_names["delivery_report"] = "delivery_report.docx"
    final_paths = {
        name: str(published / file_name)
        for name, file_name in artifact_names.items()
    }
    # Export from a deep copy that already points to the bundle being built.
    # This makes the embedded JSON/Excel status truthful without advertising
    # an incomplete bundle on the persisted job if generation later fails.
    bundle_job = job.model_copy(deep=True)
    bundle_job.invalidate_artifacts(f"superseded by artifact bundle {bundle_id}")
    bundle_job.artifacts = final_paths
    bundle_job.artifacts_revision = bundle_job.decision_revision
    bundle_job.artifacts_kind = bundle_kind
    bundle_job.artifact_bundle_id = bundle_id
    staging.mkdir(parents=True, exist_ok=False, mode=0o700)
    ensure_private_dir(staging)
    try:
        staged = {
            "annotated_report": str(export_annotated(bundle_job, staging / "annotated_report.docx")),
            "claim_ledger": str(export_ledger(bundle_job, staging / "claim_ledger.xlsx")),
        }
        if include_final:
            staged["delivery_report"] = str(_apply_replacements(bundle_job, staging / "delivery_report.docx"))
        export_json(bundle_job, staging / "audit.json")
        manifest_inputs = {name: path for name, path in staged.items()}
        manifest_inputs["audit"] = str(staging / "audit.json")
        export_manifest(bundle_job, staging / "artifact_manifest.json", manifest_inputs)
        staged["audit"] = str(staging / "audit.json")
        staged["manifest"] = str(staging / "artifact_manifest.json")
        bundle_job.artifact_hashes = {
            name: sha256_file(Path(path))
            for name, path in staged.items()
        }
        os.replace(staging, published)
        ensure_private_dir(published)
        for item in published.iterdir():
            if item.is_file():
                ensure_private_file(item)
            elif item.is_dir():
                ensure_private_dir(item)
        job.artifact_history = bundle_job.artifact_history
        job.artifacts = bundle_job.artifacts
        job.artifact_hashes = bundle_job.artifact_hashes
        job.artifacts_revision = bundle_job.artifacts_revision
        job.artifacts_kind = bundle_job.artifacts_kind
        job.artifact_bundle_id = bundle_job.artifact_bundle_id
        return final_paths
    except Exception:
        if staging.exists():
            shutil.rmtree(staging)
        raise


def export_standard_artifacts(job: AuditJob) -> dict[str, str]:
    if job.status != "completed" or not job.findings:
        raise ValueError("audit artifacts require a completed job with findings")
    job.integrity_issues = [
        *verify_snapshot_integrity(job),
        *verify_citation_integrity(job),
    ]
    if job.integrity_issues:
        raise ValueError("audit artifact export blocked: " + "; ".join(job.integrity_issues))
    return _export_bundle(job, include_final=False)


def export_final(job: AuditJob) -> Path:
    if job.status != "completed" or not job.findings:
        raise ValueError("final export requires a completed audit with findings")
    job.integrity_issues = [
        *verify_input_integrity(job),
        *verify_citation_integrity(job),
    ]
    if job.coverage_status != "complete":
        raise ValueError("final export blocked: the evidence pack was not parsed completely")
    if job.integrity_issues:
        raise ValueError("final export blocked: " + "; ".join(job.integrity_issues))
    summary = job.summary()
    if not summary["delivery_ready"]:
        raise ValueError(f"final export blocked: {summary['unresolved_high']} high-severity findings are unresolved")
    outputs = _export_bundle(job, include_final=True)
    return Path(outputs["delivery_report"])
