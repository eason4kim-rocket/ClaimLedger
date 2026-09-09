from __future__ import annotations

import hashlib
import json
import os
import re
import uuid
from pathlib import Path
from typing import Any

import pymupdf as fitz
from docx import Document
from openpyxl import load_workbook
from PIL import Image

from .config import ensure_private_dir, ensure_private_file
from .models import Locator, ParsedChunk


OCR_CACHE_VERSION = "pp-ocrv5-mobile-v1"


SUPPORTED_EVIDENCE = {".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _chunk(
    path: Path,
    file_hash: str,
    text: str,
    locator: Locator,
    confidence: float = 1.0,
    *,
    raw_text: str | None = None,
    context_text: str | None = None,
    metadata: dict[str, Any] | None = None,
) -> ParsedChunk:
    normalized = " ".join(text.split())
    return ParsedChunk(
        id=f"chunk-{uuid.uuid4().hex[:12]}",
        file_path=str(path.resolve()),
        file_name=path.name,
        file_hash=file_hash,
        source_type=path.suffix.lower().lstrip("."),
        text=normalized,
        raw_text=(raw_text if raw_text is not None else text).strip(),
        context_text=" ".join((context_text or text).split()),
        locator=locator,
        confidence=confidence,
        metadata=metadata or {},
    )


def _polygon_bbox(polygon: Any) -> list[float] | None:
    if not polygon:
        return None
    try:
        points = [point for point in polygon if isinstance(point, (list, tuple)) and len(point) >= 2]
        if not points:
            return None
        xs = [float(point[0]) for point in points]
        ys = [float(point[1]) for point in points]
        return [min(xs), min(ys), max(xs), max(ys)]
    except (TypeError, ValueError):
        return None


class OcrEngine:
    def __init__(self, enabled: bool = True):
        self.enabled = enabled
        self._engine = None
        self.error: str | None = None

    def _load(self):
        if not self.enabled:
            self.error = "OCR disabled for this profile"
            return None
        if self._engine is not None or self.error:
            return self._engine
        try:
            paddle_options: dict[str, Any] = {}
            if os.name == "nt":
                # PaddlePaddle 3.3.x currently enables the PIR + oneDNN path on
                # Windows, but PP-OCRv5 contains an ArrayAttribute that this
                # executor cannot convert. Keep the Windows CPU fallback on the
                # stable executor until the upstream implementation lands.
                os.environ.setdefault("FLAGS_enable_pir_api", "0")
                paddle_options["enable_mkldnn"] = False
            from paddleocr import PaddleOCR

            self._engine = PaddleOCR(
                text_detection_model_name="PP-OCRv5_mobile_det",
                text_recognition_model_name="PP-OCRv5_mobile_rec",
                use_doc_orientation_classify=True,
                use_doc_unwarping=False,
                use_textline_orientation=True,
                **paddle_options,
            )
        except Exception as exc:  # optional native dependency
            self.error = f"PaddleOCR unavailable: {type(exc).__name__}: {exc}"
        return self._engine

    def warmup(self) -> None:
        if self._load() is None:
            raise RuntimeError(self.error or "PaddleOCR failed to initialize")

    def read_lines(self, image_path: Path, cache_path: Path | None = None) -> list[dict[str, Any]]:
        if cache_path and cache_path.is_file():
            return json.loads(cache_path.read_text(encoding="utf-8"))
        engine = self._load()
        if engine is None:
            return []
        try:
            with Image.open(image_path) as image:
                width, height = image.size
            result = engine.predict(str(image_path))
            lines: list[dict[str, Any]] = []
            for page in result:
                payload = page.json if hasattr(page, "json") else page
                if callable(payload):
                    payload = payload()
                if isinstance(payload, dict) and "res" in payload:
                    payload = payload["res"]
                if not isinstance(payload, dict):
                    continue
                texts = list(payload.get("rec_texts", []))
                scores = list(payload.get("rec_scores", []))
                polygons = list(payload.get("rec_polys", payload.get("dt_polys", [])))
                for index, text in enumerate(texts):
                    value = str(text).strip()
                    if not value:
                        continue
                    score = float(scores[index]) if index < len(scores) else 0.0
                    polygon = polygons[index] if index < len(polygons) else None
                    lines.append(
                        {
                            "text": value,
                            "confidence": score,
                            "bbox": _polygon_bbox(polygon),
                            "canvas_width": float(width),
                            "canvas_height": float(height),
                        }
                    )
            if cache_path:
                ensure_private_dir(cache_path.parent)
                cache_path.write_text(json.dumps(lines, ensure_ascii=False, indent=2), encoding="utf-8")
                ensure_private_file(cache_path)
            return lines
        except Exception as exc:
            self.error = f"OCR failed: {type(exc).__name__}: {exc}"
            return []

    def read(self, image_path: Path) -> tuple[str, float, list[float] | None]:
        lines = self.read_lines(image_path)
        if not lines:
            return "", 0.0, None
        confidence = sum(item["confidence"] for item in lines) / len(lines)
        boxes = [item["bbox"] for item in lines if item.get("bbox")]
        bbox = None
        if boxes:
            bbox = [min(item[0] for item in boxes), min(item[1] for item in boxes), max(item[2] for item in boxes), max(item[3] for item in boxes)]
        return " ".join(item["text"] for item in lines), confidence, bbox


def parse_docx(path: Path) -> list[ParsedChunk]:
    file_hash = sha256_file(path)
    document = Document(path)
    chunks: list[ParsedChunk] = []
    section = ""
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style_name.startswith("heading") or style_name in {"title", "subtitle"}:
            section = text
        context = f"{section} {text}".strip() if section and section != text else text
        chunks.append(
            _chunk(
                path,
                file_hash,
                text,
                Locator(kind="paragraph", paragraph=index, section=section or None, anchor_precision="paragraph"),
                raw_text=text,
                context_text=context,
                metadata={"style": paragraph.style.name if paragraph.style else None},
            )
        )
    for table_index, table in enumerate(document.tables):
        headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
        for row_index, row in enumerate(table.rows):
            row_values = [cell.text.strip() for cell in row.cells]
            row_context = " | ".join(value for value in row_values if value)
            for column_index, cell in enumerate(row.cells):
                value = cell.text.strip()
                if not value:
                    continue
                header = headers[column_index] if column_index < len(headers) else ""
                context = " | ".join(item for item in [section, header, row_context] if item)
                chunks.append(
                    _chunk(
                        path,
                        file_hash,
                        value,
                        Locator(
                            kind="table_cell",
                            table=table_index,
                            row=row_index,
                            column=column_index,
                            section=section or None,
                            anchor_precision="cell",
                        ),
                        raw_text=value,
                        context_text=context,
                    )
                )
    return chunks


def _ocr_chunks_for_image(
    source_path: Path,
    image_path: Path,
    file_hash: str,
    ocr: OcrEngine,
    cache_path: Path,
    *,
    page: int | None = None,
) -> list[ParsedChunk]:
    lines = ocr.read_lines(image_path, cache_path)
    chunks: list[ParsedChunk] = []
    combined = " ".join(item["text"] for item in lines)
    for item in lines:
        kind = "pdf_page" if page is not None else "image"
        chunks.append(
            _chunk(
                source_path,
                file_hash,
                item["text"],
                Locator(
                    kind=kind,
                    page=page,
                    bbox=item.get("bbox"),
                    canvas_width=item.get("canvas_width"),
                    canvas_height=item.get("canvas_height"),
                    anchor_precision="bbox" if item.get("bbox") else "page",
                ),
                confidence=float(item["confidence"]),
                raw_text=item["text"],
                context_text=combined,
                metadata={"ocr": True},
            )
        )
    return chunks


def parse_pdf(path: Path, ocr: OcrEngine, cache_dir: Path) -> tuple[list[ParsedChunk], list[str]]:
    file_hash = sha256_file(path)
    chunks: list[ParsedChunk] = []
    warnings: list[str] = []
    document = fitz.open(path)
    if document.needs_pass:
        document.close()
        raise ValueError(f"encrypted PDF is not supported: {path.name}")
    for page_index, page in enumerate(document):
        blocks = [block for block in page.get_text("blocks") if len(block) >= 5 and str(block[4]).strip()]
        page_text = " ".join(" ".join(str(block[4]).split()) for block in blocks)
        if len(page_text) >= 20:
            for block in blocks:
                x0, y0, x1, y1, raw = block[:5]
                text = str(raw).strip()
                chunks.append(
                    _chunk(
                        path,
                        file_hash,
                        text,
                        Locator(
                            kind="pdf_page",
                            page=page_index + 1,
                            bbox=[float(x0), float(y0), float(x1), float(y1)],
                            canvas_width=float(page.rect.width),
                            canvas_height=float(page.rect.height),
                            anchor_precision="bbox",
                        ),
                        raw_text=text,
                        context_text=page_text,
                        metadata={"ocr": False},
                    )
                )
            continue
        image_path = cache_dir / f"{file_hash}-page-{page_index + 1}.png"
        if not image_path.exists():
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            pixmap.save(image_path)
            ensure_private_file(image_path)
        cache_path = cache_dir / f"{file_hash}-page-{page_index + 1}-{OCR_CACHE_VERSION}.json"
        found = _ocr_chunks_for_image(path, image_path, file_hash, ocr, cache_path, page=page_index + 1)
        if found:
            # OCR coordinates are in the rendered 2x image coordinate space.
            chunks.extend(found)
        else:
            warnings.append(f"{path.name} page {page_index + 1}: no text layer and {ocr.error or 'OCR returned no text'}")
    document.close()
    return chunks, warnings


def parse_xlsx(path: Path) -> list[ParsedChunk]:
    file_hash = sha256_file(path)
    values_book = load_workbook(path, read_only=True, data_only=True)
    formula_book = load_workbook(path, read_only=True, data_only=False)
    chunks: list[ParsedChunk] = []
    try:
        for values_sheet in values_book.worksheets:
            formula_sheet = formula_book[values_sheet.title]
            headers = {
                cell.column: str(cell.value).strip()
                for cell in next(values_sheet.iter_rows(min_row=1, max_row=1), [])
                if cell.value not in (None, "")
            }
            for row in values_sheet.iter_rows():
                row_values = [str(cell.value).strip() for cell in row if cell.value not in (None, "")]
                if not row_values:
                    continue
                row_key = row_values[0]
                labeled_row_values = []
                for row_cell in row:
                    if row_cell.value in (None, ""):
                        continue
                    label = headers.get(row_cell.column, "")
                    value_text = str(row_cell.value).strip()
                    labeled_row_values.append(f"{label}: {value_text}" if label and row_cell.row > 1 else value_text)
                row_context = " | ".join(labeled_row_values)
                for cell in row:
                    if cell.row == 1:
                        continue
                    formula_cell = formula_sheet[cell.coordinate]
                    value = cell.value
                    formula = formula_cell.value if isinstance(formula_cell.value, str) and formula_cell.value.startswith("=") else None
                    if value in (None, "") and not formula:
                        continue
                    raw = str(value if value not in (None, "") else formula).strip()
                    header = headers.get(cell.column, "")
                    unit_match = re.search(r"[（(]([^()（）]+)[）)]", header)
                    header_unit = unit_match.group(1) if unit_match else None
                    header_unit = {
                        "月": "个月",
                        "日": "天",
                    }.get(header_unit, header_unit)
                    value_with_unit = f"{raw} {header_unit}" if header_unit and not formula else raw
                    if (
                        not unit_match
                        and not formula
                        and any(
                            term in header
                            for term in ("率", "占比", "比例", "增长", "下降", "同比", "环比", "变化", "增幅", "降幅")
                        )
                    ):
                        try:
                            numeric_value = float(raw)
                        except (TypeError, ValueError):
                            numeric_value = None
                        if numeric_value is not None and 0 <= numeric_value <= 1:
                            value_with_unit = f"{numeric_value * 100:g}%"
                    quote = f"{header}: {raw}" if header and cell.row > 1 else raw
                    context = " | ".join(
                        item for item in [values_sheet.title, header, row_key, value_with_unit, row_context, formula or ""] if item
                    )
                    chunks.append(
                        _chunk(
                            path,
                            file_hash,
                            value_with_unit,
                            Locator(
                                kind="sheet_cell",
                                sheet=values_sheet.title,
                                cell=cell.coordinate,
                                anchor_precision="cell",
                            ),
                            raw_text=quote,
                            context_text=context,
                            metadata={"formula": formula, "header": header, "row_key": row_key},
                        )
                    )
    finally:
        values_book.close()
        formula_book.close()
    return chunks


def parse_image(path: Path, ocr: OcrEngine, cache_dir: Path) -> tuple[list[ParsedChunk], list[str]]:
    file_hash = sha256_file(path)
    cache_path = cache_dir / f"{file_hash}-{OCR_CACHE_VERSION}.json"
    chunks = _ocr_chunks_for_image(path, path, file_hash, ocr, cache_path)
    if not chunks:
        return [], [f"{path.name}: {ocr.error or 'OCR returned no text'}"]
    return chunks, []


def parse_evidence_directory(
    path: Path,
    cache_dir: Path,
    *,
    enable_ocr: bool = True,
) -> tuple[list[ParsedChunk], list[str], list[str]]:
    root = path.resolve()
    if not root.is_dir():
        raise ValueError(f"evidence path is not a directory: {root}")
    ensure_private_dir(cache_dir)
    ocr = OcrEngine(enabled=enable_ocr)
    chunks: list[ParsedChunk] = []
    warnings: list[str] = []
    failures: list[str] = []
    files = sorted(item for item in root.rglob("*") if item.is_file())
    for item in files:
        if item.is_symlink():
            message = f"{item.name}: symbolic links are not accepted in evidence packs"
            warnings.append(message)
            failures.append(message)
            continue
        resolved = item.resolve()
        try:
            resolved.relative_to(root)
        except ValueError:
            message = f"{item.name}: path escapes the evidence root"
            warnings.append(message)
            failures.append(message)
            continue
        suffix = item.suffix.lower()
        if suffix not in SUPPORTED_EVIDENCE:
            continue
        before = len(chunks)
        try:
            if suffix == ".docx":
                chunks.extend(parse_docx(item))
            elif suffix == ".pdf":
                parsed, found_warnings = parse_pdf(item, ocr, cache_dir)
                chunks.extend(parsed)
                warnings.extend(found_warnings)
                # A readable page must not hide an unreadable sibling page.
                # Keep the usable chunks, but mark the source pack incomplete
                # so delivery remains blocked until every PDF page is covered.
                if found_warnings:
                    failures.extend(found_warnings)
            elif suffix == ".xlsx":
                chunks.extend(parse_xlsx(item))
            else:
                parsed, found_warnings = parse_image(item, ocr, cache_dir)
                chunks.extend(parsed)
                warnings.extend(found_warnings)
                if found_warnings and not parsed:
                    failures.extend(found_warnings)
            if len(chunks) == before:
                failures.append(f"{item.name}: no readable evidence content")
        except Exception as exc:
            message = f"{item.name}: {type(exc).__name__}: {exc}"
            warnings.append(message)
            failures.append(message)
    if not chunks:
        raise ValueError("no readable evidence content was found")
    return chunks, warnings, failures
