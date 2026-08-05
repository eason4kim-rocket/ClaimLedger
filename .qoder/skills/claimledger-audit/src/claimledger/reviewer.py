from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import ensure_private_dir, ensure_private_file
from .models import AuditJob


RISK_STATUSES = {"partial", "unsupported", "conflict", "stale", "needs_review"}
VALID_LABELS = {"risk", "clean", "uncertain"}


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def export_blind_workbook(job: AuditJob, destination: Path) -> Path:
    ensure_private_dir(destination.parent)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blind review"
    headers = [
        "case_name",
        "claim_id",
        "finding_id",
        "claim_text",
        "report_locator",
        "evidence_1",
        "evidence_2",
        "evidence_3",
        "reviewer_label",
        "issue_codes",
        "notes",
    ]
    sheet.append(headers)
    claims = {item.id: item for item in job.claims}
    for finding in job.findings:
        claim = claims[finding.claim_id]
        evidence = [
            f"{item.file_name} · {item.locator.label()} · {item.quote[:320]}"
            for item in finding.evidence[:3]
        ]
        sheet.append(
            [
                job.request.case_name,
                claim.id,
                finding.id,
                _safe_cell(claim.text),
                claim.locator.label(),
                *[_safe_cell(item) for item in evidence],
                *([""] * (3 - len(evidence))),
                "",
                "",
                "",
            ]
        )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244AA5")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [18, 16, 16, 48, 24, 46, 46, 46, 18, 28, 36]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Allowed reviewer_label", "Meaning"])
    instructions.append(["risk", "The claim has a delivery-relevant evidence problem."])
    instructions.append(["clean", "The claim is adequately supported."])
    instructions.append(["uncertain", "The reviewer cannot decide without escalation."])
    instructions.append(["", "Do not infer labels from formatting; system status is intentionally omitted."])
    workbook.save(destination)
    return ensure_private_file(destination)


def import_blind_workbook(job: AuditJob, source: Path) -> dict:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["Blind review"]
        headers = [str(cell.value or "") for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        indexes = {name: headers.index(name) for name in ("finding_id", "reviewer_label")}
        labels: dict[str, str] = {}
        issue_labels: dict[str, list[str]] = {}
        issue_index = headers.index("issue_codes") if "issue_codes" in headers else None
        for row in sheet.iter_rows(min_row=2, values_only=True):
            finding_id = str(row[indexes["finding_id"]] or "").strip()
            label = str(row[indexes["reviewer_label"]] or "").strip().lower()
            if not finding_id or not label:
                continue
            if label not in VALID_LABELS:
                raise ValueError(f"invalid reviewer_label for {finding_id}: {label}")
            labels[finding_id] = label
            if issue_index is not None:
                issue_labels[finding_id] = [
                    item.strip()
                    for item in str(row[issue_index] or "").split(",")
                    if item.strip()
                ]
    finally:
        workbook.close()

    findings = {item.id: item for item in job.findings}
    unknown = sorted(set(labels) - set(findings))
    if unknown:
        raise ValueError(f"blind workbook contains findings from another job: {', '.join(unknown[:5])}")
    tp = fp = fn = tn = uncertain = agreements = compared = 0
    disagreements = []
    issue_tp = issue_fp = issue_fn = 0
    for finding_id, label in labels.items():
        if label == "uncertain":
            uncertain += 1
            continue
        predicted_risk = findings[finding_id].status.value in RISK_STATUSES
        human_risk = label == "risk"
        compared += 1
        agreements += int(predicted_risk == human_risk)
        if predicted_risk and human_risk:
            tp += 1
        elif predicted_risk and not human_risk:
            fp += 1
        elif not predicted_risk and human_risk:
            fn += 1
        else:
            tn += 1
        if predicted_risk != human_risk:
            disagreements.append(
                {
                    "finding_id": finding_id,
                    "system": "risk" if predicted_risk else "clean",
                    "reviewer": label,
                }
            )
        expected = set(issue_labels.get(finding_id, []))
        if expected:
            actual = {item.value for item in findings[finding_id].issue_codes}
            issue_tp += len(expected & actual)
            issue_fp += len(actual - expected)
            issue_fn += len(expected - actual)

    def ratio(left: int, right: int) -> float:
        return round(left / right, 4) if right else 0.0

    precision = ratio(tp, tp + fp)
    recall = ratio(tp, tp + fn)
    f1 = round(2 * precision * recall / (precision + recall), 4) if precision + recall else 0.0
    issue_precision = ratio(issue_tp, issue_tp + issue_fp)
    issue_recall = ratio(issue_tp, issue_tp + issue_fn)
    issue_f1 = (
        round(2 * issue_precision * issue_recall / (issue_precision + issue_recall), 4)
        if issue_precision + issue_recall
        else 0.0
    )
    return {
        "job_id": job.id,
        "labeled": len(labels),
        "compared": compared,
        "uncertain": uncertain,
        "agreement": ratio(agreements, compared),
        "confusion": {"tp": tp, "fp": fp, "fn": fn, "tn": tn},
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "issue_code_precision": issue_precision,
        "issue_code_recall": issue_recall,
        "issue_code_f1": issue_f1,
        "disagreements": disagreements,
    }


def write_blind_metrics(payload: dict, destination: Path) -> Path:
    ensure_private_dir(destination.parent)
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return ensure_private_file(destination)
