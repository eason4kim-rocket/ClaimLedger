from __future__ import annotations

import html
import io
import json
import mimetypes
import secrets
from pathlib import Path

import pymupdf as fitz
from docx import Document
from fastapi import BackgroundTasks, Cookie, Depends, FastAPI, Header, HTTPException, Query
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, Response
from openpyxl import load_workbook

from . import __version__
from .config import api_token, jobs_dir
from .engine import new_job, recheck_replacement, run_audit
from .exporters import export_final, export_standard_artifacts
from .localization import (
    ACTION_ZH,
    COVERAGE_ZH,
    ISSUE_ZH,
    OUTCOME_ZH,
    RECHECK_ZH,
    RELATION_ZH,
    SEVERITY_ZH,
    STATUS_ZH,
    WAIVER_SCOPE_ZH,
    locator_zh,
    message_zh,
    zh,
)
from .models import (
    AuditJob,
    Decision,
    DecisionRequest,
    Evidence,
    JobRequest,
    Locator,
    PolicyDraftRequest,
    PolicyTransitionRequest,
    ReviewerProfile,
    ReviewerProfileRequest,
)
from .parsers import sha256_file
from .storage import JobStore


SESSION_COOKIE = "claimledger_session"
ARTIFACT_FILENAMES = {
    "annotated_report": "annotated_report.docx",
    "claim_ledger": "claim_ledger.xlsx",
    "audit": "audit.json",
    "manifest": "artifact_manifest.json",
    "delivery_report": "delivery_report.docx",
}


def _security_headers(*, script_nonce: str | None = None) -> dict[str, str]:
    script_policy = f"'nonce-{script_nonce}'" if script_nonce else "'none'"
    return {
        "Cache-Control": "no-store",
        "Content-Security-Policy": (
            "default-src 'none'; "
            f"script-src {script_policy}; "
            "style-src 'unsafe-inline'; "
            "img-src 'self' data:; "
            "frame-src 'self'; "
            "connect-src 'self'; "
            "base-uri 'none'; "
            "form-action 'none'; "
            "frame-ancestors 'none'"
        ),
        "Referrer-Policy": "no-referrer",
        "X-Content-Type-Options": "nosniff",
    }


def _binary_security_headers() -> dict[str, str]:
    return {
        "Cache-Control": "no-store",
        "Referrer-Policy": "no-referrer",
        "X-Content-Type-Options": "nosniff",
    }


def create_app(store: JobStore | None = None) -> FastAPI:
    app = FastAPI(title="ClaimLedger", version=__version__, docs_url="/api/docs")
    app.state.store = store or JobStore()
    app.state.token = api_token()

    def authorize(
        authorization: str | None = Header(default=None),
        token: str | None = Query(default=None),
        claimledger_session: str | None = Cookie(default=None, alias=SESSION_COOKIE),
    ) -> None:
        header_value = authorization if isinstance(authorization, str) else ""
        cookie_value = claimledger_session if isinstance(claimledger_session, str) else ""
        query_value = token if isinstance(token, str) else ""
        supplied = (
            header_value.removeprefix("Bearer ").strip()
            or cookie_value
            or query_value
        )
        if supplied != app.state.token:
            raise HTTPException(status_code=401, detail="本地会话令牌无效")

    def process(job_id: str) -> None:
        job = app.state.store.get(job_id)
        if not job:
            return
        try:
            job, _chunks = run_audit(job, app.state.store)
            export_standard_artifacts(job)
        except Exception as exc:
            job.status = "failed"
            job.error = f"{type(exc).__name__}: {exc}"
        app.state.store.save(job)

    @app.get("/api/v1/health")
    def health() -> dict:
        return {"status": "ok", "service": "claimledger", "version": __version__}

    @app.get("/api/v1/profiles", dependencies=[Depends(authorize)])
    def list_profiles() -> dict:
        return {
            "items": [
                item.model_dump(mode="json")
                for item in app.state.store.list_profiles()
            ]
        }

    @app.post("/api/v1/profiles", dependencies=[Depends(authorize)])
    def create_profile(request: ReviewerProfileRequest) -> dict:
        try:
            profile = app.state.store.create_profile(
                ReviewerProfile(**request.model_dump())
            )
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        return profile.model_dump(mode="json")

    @app.get("/api/v1/profiles/{profile_id}", dependencies=[Depends(authorize)])
    def get_profile(profile_id: str) -> dict:
        profile = app.state.store.get_profile(profile_id)
        if not profile:
            raise HTTPException(status_code=404, detail="未找到审阅配置")
        payload = profile.model_dump(mode="json")
        payload["policies"] = [
            item.model_dump(mode="json")
            for item in app.state.store.list_policies(profile.id)
        ]
        return payload

    @app.delete("/api/v1/profiles/{profile_id}", dependencies=[Depends(authorize)])
    def delete_profile(profile_id: str) -> dict:
        if not app.state.store.delete_profile(profile_id):
            raise HTTPException(status_code=404, detail="未找到审阅配置")
        return {"deleted": True, "profile": profile_id}

    @app.get(
        "/api/v1/profiles/{profile_id}/memories",
        dependencies=[Depends(authorize)],
    )
    def list_memories(profile_id: str, include_expired: bool = False) -> dict:
        profile = app.state.store.get_profile(profile_id)
        if not profile:
            raise HTTPException(status_code=404, detail="未找到审阅配置")
        return {
            "items": [
                item.model_dump(mode="json")
                for item in app.state.store.list_memories(
                    profile.id,
                    include_expired=include_expired,
                )
            ]
        }

    @app.get(
        "/api/v1/profiles/{profile_id}/policies",
        dependencies=[Depends(authorize)],
    )
    def list_policies(profile_id: str) -> dict:
        profile = app.state.store.get_profile(profile_id)
        if not profile:
            raise HTTPException(status_code=404, detail="未找到审阅配置")
        return {
            "items": [
                item.model_dump(mode="json")
                for item in app.state.store.list_policies(profile.id)
            ]
        }

    @app.post(
        "/api/v1/profiles/{profile_id}/policies",
        dependencies=[Depends(authorize)],
    )
    def create_policy(profile_id: str, request: PolicyDraftRequest) -> dict:
        profile = app.state.store.get_profile(profile_id)
        if not profile:
            raise HTTPException(status_code=404, detail="未找到审阅配置")
        memory = app.state.store.get_memory(request.source_memory_id)
        if not memory or memory.profile_id != profile.id:
            raise HTTPException(status_code=400, detail="该历史先例不属于当前审阅配置")
        from .memory import draft_policy

        try:
            policy = draft_policy(
                app.state.store,
                request.source_memory_id,
                actor=request.actor,
                reason=request.reason,
                kind=request.kind,
            )
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        return policy.model_dump(mode="json")

    @app.post("/api/v1/policies/{policy_id}/activate", dependencies=[Depends(authorize)])
    def activate_policy(policy_id: str, request: PolicyTransitionRequest) -> dict:
        from .memory import transition_policy

        try:
            policy = transition_policy(
                app.state.store,
                policy_id,
                target="active",
                actor=request.actor,
                reason=request.reason,
            )
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        return policy.model_dump(mode="json")

    @app.post("/api/v1/policies/{policy_id}/retire", dependencies=[Depends(authorize)])
    def retire_policy(policy_id: str, request: PolicyTransitionRequest) -> dict:
        from .memory import transition_policy

        try:
            policy = transition_policy(
                app.state.store,
                policy_id,
                target="retired",
                actor=request.actor,
                reason=request.reason,
            )
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        return policy.model_dump(mode="json")

    @app.post("/api/v1/jobs", dependencies=[Depends(authorize)])
    def create_job(request: JobRequest, tasks: BackgroundTasks) -> dict:
        try:
            job = new_job(request, app.state.store)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        app.state.store.save(job)
        tasks.add_task(process, job.id)
        return {"job_id": job.id, "status": job.status, "review_url": f"/review/{job.id}?token={app.state.token}"}

    @app.get("/api/v1/jobs/{job_id}", dependencies=[Depends(authorize)])
    def get_job(job_id: str) -> dict:
        job = _require_job(app.state.store, job_id)
        payload = job.model_dump(mode="json", exclude={"claims", "findings"})
        payload["summary"] = job.summary()
        return payload

    @app.get("/api/v1/jobs/{job_id}/claims", dependencies=[Depends(authorize)])
    def get_claims(job_id: str) -> dict:
        job = _require_job(app.state.store, job_id)
        claims = {claim.id: claim for claim in job.claims}
        latest = job.latest_decisions()
        return {
            "items": [
                {
                    "claim": claims[finding.claim_id].model_dump(mode="json"),
                    "finding": finding.model_dump(mode="json"),
                    "latest_decision": latest.get(finding.id).model_dump(mode="json") if latest.get(finding.id) else None,
                }
                for finding in job.findings
            ]
        }

    @app.post("/api/v1/jobs/{job_id}/decisions", dependencies=[Depends(authorize)])
    def decide(job_id: str, request: DecisionRequest) -> dict:
        decision = request.to_decision()

        def apply(current: AuditJob) -> dict:
            finding = next((item for item in current.findings if item.id == decision.finding_id), None)
            if not finding:
                raise HTTPException(status_code=400, detail="该风险发现不属于当前任务")
            valid_evidence_ids = {item.id for item in finding.evidence}
            if not set(decision.selected_evidence_ids).issubset(valid_evidence_ids):
                raise HTTPException(status_code=400, detail="所选证据不属于当前风险发现")
            if decision.action == "replace" and decision.replacement_text:
                passed, message = recheck_replacement(current, decision.finding_id, decision.replacement_text)
                decision.recheck_status = "passed" if passed else "failed"
                decision.recheck_message = message
            current.record_decision(decision)
            return {"decision": decision.model_dump(mode="json"), "summary": current.summary()}

        updated = app.state.store.mutate(job_id, apply)
        if updated is None:
            raise HTTPException(status_code=404, detail="未找到审计任务")
        if updated[0].policy_snapshot.profile_id:
            from .memory import memory_from_decision

            finding = next(item for item in updated[0].findings if item.id == decision.finding_id)
            claim = next(item for item in updated[0].claims if item.id == finding.claim_id)
            event = memory_from_decision(updated[0], finding, claim, decision)
            if event:
                app.state.store.add_memory(event)
        return updated[1]

    @app.post("/api/v1/jobs/{job_id}/exports", dependencies=[Depends(authorize)])
    def export(job_id: str, final: bool = False) -> dict:
        def build(current: AuditJob) -> dict:
            if final:
                export_final(current)
            else:
                export_standard_artifacts(current)
            return {
                "available_artifacts": sorted(current.artifacts),
                "artifact_bundle_id": current.artifact_bundle_id,
                "artifact_kind": current.artifacts_kind,
                "summary": current.summary(),
            }

        try:
            updated = app.state.store.mutate(job_id, build)
            if updated is None:
                raise HTTPException(status_code=404, detail="未找到审计任务")
            return updated[1]
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @app.get(
        "/api/v1/jobs/{job_id}/artifacts/{kind}",
        dependencies=[Depends(authorize)],
    )
    def download_artifact(job_id: str, kind: str) -> FileResponse:
        job = _require_job(app.state.store, job_id)
        source = _resolve_artifact_path(job, kind)
        media_type = mimetypes.guess_type(source.name)[0] or "application/octet-stream"
        return FileResponse(
            source,
            filename=source.name,
            media_type=media_type,
            content_disposition_type="attachment",
            headers=_binary_security_headers(),
        )

    @app.get("/review/{job_id}", response_class=HTMLResponse)
    def review(
        job_id: str,
        token: str | None = Query(default=None),
        claimledger_session: str | None = Cookie(default=None, alias=SESSION_COOKIE),
    ) -> Response:
        if token is not None:
            authorize(token=token)
            response = RedirectResponse(
                url=f"/review/{job_id}",
                status_code=303,
                headers=_binary_security_headers(),
            )
            response.set_cookie(
                SESSION_COOKIE,
                app.state.token,
                httponly=True,
                samesite="strict",
                secure=False,
                path="/",
            )
            return response
        authorize(claimledger_session=claimledger_session)
        nonce = secrets.token_urlsafe(18)
        content = render_review(
            _require_job(app.state.store, job_id),
            script_nonce=nonce,
            store=app.state.store,
        )
        return HTMLResponse(content=content, headers=_security_headers(script_nonce=nonce))

    @app.get(
        "/review/{job_id}/evidence/{evidence_id}",
        response_class=HTMLResponse,
        dependencies=[Depends(authorize)],
    )
    def evidence_view(job_id: str, evidence_id: str) -> HTMLResponse:
        job = _require_job(app.state.store, job_id)
        evidence = _find_evidence(job, evidence_id)
        source = _resolve_evidence_path(job, evidence)
        raw_url = f"/review/{job.id}/evidence/{evidence.id}/raw"
        page_url = f"/review/{job.id}/evidence/{evidence.id}/page.png"
        content = render_source_view(source, evidence.locator, evidence.quote, raw_url, "原始证据", page_url)
        return HTMLResponse(content=content, headers=_security_headers())

    @app.get(
        "/review/{job_id}/evidence/{evidence_id}/raw",
        dependencies=[Depends(authorize)],
    )
    def evidence_raw(job_id: str, evidence_id: str) -> FileResponse:
        job = _require_job(app.state.store, job_id)
        evidence = _find_evidence(job, evidence_id)
        source = _resolve_evidence_path(job, evidence)
        media_type = mimetypes.guess_type(source.name)[0] or "application/octet-stream"
        return FileResponse(
            source,
            media_type=media_type,
            content_disposition_type="inline",
            headers=_binary_security_headers(),
        )

    @app.get(
        "/review/{job_id}/evidence/{evidence_id}/page.png",
        dependencies=[Depends(authorize)],
    )
    def evidence_page_image(job_id: str, evidence_id: str) -> Response:
        job = _require_job(app.state.store, job_id)
        evidence = _find_evidence(job, evidence_id)
        source = _resolve_evidence_path(job, evidence)
        if source.suffix.lower() != ".pdf":
            raise HTTPException(status_code=400, detail="该来源文件不是 PDF")
        document = fitz.open(source)
        try:
            index = max(0, (evidence.locator.page or 1) - 1)
            pixmap = document[index].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            return Response(
                content=pixmap.tobytes("png"),
                media_type="image/png",
                headers=_binary_security_headers(),
            )
        finally:
            document.close()

    @app.get(
        "/review/{job_id}/report/{finding_id}",
        response_class=HTMLResponse,
        dependencies=[Depends(authorize)],
    )
    def report_view(job_id: str, finding_id: str) -> HTMLResponse:
        job = _require_job(app.state.store, job_id)
        finding = next((item for item in job.findings if item.id == finding_id), None)
        if not finding:
            raise HTTPException(status_code=404, detail="未找到风险发现")
        claim = next((item for item in job.claims if item.id == finding.claim_id), None)
        if not claim:
            raise HTTPException(status_code=404, detail="未找到报告结论")
        report_record = next((item for item in job.source_inventory if item.role == "report"), None)
        report = Path(report_record.snapshot_path if report_record else (job.report_snapshot or job.request.report_path)).resolve()
        if not report.is_file() or (report_record and sha256_file(report) != report_record.file_hash):
            raise HTTPException(status_code=404, detail="待审报告快照不可用或文件哈希已变化")
        content = render_source_view(report, claim.locator, claim.text, None, "报告原文")
        return HTMLResponse(content=content, headers=_security_headers())

    return app


def _require_job(store: JobStore, job_id: str) -> AuditJob:
    job = store.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="未找到审计任务")
    return job


def _find_evidence(job: AuditJob, evidence_id: str) -> Evidence:
    for finding in job.findings:
        for evidence in finding.evidence:
            if evidence.id == evidence_id:
                return evidence
    raise HTTPException(status_code=404, detail="未找到证据")


def _resolve_artifact_path(job: AuditJob, kind: str) -> Path:
    expected_name = ARTIFACT_FILENAMES.get(kind)
    if not expected_name:
        raise HTTPException(status_code=404, detail="不支持该交付物类型")
    if kind == "delivery_report" and job.artifacts_kind != "final":
        raise HTTPException(status_code=404, detail="最终交付稿尚未生成")
    configured = job.artifacts.get(kind)
    expected_hash = job.artifact_hashes.get(kind)
    if not configured:
        raise HTTPException(status_code=404, detail="交付物尚不可用")
    if not expected_hash:
        raise HTTPException(
            status_code=409,
            detail="该交付物生成于下载校验功能之前，请重新导出任务",
        )
    if not job.artifact_bundle_id:
        raise HTTPException(status_code=409, detail="交付包缺少唯一标识")
    expected_root = (
        jobs_dir()
        / job.id
        / "artifacts"
        / job.artifact_bundle_id
    ).resolve()
    candidate = Path(configured).expanduser().resolve()
    try:
        candidate.relative_to(expected_root)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail="交付物路径超出当前任务目录") from exc
    if candidate.name != expected_name or candidate.parent != expected_root:
        raise HTTPException(status_code=409, detail="交付物路径与声明类型不一致")
    if not candidate.is_file() or candidate.is_symlink():
        raise HTTPException(status_code=404, detail="交付物文件不可用")
    if sha256_file(candidate) != expected_hash:
        raise HTTPException(status_code=409, detail="交付物导出后哈希发生变化")
    return candidate


def _resolve_evidence_path(job: AuditJob, evidence: Evidence) -> Path:
    source_root = Path(job.sources_snapshot or job.request.sources_path).expanduser().resolve()
    candidates: list[Path] = []
    if evidence.file_path:
        candidates.append(Path(evidence.file_path).expanduser().resolve())
    if source_root.is_dir():
        candidates.extend(item.resolve() for item in source_root.rglob(evidence.file_name) if item.is_file())
    for candidate in candidates:
        try:
            candidate.relative_to(source_root)
        except ValueError:
            continue
        if candidate.is_file() and not candidate.is_symlink() and sha256_file(candidate) == evidence.file_hash:
            return candidate
    raise HTTPException(status_code=404, detail="证据快照不可用或文件哈希已变化")


def _base_view(title: str, body: str, source: Path, locator: Locator, quote: str) -> str:
    return f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width"><title>{html.escape(title)}</title><style>
body{{margin:0;background:#f3f5f8;color:#172033;font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}}
main{{max-width:1180px;margin:auto;padding:28px}} .card{{background:white;border:1px solid #dfe3ea;border-radius:14px;padding:20px}}
.meta{{color:#687085}} mark{{background:#fff1a8;padding:2px 4px}} iframe{{width:100%;height:76vh;border:1px solid #dfe3ea}}
.image-stage{{position:relative;display:inline-block;max-width:100%}} .image-stage img{{display:block;max-width:100%;height:auto}}
.ocr-box{{position:absolute;border:3px solid #e23b3b;background:#e23b3b20;box-sizing:border-box}}
table{{border-collapse:collapse;width:100%}} td{{border:1px solid #dfe3ea;padding:7px}} td.hit{{background:#fff1a8;font-weight:700}}
.context{{padding:10px;border-left:3px solid #3157d5;margin:8px 0}} .focus{{background:#fff1a8}}
</style></head><body><main><h1>{html.escape(title)}</h1><p class="meta">{html.escape(source.name)} · {html.escape(locator_zh(locator))}</p>
<div class="card">{body}</div><h2>审计引用</h2><blockquote>{html.escape(quote)}</blockquote></main></body></html>"""


def _boxed_image(url: str, locator: Locator, *, pdf_scale: float = 1.0) -> str:
    if not locator.bbox or not locator.canvas_width or not locator.canvas_height:
        return f'<img style="max-width:100%;height:auto" src="{html.escape(url)}" alt="来源文件预览">'
    x0, y0, x1, y1 = locator.bbox
    width = locator.canvas_width
    height = locator.canvas_height
    left, top = 100 * x0 / width, 100 * y0 / height
    box_width, box_height = 100 * (x1 - x0) / width, 100 * (y1 - y0) / height
    return (
        f'<div class="image-stage"><img src="{html.escape(url)}" alt="来源文件预览">'
        f'<span class="ocr-box" style="left:{left:.3f}%;top:{top:.3f}%;width:{box_width:.3f}%;height:{box_height:.3f}%"></span></div>'
    )


def render_source_view(
    source: Path,
    locator: Locator,
    quote: str,
    raw_url: str | None,
    title: str,
    page_image_url: str | None = None,
) -> str:
    suffix = source.suffix.lower()
    if suffix == ".pdf" and raw_url:
        if locator.bbox and page_image_url:
            body = _boxed_image(page_image_url, locator)
        else:
            body = f'<iframe src="{html.escape(raw_url)}#page={locator.page or 1}"></iframe>'
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"} and raw_url:
        body = _boxed_image(raw_url, locator)
    elif suffix == ".xlsx":
        body = _render_xlsx_context(source, locator)
    elif suffix == ".docx":
        body = _render_docx_context(source, locator)
    else:
        body = "<p>暂不支持预览此类来源文件，请在本机使用对应软件打开。</p>"
    return _base_view(title, body, source, locator, quote)


def _render_docx_context(source: Path, locator: Locator) -> str:
    document = Document(source)
    if locator.kind == "table_cell" and None not in (locator.table, locator.row, locator.column):
        table = document.tables[locator.table]
        rows = []
        for row_index, row in enumerate(table.rows):
            cells = []
            for column_index, cell in enumerate(row.cells):
                css = ' class="hit"' if (row_index, column_index) == (locator.row, locator.column) else ""
                cells.append(f"<td{css}>{html.escape(cell.text)}</td>")
            rows.append(f"<tr>{''.join(cells)}</tr>")
        return f"<table>{''.join(rows)}</table>"
    index = locator.paragraph or 0
    start, end = max(0, index - 2), min(len(document.paragraphs), index + 3)
    rendered = []
    for item in range(start, end):
        text = document.paragraphs[item].text
        if item == index and locator.char_start is not None and locator.char_end is not None:
            text_html = (
                html.escape(text[:locator.char_start])
                + "<mark>"
                + html.escape(text[locator.char_start:locator.char_end])
                + "</mark>"
                + html.escape(text[locator.char_end:])
            )
        else:
            text_html = html.escape(text)
        rendered.append(f'<p class="context{" focus" if item == index else ""}">{text_html}</p>')
    return "".join(rendered)


def _render_xlsx_context(source: Path, locator: Locator) -> str:
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook[locator.sheet] if locator.sheet in workbook.sheetnames else workbook.active
        target = sheet[locator.cell or "A1"]
        min_row, max_row = max(1, target.row - 2), min(sheet.max_row, target.row + 2)
        min_col, max_col = max(1, target.column - 2), min(sheet.max_column, target.column + 4)
        rows = []
        for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            cells = []
            for cell in row:
                css = ' class="hit"' if cell.coordinate == target.coordinate else ""
                cells.append(f"<td{css}>{html.escape(str(cell.value) if cell.value is not None else '')}</td>")
            rows.append(f"<tr>{''.join(cells)}</tr>")
        return f"<p><strong>{html.escape(sheet.title)}!{html.escape(target.coordinate)}</strong></p><table>{''.join(rows)}</table>"
    finally:
        workbook.close()


def _artifact_download_panel(job: AuditJob) -> str:
    labels = {
        "annotated_report": "下载带批注报告",
        "claim_ledger": "下载证据台账",
        "audit": "下载审计 JSON",
        "manifest": "下载产物清单",
        "delivery_report": "下载最终交付稿",
    }
    links = [
        (
            f'<a class="artifact-link" href="/api/v1/jobs/{html.escape(job.id, quote=True)}'
            f'/artifacts/{kind}">{label}</a>'
        )
        for kind, label in labels.items()
        if kind in job.artifacts
    ]
    if not links:
        return '<p class="artifact-empty">当前决定版本尚无有效产物；完成复核后生成新的交付包。</p>'
    return '<div class="artifact-links">' + "".join(links) + "</div>"


def render_review(
    job: AuditJob,
    *,
    script_nonce: str | None = None,
    store: JobStore | None = None,
) -> str:
    claims = {claim.id: claim for claim in job.claims}
    latest = job.latest_decisions()
    histories: dict[str, list[Decision]] = {}
    for decision in job.decisions:
        histories.setdefault(decision.finding_id, []).append(decision)
    rows: list[str] = []
    for finding in job.findings:
        claim = claims[finding.claim_id]
        decision = latest.get(finding.id)
        resolved = job.decision_resolves(decision)
        evidence_cards = []
        for item in finding.evidence[:5]:
            context = (
                f'<details><summary>查看行/页上下文</summary><p class="context">{html.escape(item.context_text)}</p></details>'
                if item.context_text and item.context_text != item.quote
                else ""
            )
            evidence_cards.append(
                f'<div class="evidence"><label><input type="checkbox" name="evidence-{finding.id}" value="{html.escape(item.id)}"> '
                f'<strong>#{item.rank} {html.escape(item.file_name)}</strong> · {html.escape(locator_zh(item.locator))} · '
                f'{html.escape(zh(RELATION_ZH, item.relation.value))} · 匹配分 {item.score:.2f}</label>'
                f'<blockquote>{html.escape(item.quote)}</blockquote>'
                f'{context}'
                f'<a class="jump" target="_blank" href="/review/{job.id}/evidence/{item.id}">打开原始证据定位 ↗</a></div>'
            )
        checks = "".join(
            f"<tr><td>{html.escape(zh(ISSUE_ZH, item.code.value))}</td>"
            f"<td>{html.escape(zh(OUTCOME_ZH, item.outcome))}</td><td>{html.escape(item.claim_value or '')}</td>"
            f"<td>{html.escape(item.evidence_value or '')}</td><td>{html.escape(message_zh(item.message))}</td></tr>"
            for item in finding.checks
        )
        history_items = []
        for item in histories.get(finding.id, []):
            waiver_detail = ""
            if item.action == "waive":
                expiry = item.waiver_expires_at or "无到期日（永久豁免）"
                waiver_detail = f" · {zh(WAIVER_SCOPE_ZH, item.waiver_scope)} · {expiry}"
            history_items.append(
                f"<li>{html.escape(item.decided_at)} · {html.escape(item.actor)} · "
                f"<strong>{html.escape(zh(ACTION_ZH, item.action))}</strong> · "
                f"{html.escape(item.reason or message_zh(item.recheck_message) or '')}"
                f"{html.escape(waiver_detail)}</li>"
            )
        history = "".join(history_items) or "<li>尚未人工处理</li>"
        memory_cards = []
        for match in finding.memory_matches:
            factors = " · ".join(match.match_factors)
            draft_button = (
                f'<button type="button" data-policy-draft="{html.escape(match.memory_id)}">'
                "生成策略草稿</button>"
                if match.action != "waive" and job.policy_snapshot.profile_id
                else ""
            )
            warning = (
                "豁免仅供回看，绝不作为事实或支持证据。"
                if match.action == "waive"
                else "历史先例仅供建议，不改变当前审计状态。"
            )
            memory_cards.append(
                f'<div class="memory"><strong>{html.escape(zh(ACTION_ZH, match.action))} · {match.score:.2f}</strong>'
                f'<p>{html.escape(match.claim_excerpt)}</p>'
                f'<small>{html.escape(factors)} · {html.escape(match.created_at)}</small>'
                f'<p class="memory-warning">{html.escape(warning)}</p>{draft_button}</div>'
            )
        memory_section = (
            f"<details open><summary>相似历史先例（{len(memory_cards)}）</summary>"
            f"{''.join(memory_cards)}</details>"
            if memory_cards
            else "<details><summary>相似历史先例（0）</summary><p>没有达到安全阈值的历史先例。</p></details>"
        )
        rows.append(
            f"""
            <article id="{html.escape(finding.id, quote=True)}" class="finding" data-status="{finding.status.value}" data-severity="{finding.severity}" data-resolved="{'yes' if resolved else 'no'}" data-suggested-replacement="{html.escape(finding.suggested_replacement or '', quote=True)}">
              <header><span class="badge {finding.status.value}">{html.escape(zh(STATUS_ZH, finding.status.value))}</span>
              <span class="severity">{html.escape(zh(SEVERITY_ZH, finding.severity))}</span>
              <span>{html.escape('、'.join(zh(ISSUE_ZH, item.value) for item in finding.issue_codes))}</span>
              <span class="resolution {'done' if resolved else 'open'}">{'已解决' if resolved else '待处理'}</span></header>
              <div class="columns"><section><h3>报告结论</h3><p class="claim">{html.escape(claim.text)}</p>
              <small>{html.escape(locator_zh(claim.locator))}</small>
              <p><a class="jump" target="_blank" href="/review/{job.id}/report/{finding.id}">打开报告原文定位 ↗</a></p>
              <p>{html.escape(message_zh(finding.explanation))}</p>
              <h4>逐项核对</h4><table><thead><tr><th>规则</th><th>结果</th><th>报告值</th><th>证据值</th><th>说明</th></tr></thead><tbody>{checks}</tbody></table>
              </section><section><h3>前 5 条候选证据</h3>{''.join(evidence_cards) or '<em>未找到候选证据</em>'}</section></div>
              {memory_section}
              <details><summary>人工决策历史（{len(histories.get(finding.id, []))}）</summary><ul>{history}</ul></details>
              <div class="actions">
                <button type="button" data-decision-action="accept" data-finding-id="{finding.id}">确认发现（不解除风险）</button>
                <button type="button" data-decision-action="reject" data-finding-id="{finding.id}">驳回误报</button>
                <button type="button" data-decision-action="replace" data-finding-id="{finding.id}">修正报告并复验</button>
                <button type="button" data-decision-action="waive" data-finding-id="{finding.id}">正式豁免</button>
              </div>
            </article>
            """
        )
    summary = job.summary()
    summary_text = (
        f"结论 {summary['total']} 条 · 证据充分 {summary['supported']} 条 · "
        f"证据冲突 {summary['conflict']} 条 · 部分支持 {summary['partial']} 条 · "
        f"缺少证据 {summary['unsupported']} 条 · 过期 {summary['stale']} 条 · "
        f"待复核 {summary['needs_review']} 条 · 未处理高风险 {summary['unresolved_high']} 条"
    )
    warnings = "".join(f"<li>{html.escape(item)}</li>" for item in job.warnings)
    policies = (
        store.list_policies(job.policy_snapshot.profile_id)
        if store and job.policy_snapshot.profile_id
        else []
    )
    policy_rows = []
    for policy in policies:
        transition = ""
        if policy.state == "draft":
            transition = f'<button type="button" data-policy-activate="{html.escape(policy.id)}">批准并激活</button>'
        elif policy.state == "active":
            transition = f'<button type="button" data-policy-retire="{html.escape(policy.id)}">停用策略</button>'
        state_label = {"draft": "草稿", "active": "已激活", "retired": "已停用"}.get(
            policy.state, policy.state
        )
        raw_kind = policy.rule.get("kind", "attention")
        kind_label = {"attention": "提高关注度", "severity": "调整风险级别"}.get(
            raw_kind, raw_kind
        )
        policy_rows.append(
            f"<li><code>{html.escape(policy.id)}</code> · v{policy.version} · "
            f"<strong>{html.escape(state_label)}</strong> · {html.escape(kind_label)}"
            f" · {html.escape(policy.hash[:12])}{transition}</li>"
        )
    role_label = {
        "generic": "通用",
        "audit": "审计",
        "operations": "运营",
        "consulting": "咨询",
    }.get(job.policy_snapshot.role_template, job.policy_snapshot.role_template)
    profile_panel = (
        f'<section class="profile"><strong>审阅配置：</strong>{html.escape(job.policy_snapshot.profile_name or "")}'
        f' · 岗位模板 {html.escape(role_label)} v{html.escape(job.policy_snapshot.role_version)}'
        f' · 策略快照 {html.escape((job.policy_snapshot.policy_hash or "")[:12])}'
        f'<details><summary>策略版本（{len(policy_rows)}）</summary><ul>{"".join(policy_rows) or "<li>尚无策略</li>"}</ul></details>'
        f'<p>策略只影响新任务；历史先例和策略均不能自动关闭风险。</p></section>'
        if job.policy_snapshot.profile_id
        else '<section class="profile"><strong>审阅配置：</strong>未启用个人记忆；本任务仅使用通用规则。</section>'
    )
    artifact_panel = _artifact_download_panel(job)
    audit_complete = job.status == "completed"
    review_complete = summary["unresolved_high"] == 0
    final_ready = bool(summary["delivery_ready"])
    export_disabled = "" if final_ready else " disabled aria-disabled=\"true\""
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>ClaimLedger · {html.escape(job.request.case_name)}</title>
<style>
:root{{--ink:#142638;--muted:#607184;--paper:#f3f7fa;--panel:#fff;--line:#d9e3eb;--accent:#0e9488;--navy:#0b1f33;--danger:#c43d4f;--ok:#177d70;--amber:#d99213;--shadow:0 12px 32px #0b1f3312;color-scheme:light dark}}
*{{box-sizing:border-box}} body{{margin:0;background:var(--paper);color:var(--ink);font:14px/1.58 -apple-system,BlinkMacSystemFont,"PingFang SC","Microsoft YaHei",sans-serif}}
main{{max-width:1380px;margin:auto;padding:96px 32px 32px}} h1{{margin:0;font-size:24px}} .summary{{color:var(--muted);margin:3px 0 14px}}
.brandbar{{position:fixed;z-index:20;left:0;right:0;top:0;display:flex;justify-content:space-between;align-items:center;padding:13px max(24px,calc((100vw - 1380px)/2 + 32px));background:linear-gradient(120deg,#0b1f33,#123953);color:#fff;box-shadow:0 7px 24px #0b1f3330}}
.brandbar strong{{font-size:17px}} .brandbar small{{opacity:.72}} .offline{{display:inline-flex;align-items:center;gap:7px;padding:6px 10px;border-radius:999px;background:#ffffff14}}
.offline::before{{content:"";width:7px;height:7px;border-radius:50%;background:#43ddc6;box-shadow:0 0 0 4px #43ddc625}}
.metrics{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:14px 0}} .metric{{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:13px;box-shadow:var(--shadow)}}
.metric strong{{display:block;color:var(--navy);font-size:24px}} .metric span{{color:var(--muted);font-size:12px}}
.distribution{{display:flex;flex-wrap:wrap;gap:7px;margin:-2px 0 14px}} .distribution span{{padding:5px 9px;border:1px solid var(--line);border-radius:999px;background:var(--panel);color:var(--muted);font-size:12px}}
.gate{{display:flex;justify-content:space-between;gap:16px;align-items:center;background:var(--panel);border:1px solid var(--line);border-left:5px solid {'var(--ok)' if summary['delivery_ready'] else 'var(--danger)'};border-radius:12px;padding:14px 18px;box-shadow:var(--shadow)}}
.journey{{display:grid;grid-template-columns:repeat(3,1fr);gap:10px;margin:14px 0}} .step{{background:var(--panel);border:1px solid var(--line);border-radius:10px;padding:12px}}
.step.done{{border-color:#79b995;background:#edf9f5}} .step.current{{border-color:var(--amber);background:#fff8e8}} .step strong{{display:block}}
.finding{{scroll-margin-top:90px;background:var(--panel);border:1px solid var(--line);border-radius:14px;padding:18px;margin:14px 0;box-shadow:var(--shadow)}}
.columns{{display:grid;grid-template-columns:1fr 1.08fr;gap:22px}} blockquote{{margin:8px 0;padding:10px;border-left:3px solid var(--accent);background:#f7f9ff;max-height:150px;overflow:auto}}
.badge{{padding:3px 8px;border-radius:999px;font-weight:700}} .supported{{background:#d9ead3}} .partial{{background:#fff2cc}}
.unsupported,.conflict{{background:#f4cccc}} .stale{{background:#fce5cd}} .needs_review{{background:#d9d2e9}}
.severity{{margin:0 8px;color:var(--muted)}} .resolution{{float:right;border-radius:999px;padding:3px 8px}} .resolution.done{{background:#d9ead3}} .resolution.open{{background:#f4cccc}}
button{{margin:8px 8px 0 0;padding:9px 13px;border:1px solid var(--line);border-radius:8px;background:var(--panel);color:var(--ink);cursor:pointer}} button.primary{{background:var(--accent);border-color:var(--accent);color:#fff}} button.danger{{color:var(--danger)}} button:disabled{{cursor:not-allowed;opacity:.45}}
.jump{{display:inline-block;color:var(--accent);font-weight:650;text-decoration:none}} .filters{{display:flex;flex-wrap:wrap;gap:8px;margin:14px 0 22px}}
.filters button.active{{background:var(--accent);color:white;border-color:var(--accent)}} .evidence{{border:1px solid var(--line);border-radius:10px;padding:10px;margin:8px 0}}
table{{border-collapse:collapse;width:100%;font-size:13px}} th,td{{border:1px solid var(--line);padding:6px;text-align:left;vertical-align:top}} th{{background:#eef2f8}}
.claim{{font-size:17px;font-weight:650}} .warnings{{color:#8a4b08}}
.profile,.memory{{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:12px;margin:12px 0}}
.memory{{background:#f8f6ff;border-left:4px solid #6b55c5}} .memory p{{margin:5px 0}} .memory-warning{{color:#6a3f00}}
.artifacts{{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:14px;margin:14px 0}} .artifact-links{{display:flex;flex-wrap:wrap;gap:9px;margin-top:8px}}
.artifact-link{{display:inline-block;background:#eef3ff;color:var(--accent);border:1px solid #c9d5f2;border-radius:8px;padding:9px 12px;text-decoration:none;font-weight:650}}
.artifact-empty{{color:var(--muted)}} .modal-backdrop{{position:fixed;z-index:40;inset:0;display:grid;place-items:center;padding:18px;background:#071521a8;backdrop-filter:blur(4px)}} .modal-backdrop[hidden]{{display:none}}
.modal{{width:min(640px,100%);max-height:90vh;overflow:auto;background:var(--panel);border:1px solid var(--line);border-radius:16px;padding:20px;box-shadow:0 24px 80px #0006}} .modal h2{{margin:0 0 4px}} .modal-note{{color:var(--muted);margin:0 0 14px}}
.field{{display:block;margin:11px 0}} .field span{{display:block;margin-bottom:4px;font-weight:650}} .field textarea,.field input{{width:100%;padding:10px;border:1px solid var(--line);border-radius:8px;background:var(--panel);color:var(--ink);font:inherit}} .field textarea{{min-height:92px;resize:vertical}}
.diff{{display:grid;grid-template-columns:1fr 1fr;gap:9px}} .diff pre{{white-space:pre-wrap;margin:0;padding:10px;border-radius:8px;background:var(--paper);font:inherit}} .modal-actions{{display:flex;justify-content:flex-end;gap:8px}}
.toast{{position:fixed;z-index:50;right:20px;bottom:20px;max-width:420px;padding:12px 15px;border-radius:10px;background:var(--navy);color:#fff;box-shadow:var(--shadow)}} .toast[hidden]{{display:none}}
@media(prefers-color-scheme:dark){{:root{{--ink:#edf4fa;--muted:#a9b7c3;--paper:#0f171f;--panel:#16212b;--line:#2c3c4b;--accent:#35c4b5;--navy:#d9e8f5;--danger:#ff7888;--ok:#35c4b5;--amber:#f4b64b;--shadow:0 12px 32px #0004}} .step.done{{background:#123d3b}} .step.current{{background:#3e3219}} blockquote,th,.artifact-link,.memory{{background:#1d2b37}}}}
@media(max-width:850px){{.columns,.journey{{grid-template-columns:1fr}} .metrics{{grid-template-columns:repeat(2,1fr)}} main{{padding:88px 18px 18px}} .gate{{display:block}} .diff{{grid-template-columns:1fr}}}}
@media(max-width:430px){{.metrics{{grid-template-columns:1fr 1fr}} .brandbar small{{display:none}} .finding{{padding:13px}}}}
</style></head><body>
<header class="brandbar"><div><strong>ClaimLedger 可信交付台</strong><br><small>报告证据审计与人工复核</small></div><span class="offline">本地离线</span></header>
<main id="claimledger-app" data-job-id="{html.escape(job.id, quote=True)}"><h1>{html.escape(job.request.case_name)}</h1><div class="summary">任务 {html.escape(job.id)} · {html.escape(summary_text)}</div>
<section class="metrics" aria-label="审计指标">
<div class="metric"><strong>{summary['total']}</strong><span>候选发现</span></div>
<div class="metric"><strong>{int(summary['total']) - int(summary['supported'])}</strong><span>风险发现</span></div>
<div class="metric"><strong>{summary['unresolved_high']}</strong><span>未处理高风险</span></div>
<div class="metric"><strong>{summary['decision_events']}</strong><span>人工决定</span></div>
</section>
<section class="distribution" aria-label="风险分布">
<span>证据充分 {summary['supported']}</span><span>证据冲突 {summary['conflict']}</span>
<span>缺少证据 {summary['unsupported']}</span><span>部分支持 {summary['partial']}</span>
<span>证据过期 {summary['stale']}</span><span>需要复核 {summary['needs_review']}</span>
</section>
<section class="journey" aria-label="交付进度">
<div class="step {'done' if audit_complete else 'current'}"><strong>1. 自动审计</strong>{'已完成并生成候选发现' if audit_complete else '正在解析和核证'}</div>
<div class="step {'done' if review_complete else 'current'}"><strong>2. 人工复核</strong>{'高风险已全部处理' if review_complete else f"仍有 {summary['unresolved_high']} 项高风险待处理"}</div>
<div class="step {'done' if final_ready else ''}"><strong>3. 可信交付</strong>{'可以生成最终稿' if final_ready else '等待人工复核完成'}</div>
</section>
<section class="gate"><div><strong>{'可以生成可信交付版本' if summary['delivery_ready'] else '交付闸门已锁定'}</strong><br>
证据覆盖：{html.escape(zh(COVERAGE_ZH, job.coverage_status))} · 未解决高风险：{summary['unresolved_high']}</div>
<button type="button" class="primary" data-export-final{export_disabled}>生成最终交付包</button></section>
<section class="artifacts"><strong>当前有效交付物</strong>{artifact_panel}</section>
{profile_panel}
{f'<details class="warnings"><summary>解析警告（{len(job.warnings)}）</summary><ul>{warnings}</ul></details>' if job.warnings else ''}
<div class="filters"><button type="button" class="active" data-filter="all">全部</button>
<button type="button" data-filter="unresolved">仅未解决</button><button type="button" data-filter="high">高风险</button>
<button type="button" data-filter="conflict">冲突</button><button type="button" data-filter="unsupported">无来源</button>
<button type="button" data-filter="partial">弱支持</button><button type="button" data-filter="stale">过期</button>
<button type="button" data-filter="needs_review">待复核</button><button type="button" data-filter="supported">已支持</button></div>
{''.join(rows) or '<p>任务仍在运行或没有发现。</p>'}
<div class="modal-backdrop" id="review-modal" hidden>
  <section class="modal" role="dialog" aria-modal="true" aria-labelledby="modal-title">
    <h2 id="modal-title">记录人工决定</h2>
    <p class="modal-note" id="modal-note">此操作将写入不可变审计轨迹。</p>
    <form id="review-form">
      <input type="hidden" id="modal-mode">
      <input type="hidden" id="modal-finding-id">
      <input type="hidden" id="modal-policy-id">
      <div class="diff" id="replacement-diff" hidden>
        <label class="field"><span>报告原文</span><pre id="modal-original"></pre></label>
        <label class="field"><span>建议修正</span><pre id="modal-suggested"></pre></label>
      </div>
      <label class="field" id="replacement-field" hidden><span>修正后的结论（提交后立即重新核证）</span><textarea id="modal-replacement"></textarea></label>
      <label class="field"><span id="reason-label">理由（必填）</span><textarea id="modal-reason" required></textarea></label>
      <label class="field" id="owner-field" hidden><span>风险责任人（必填）</span><input id="modal-owner" value="当前用户"></label>
      <label class="field" id="expiry-field" hidden><span>豁免有效期（留空表示永久，格式 YYYY-MM-DD）</span><input id="modal-expiry" type="date"></label>
      <div class="modal-actions"><button type="button" data-modal-cancel>取消</button><button type="submit" class="primary">确认写入</button></div>
    </form>
  </section>
</div>
<div class="toast" id="toast" hidden role="status"></div>
<script{f' nonce="{html.escape(script_nonce, quote=True)}"' if script_nonce else ''}>
const appRoot=document.getElementById('claimledger-app');
const job=appRoot.dataset.jobId;
const modal=document.getElementById('review-modal');
const form=document.getElementById('review-form');
const modeField=document.getElementById('modal-mode');
const findingField=document.getElementById('modal-finding-id');
const policyField=document.getElementById('modal-policy-id');
const reasonField=document.getElementById('modal-reason');
const replacementField=document.getElementById('modal-replacement');
const ownerField=document.getElementById('modal-owner');
const expiryField=document.getElementById('modal-expiry');
const profile={json.dumps(job.policy_snapshot.profile_id)};
function selectedEvidence(id){{return [...document.querySelectorAll(`input[name="evidence-${{id}}"]:checked`)].map(item=>item.value)}}
function showToast(message){{
 const toast=document.getElementById('toast');toast.textContent=message;toast.hidden=false;
 window.setTimeout(()=>{{toast.hidden=true}},4500);
}}
async function send(finding_id,action,replacement_text=null,reason=null,extra={{}}){{
 const payload={{finding_id,action,replacement_text,reason,selected_evidence_ids:selectedEvidence(finding_id),...extra}};
 const response=await fetch(`/api/v1/jobs/${{job}}/decisions`,{{method:'POST',headers:{{'Content-Type':'application/json'}},body:JSON.stringify(payload)}});
 if(!response.ok){{showToast(await response.text());return false}} location.reload();return true;
}}
function acceptFinding(id){{send(id,'accept',null,'确认系统发现成立；尚未解除风险')}}
function openModal(mode,id){{
 modeField.value=mode;findingField.value=mode.startsWith('policy-')?'':id;policyField.value=mode.startsWith('policy-')?id:'';
 const article=id?document.getElementById(id):null;
 const original=article?.querySelector('.claim')?.textContent?.trim()||'';
 const suggested=article?.dataset.suggestedReplacement||'';
 reasonField.value=mode==='replace'?'根据证据修正':'';
 replacementField.value=suggested||original;
 ownerField.value='当前用户';expiryField.value='';
 const replace=mode==='replace';const waive=mode==='waive';
 document.getElementById('replacement-diff').hidden=!replace;
 document.getElementById('replacement-field').hidden=!replace;
 document.getElementById('owner-field').hidden=!waive;
 document.getElementById('expiry-field').hidden=!waive;
 replacementField.required=replace;ownerField.required=waive;
 document.getElementById('modal-original').textContent=original;
 document.getElementById('modal-suggested').textContent=suggested||'尚无系统建议，请根据证据填写。';
 const titles={{reject:'驳回误报',replace:'修正并重新核证',waive:'正式豁免','policy-draft':'生成策略草稿','policy-activate':'批准并激活策略','policy-retire':'停用策略'}};
 document.getElementById('modal-title').textContent=titles[mode]||'记录人工决定';
 document.getElementById('modal-note').textContent=replace?'修正内容只有复验通过后才会解除交付阻断。':waive?'豁免不会成为事实证据，请明确责任人与有效期。':'理由将写入不可变审计轨迹。';
 modal.hidden=false;reasonField.focus();
}}
function closeModal(){{modal.hidden=true;form.reset()}}
function filterItems(mode,button){{document.querySelectorAll('.filters button').forEach(item=>item.classList.remove('active'));button.classList.add('active');
 document.querySelectorAll('.finding').forEach(item=>{{item.hidden=!(mode==='all'||item.dataset.status===mode||item.dataset.severity===mode||(mode==='unresolved'&&item.dataset.resolved==='no'))}})}}
async function exportFinal(){{const response=await fetch(`/api/v1/jobs/${{job}}/exports?final=true`,{{method:'POST'}}); const body=await response.json(); if(!response.ok)showToast(body.detail||JSON.stringify(body));else location.reload()}}
async function draftPolicy(memoryId){{
 openModal('policy-draft',memoryId);
}}
async function transitionPolicy(id,action){{
 openModal(`policy-${{action}}`,id);
}}
form.addEventListener('submit',async event=>{{
 event.preventDefault();
 const mode=modeField.value;const reason=reasonField.value.trim();
 if(!reason){{showToast('请填写理由。');reasonField.focus();return}}
 if(mode==='replace'){{
   const replacement=replacementField.value.trim();
   if(!replacement){{showToast('请填写修正后的结论。');replacementField.focus();return}}
   await send(findingField.value,'replace',replacement,reason);return;
 }}
 if(mode==='reject'){{await send(findingField.value,'reject',null,reason);return}}
 if(mode==='waive'){{
   const owner=ownerField.value.trim();
   if(!owner){{showToast('请填写风险责任人。');ownerField.focus();return}}
   await send(findingField.value,'waive',null,reason,{{risk_owner:owner,waiver_expires_at:expiryField.value||null}});return;
 }}
 let response;
 if(mode==='policy-draft'){{
   response=await fetch(`/api/v1/profiles/${{profile}}/policies`,{{method:'POST',headers:{{'Content-Type':'application/json'}},body:JSON.stringify({{source_memory_id:policyField.value,kind:'attention',actor:'local-user',reason}})}});
 }}else{{
   const action=mode.replace('policy-','');
   response=await fetch(`/api/v1/policies/${{policyField.value}}/${{action}}`,{{method:'POST',headers:{{'Content-Type':'application/json'}},body:JSON.stringify({{actor:'local-user',reason}})}});
 }}
 const body=await response.json();if(!response.ok)showToast(body.detail||JSON.stringify(body));else location.reload();
}});
modal.addEventListener('click',event=>{{if(event.target===modal)closeModal()}});
document.querySelector('[data-modal-cancel]').addEventListener('click',closeModal);
document.addEventListener('keydown',event=>{{if(event.key==='Escape'&&!modal.hidden)closeModal()}});
appRoot.addEventListener('click',event=>{{
 const button=event.target.closest('button');
 if(!button||!appRoot.contains(button))return;
 if(button.dataset.filter){{filterItems(button.dataset.filter,button);return}}
 if(button.hasAttribute('data-export-final')){{exportFinal();return}}
 if(button.dataset.policyDraft){{draftPolicy(button.dataset.policyDraft);return}}
 if(button.dataset.policyActivate){{transitionPolicy(button.dataset.policyActivate,'activate');return}}
 if(button.dataset.policyRetire){{transitionPolicy(button.dataset.policyRetire,'retire');return}}
 const action=button.dataset.decisionAction; const id=button.dataset.findingId;
 if(!action||!id)return;
 if(action==='accept'){{acceptFinding(id);return}}
 if(action==='reject'||action==='waive'||action==='replace'){{openModal(action,id)}}
}});
</script></main></body></html>"""


app = create_app()
