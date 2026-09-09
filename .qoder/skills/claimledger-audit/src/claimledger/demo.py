from __future__ import annotations

import json
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageEnhance, ImageFilter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DISCLOSURE = "公开问题原型驱动｜脱敏合成测试文件｜不含客户数据"
WATERMARK = DISCLOSURE


def _set_style_font(style, name: str, size: float, color: str, *, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        style.element.rPr.rFonts.set(qn(key), name)


def _add_page_field(paragraph) -> None:
    paragraph.add_run("  ·  第 ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    paragraph.add_run(" 页")


def _add_bottom_rule(paragraph, color: str = "2E74B5") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _style_document(
    document: Document,
    title: str,
    subtitle: str | None = None,
    *,
    metadata: list[tuple[str, str]] | None = None,
) -> None:
    """Apply the decision-memo preset and memo-masthead pattern."""

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    styles = document.styles
    body_font = "Arial Unicode MS"
    _set_style_font(styles["Normal"], body_font, 10.5, "20252B")
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.1
    _set_style_font(styles["Title"], body_font, 23, "111827", bold=True)
    styles["Title"].paragraph_format.space_after = Pt(4)
    title_properties = styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 12, 6),
        ("Heading 2", 13, "2E74B5", 10, 5),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        _set_style_font(styles[name], body_font, size, color, bold=True)
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True
    for name in ("List Number", "List Bullet"):
        _set_style_font(styles[name], body_font, 10.5, "20252B")
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.text = "CLAIMLEDGER  ·  EVIDENCE ASSURANCE DEMO"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.name = body_font
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.color.rgb = RGBColor.from_string("667085")

    paragraph = document.add_paragraph()
    paragraph.style = styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(title)
    if subtitle:
        note = document.add_paragraph(subtitle)
        note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        note.paragraph_format.space_after = Pt(10)
        note.runs[0].font.name = body_font
        note.runs[0].font.color.rgb = RGBColor.from_string("475467")
        note.runs[0].font.size = Pt(12)
    for label, value in metadata or []:
        row = document.add_paragraph()
        row.style = styles["Caption"]
        row.paragraph_format.space_before = Pt(0)
        row.paragraph_format.space_after = Pt(2)
        row.paragraph_format.line_spacing = 1.0
        label_run = row.add_run(f"{label}：")
        label_run.bold = True
        label_run.font.name = body_font
        label_run.font.size = Pt(9.5)
        value_run = row.add_run(value)
        value_run.font.name = body_font
        value_run.font.size = Pt(9.5)
        value_run.font.color.rgb = RGBColor.from_string("344054")
    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(10)
    _add_bottom_rule(rule)
    banner = document.add_paragraph(WATERMARK)
    banner.alignment = WD_ALIGN_PARAGRAPH.LEFT
    banner.paragraph_format.space_after = Pt(10)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "FFF3E8")
    banner._p.get_or_add_pPr().append(shading)
    banner.runs[0].font.bold = True
    banner.runs[0].font.name = body_font
    banner.runs[0].font.size = Pt(9)
    banner.runs[0].font.color.rgb = RGBColor.from_string("9A3412")
    footer = section.footer.paragraphs[0]
    footer.text = WATERMARK
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_field(footer)
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(128, 128, 128)


def _save_docx(path: Path, title: str, paragraphs: list[str], *, subtitle: str | None = None) -> Path:
    document = Document()
    _style_document(document, title, subtitle)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


def _save_pdf(path: Path, title: str, lines: list[str]) -> Path:
    page_rect = fitz.paper_rect("a4")
    writer = fitz.open()
    page = writer.new_page(width=page_rect.width, height=page_rect.height)
    y = 60

    def add_header(current_page, continuation: bool = False) -> float:
        current_title = f"{title}（续）" if continuation else title
        current_page.insert_text((54, 60), current_title, fontname="china-s", fontsize=17, color=(0.07, 0.11, 0.18))
        current_page.insert_text((54, 88), WATERMARK, fontname="china-s", fontsize=9, color=(0.60, 0.20, 0.05))
        current_page.draw_line((54, 100), (page_rect.width - 54, 100), color=(0.18, 0.45, 0.71), width=1)
        return 128

    y = add_header(page)
    for line in lines:
        if y > page_rect.height - 72:
            page = writer.new_page(width=page_rect.width, height=page_rect.height)
            y = add_header(page, continuation=True)
        for start in range(0, len(line), 44):
            page.insert_text(
                (54, y),
                line[start:start + 44],
                fontname="china-s",
                fontsize=11,
                color=(0.10, 0.12, 0.18),
            )
            y += 19
        y += 4
    for index, current_page in enumerate(writer, start=1):
        footer = f"{WATERMARK}  ·  第 {index}/{len(writer)} 页"
        width = fitz.get_text_length(footer, fontname="china-s", fontsize=8)
        current_page.insert_text(
            ((page_rect.width - width) / 2, page_rect.height - 28),
            footer,
            fontname="china-s",
            fontsize=8,
            color=(0.45, 0.45, 0.45),
        )
    writer.set_metadata({"title": title, "subject": WATERMARK, "author": "ClaimLedger"})
    writer.save(path)
    writer.close()
    return path


def _scan_pdf(text_pdf: Path, scan_pdf: Path, *, low_quality: bool = False) -> Path:
    source = fitz.open(text_pdf)
    output = fitz.open()
    for page in source:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.65, 1.65), alpha=False)
        image = Image.open(__import__("io").BytesIO(pixmap.tobytes("png"))).convert("L")
        if low_quality:
            image = image.resize((max(500, image.width // 2), max(700, image.height // 2)))
            image = ImageEnhance.Contrast(image).enhance(0.72).filter(ImageFilter.GaussianBlur(0.8))
        buffer = __import__("io").BytesIO()
        image.save(buffer, format="JPEG", quality=55 if low_quality else 82)
        page_out = output.new_page(width=image.width, height=image.height)
        page_out.insert_image(page_out.rect, stream=buffer.getvalue())
    output.save(scan_pdf)
    output.close()
    source.close()
    text_pdf.unlink()
    return scan_pdf


def _pdf_to_image(text_pdf: Path, image_path: Path, *, low_quality: bool = False) -> Path:
    source = fitz.open(text_pdf)
    pixmap = source[0].get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
    image = Image.open(__import__("io").BytesIO(pixmap.tobytes("png"))).convert("RGB")
    if low_quality:
        image = image.resize((max(520, image.width // 3), max(720, image.height // 3)))
        image = ImageEnhance.Contrast(image).enhance(0.62).filter(ImageFilter.GaussianBlur(1.1))
    image.save(image_path)
    source.close()
    text_pdf.unlink()
    return image_path


def _style_workbook(workbook: Workbook) -> None:
    workbook.properties.title = "ClaimLedger evidence-assurance demonstration"
    workbook.properties.subject = DISCLOSURE
    workbook.properties.creator = "ClaimLedger"
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.oddHeader.right.text = "CLAIMLEDGER · EVIDENCE ASSURANCE DEMO"
        sheet.oddFooter.center.text = DISCLOSURE
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(32, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def _procurement_demo(output: Path) -> dict[str, str]:
    evidence = output / "evidence"
    evidence.mkdir(parents=True, exist_ok=True)
    report_claims = [
        "评估期间为2025年7月1日至2026年6月30日。",
        "年度需求为1200万个750mL双淋膜纸餐盒。",
        "年度预算为1080万元。",
        "采购委员会已批准供应商B占70%、供应商A占30%的正式份额。",
        "切换供应商B每年可节省168万元。",
        "供应商B的综合到岸成本最低。",
        "三份报价均有效至2026年9月30日。",
        "供应商A报价0.82元/个，含税含运。",
        "供应商B报价0.68元/个，含税。",
        "供应商C报价0.77元/个，含税含运。",
        "供应商B报价包含运费。",
        "三家比价对象均为750mL双淋膜产品。",
        "供应商B的750mL产品最小起订量为10万件。",
        "供应商B付款条件为30%预付、70%验收后30日支付。",
        "总拥有成本测算已计入4万元制版费。",
        "供应商A准时交付率为93.2%。",
        "供应商B准时交付率为96.8%。",
        "供应商C准时交付率为90.1%。",
        "供应商B平均交付周期为10个自然日。",
        "供应商B在2026年4月发生过一次重大延期。",
        "供应商B质量缺陷率为0.6%。",
        "供应商B质量表现为三家最佳。",
        "供应商B的5万件试单已验收。",
        "供应商B投诉平均关闭时间为2.1天。",
        "评估期内没有食品安全事件。",
        "当前检测报告是最新有效报告。",
        "检测报告样品与750mL双淋膜采购品一致。",
        "检测机构CMA范围覆盖该检测方法。",
        "当前FSC证书有效至2027年。",
        "FSC证书持有人就是供应商B合同主体博远包装有限公司。",
        "供应商B的ISO 9001证书有效至2027年。",
        "供应商B注册资本为3000万元。",
        "供应商B月产能为800万件并已独立核验。",
        "经审计月产能为600万件。",
        "供应商B当前产能利用率为62%。",
        "供应商B现场有两条备用产线。",
        "供应商B对原材料供应商每年复审。",
        "历史三批采购累计付款58.4万元。",
        "三笔款项均支付给合同主体博远包装有限公司。",
        "所有评审人员均提交利益冲突声明。",
        "70%与30%的份额分配足以控制断供风险。",
        "供应商B可以无条件续签。",
        "本次证据截止日为2026年6月30日。",
        "报告只引用了截止日前收到的17组材料。",
    ]
    report = Document()
    _style_document(
        report,
        "食品包装供应商续签决策报告",
        "采购委员会审议稿｜AI 初稿待证据审计",
        metadata=[
            ("送审对象", "采购委员会"),
            ("案例编号", "CL-PROC-001"),
            ("证据截止日", "2026-06-30"),
            ("交付状态", "存在高风险结论，未经人工闭环不得交付"),
        ],
    )
    report.add_heading("一、决策结论", level=1)
    for claim in report_claims[:15]:
        report.add_paragraph(claim, style="List Number")
    report.add_heading("二、履约、质量与资质", level=1)
    for claim in report_claims[15:37]:
        report.add_paragraph(claim, style="List Number")
    report.add_heading("三、付款、治理与交付建议", level=1)
    for claim in report_claims[37:]:
        report.add_paragraph(claim, style="List Number")
    report_path = output / "report.docx"
    report.save(report_path)

    _save_docx(
        evidence / "01_审计口径与决策门槛.docx",
        "审计口径与决策门槛",
        [
            "评估期间为2025年7月1日至2026年6月30日，证据截止日为2026年6月30日。",
            "年度需求为1200万个750mL双淋膜纸餐盒，年度预算为1080万元。",
            "报价必须统一到同规格、含税、含运、包含一次性费用的年度总拥有成本口径。",
            "价格、证书、合同主体和检测范围风险关闭前，不得无条件续签供应商B。",
            "本轮共登记17组证据材料，均在截止日前收到。",
        ],
    )

    tco = Workbook()
    sheet = tco.active
    sheet.title = "TCO复核"
    sheet.append([
        "供应商", "产品规格", "淋膜", "报价(元/个)", "税口径", "运费(元/个)", "运费口径",
        "制版费(元)", "同口径成本(元/个)", "年度成本(元)", "排名",
    ])
    sheet.append(["供应商A", "750mL", "双淋膜", 0.82, "含税含运", 0, "已含", 0, 0.82, 9840000, 2])
    sheet.append(["供应商B", "750mL", "双淋膜", 0.76, "未税", 0.018, "另计", 40000, 0.8801, 10561600, 3])
    sheet.append(["供应商C", "750mL", "双淋膜", 0.77, "含税含运", 0, "已含", 0, 0.77, 9240000, 1])
    demand = tco.create_sheet("年度需求")
    demand.append(["项目", "数量(个)", "单位"])
    demand.append(["750mL双淋膜纸餐盒", 12000000, "个"])
    formula = tco.create_sheet("公式说明")
    formula.append(["项目", "公式", "结果"])
    formula.append(["供应商B同口径成本", "0.76×1.13+0.018+40000÷12000000", 0.8801])
    formula.append([
        "B相对A年度成本增加额",
        "(0.76×1.13+0.018+40000÷12000000-0.82)×12000000",
        721600,
    ])
    formula.append(["AI初稿错误节省额", "(0.82-0.68)×12000000", 1680000])
    _style_workbook(tco)
    tco_path = evidence / "02_三家供应商报价与TCO测算.xlsx"
    tco.save(tco_path)

    kpi = Workbook()
    kpi_sheet = kpi.active
    kpi_sheet.title = "履约KPI"
    kpi_sheet.append(["供应商", "准时交付率", "平均交付周期", "质量缺陷率", "投诉关闭天数", "统计口径"])
    kpi_sheet.append(["供应商A", 0.932, "13个自然日", 0.012, 3.4, "全量订单"])
    kpi_sheet.append(["供应商B", 0.884, "15个工作日", 0.018, 2.1, "全量订单，含延期订单"])
    kpi_sheet.append(["供应商C", 0.901, "14个自然日", 0.005, 2.8, "全量订单"])
    filtered = kpi.create_sheet("旧版过滤汇总")
    filtered.append(["供应商", "准时交付率", "质量缺陷率", "问题"])
    filtered.append(["供应商B", 0.968, 0.006, "错误过滤延期与退货批次"])
    _style_workbook(kpi)
    kpi.save(evidence / "03_历史采购订单与履约KPI.xlsx")

    _save_pdf(
        evidence / "04_供应商A报价单.pdf",
        "安禾包装有限公司报价单",
        ["产品：750mL双淋膜纸餐盒", "报价：0.82元/个，含税含运", "报价有效期至2026年9月30日"],
    )
    temp = evidence / "_05_text.pdf"
    _save_pdf(
        temp,
        "博远包装有限公司报价单",
        [
            "650mL单淋膜纸餐盒报价0.68元/个，未税，运费另计。",
            "750mL双淋膜纸餐盒报价0.76元/个，未税，运费另计0.018元/个。",
            "750mL产品最小起订量为10万件，制版费4万元。",
            "付款条件为30%预付、70%验收后30日支付。",
            "报价有效期至2026年8月31日。",
        ],
    )
    _scan_pdf(temp, evidence / "05_供应商B报价单_扫描.pdf")
    _save_pdf(
        evidence / "06_供应商C报价单.pdf",
        "诚新包装有限公司报价单",
        ["产品：750mL双淋膜纸餐盒", "报价：0.77元/个，含税含运", "报价有效期至2026年9月30日"],
    )
    temp = evidence / "_07_text.pdf"
    _save_pdf(
        temp,
        "产品检测报告",
        ["签发日期：2024年5月20日", "送检单位：博远包装有限公司", "样品：650mL单淋膜纸餐盒", "结论：所检项目符合送检样品对应方法要求"],
    )
    _scan_pdf(temp, evidence / "07_供应商B产品检测报告_扫描.pdf")
    temp = evidence / "_08_text.pdf"
    _save_pdf(temp, "CMA能力附表", ["机构具备食品接触用纸制品部分项目能力。", "关键方法编号：GB 4806.8-2022（扫描模糊，需人工核对范围）"])
    _pdf_to_image(temp, evidence / "08_检测机构CMA能力附表_低清.tiff", low_quality=True)
    temp = evidence / "_09_text.pdf"
    _save_pdf(temp, "FSC证书", ["证书持有人：博远供应链有限公司（关联公司）", "有效期至2025年12月31日", "状态：已失效"])
    _pdf_to_image(temp, evidence / "09_供应商B提交的FSC证书.jpg")
    _save_pdf(evidence / "10_供应商B_ISO9001证书.pdf", "ISO 9001证书", ["证书持有人：博远包装有限公司", "有效期至2027年10月31日", "证书状态：有效"])

    _save_docx(
        evidence / "11_供应商主体与关联关系说明.docx",
        "供应商主体与关联关系说明",
        [
            "合同主体为博远包装有限公司，注册资本为3000万元。",
            "博远供应链有限公司是关联公司，不是本次合同主体。",
            "证书、合同、付款收款人必须分别核对，不得因为关联关系视为同一主体。",
        ],
    )
    _save_docx(
        evidence / "12_现场验厂纪要.docx",
        "现场验厂纪要",
        [
            "现场看到两条备用产线。",
            "供应商自述月产能为800万件，本次未获得独立审计或设备利用记录。",
            "原材料供应商执行年度复审，抽查到2025年度复审记录。",
        ],
    )
    trial = Workbook()
    trial_sheet = trial.active
    trial_sheet.title = "试单验收"
    trial_sheet.append(["供应商", "试单数量(件)", "验收状态", "验收日期"])
    trial_sheet.append(["供应商B", 50000, "已验收", "2026-03-18"])
    _style_workbook(trial)
    trial.save(evidence / "13_试单验收记录.xlsx")
    _save_docx(
        evidence / "14_客户投诉与8D整改记录.docx",
        "投诉与整改记录",
        [
            "供应商B在2026年4月发生一次重大延期，原因是主线设备故障。",
            "供应商B投诉平均关闭时间为2.1天。",
            "评估期内没有食品安全事件。",
        ],
    )
    temp = evidence / "_15_text.pdf"
    _save_pdf(temp, "历史框架合同", ["合同主体：博远包装有限公司", "交付周期：15个工作日", "付款：30%预付，70%验收后30日支付"])
    _scan_pdf(temp, evidence / "15_历史框架合同_扫描.pdf")

    receipts = [
        ("16a_银行电子回单.png", "博远包装有限公司", "18万元"),
        ("16b_银行电子回单.png", "博远供应链有限公司（关联公司）", "19.6万元"),
        ("16c_银行电子回单_低清.png", "博远包装有限公司", "20.8万元"),
    ]
    for index, (name, recipient, amount) in enumerate(receipts):
        temp = evidence / f"_{name}.pdf"
        _save_pdf(temp, "银行电子回单", [f"收款人：{recipient}", f"付款金额：{amount}", f"用途：第{index + 1}批纸餐盒采购款", "账号：演示脱敏 62******88"])
        _pdf_to_image(temp, evidence / name, low_quality="低清" in name)

    _save_docx(
        evidence / "17_采购评审会议纪要及利益冲突声明.docx",
        "采购评审会议纪要及利益冲突声明",
        [
            "委员会仅批准供应商B先试供30%，供应商A保留70%；价格、证书和主体风险关闭后再议正式份额。",
            "份额分配尚未完成断供压力测试或替代方案测算。",
            "利益冲突声明共6名评审人，已收5份，仍缺1名审批人签署。",
            "本次共登记17组材料，均在2026年6月30日前收到。",
        ],
    )

    statuses = [
        "supported", "supported", "supported", "conflict", "conflict", "conflict", "partial", "supported",
        "conflict", "supported", "conflict", "partial", "supported", "supported", "supported", "supported",
        "conflict", "supported", "conflict", "supported", "conflict", "conflict", "supported", "supported",
        "supported", "stale", "conflict", "needs_review", "stale", "conflict", "supported", "supported",
        "conflict", "unsupported", "unsupported", "supported", "supported", "supported", "conflict", "conflict",
        "partial", "conflict", "supported", "supported",
    ]
    accepted_overrides = {
        7: ["partial", "conflict"],
        12: ["partial", "conflict"],
        27: ["conflict", "partial"],
        33: ["conflict", "partial"],
        34: ["unsupported", "conflict"],
        35: ["unsupported", "partial", "conflict"],
        39: ["conflict", "partial"],
        40: ["conflict", "partial"],
        41: ["partial", "conflict"],
        42: ["conflict", "partial"],
    }
    evidence_files = [
        "01_审计口径与决策门槛.docx",
        "01_审计口径与决策门槛.docx",
        "01_审计口径与决策门槛.docx",
        "17_采购评审会议纪要及利益冲突声明.docx",
        "02_三家供应商报价与TCO测算.xlsx",
        "02_三家供应商报价与TCO测算.xlsx",
        "05_供应商B报价单_扫描.pdf",
        "04_供应商A报价单.pdf",
        "02_三家供应商报价与TCO测算.xlsx",
        "06_供应商C报价单.pdf",
        "02_三家供应商报价与TCO测算.xlsx",
        "05_供应商B报价单_扫描.pdf",
        "05_供应商B报价单_扫描.pdf",
        "05_供应商B报价单_扫描.pdf",
        "02_三家供应商报价与TCO测算.xlsx",
        "03_历史采购订单与履约KPI.xlsx",
        "03_历史采购订单与履约KPI.xlsx",
        "03_历史采购订单与履约KPI.xlsx",
        "03_历史采购订单与履约KPI.xlsx",
        "14_客户投诉与8D整改记录.docx",
        "03_历史采购订单与履约KPI.xlsx",
        "03_历史采购订单与履约KPI.xlsx",
        "13_试单验收记录.xlsx",
        "14_客户投诉与8D整改记录.docx",
        "14_客户投诉与8D整改记录.docx",
        "07_供应商B产品检测报告_扫描.pdf",
        "07_供应商B产品检测报告_扫描.pdf",
        "08_检测机构CMA能力附表_低清.tiff",
        "09_供应商B提交的FSC证书.jpg",
        "09_供应商B提交的FSC证书.jpg",
        "10_供应商B_ISO9001证书.pdf",
        "11_供应商主体与关联关系说明.docx",
        "12_现场验厂纪要.docx",
        "12_现场验厂纪要.docx",
        "12_现场验厂纪要.docx",
        "12_现场验厂纪要.docx",
        "12_现场验厂纪要.docx",
        "16a_银行电子回单.png",
        "16b_银行电子回单.png",
        "17_采购评审会议纪要及利益冲突声明.docx",
        "17_采购评审会议纪要及利益冲突声明.docx",
        "01_审计口径与决策门槛.docx",
        "01_审计口径与决策门槛.docx",
        "01_审计口径与决策门槛.docx",
    ]
    evidence_quotes = [
        "评估期间为2025年7月1日至2026年6月30日，证据截止日为2026年6月30日。",
        "年度需求为1200万个750mL双淋膜纸餐盒，年度预算为1080万元。",
        "年度需求为1200万个750mL双淋膜纸餐盒，年度预算为1080万元。",
        "委员会仅批准供应商B先试供30%，供应商A保留70%；价格、证书和主体风险关闭后再议正式份额。",
        "结果: 721600",
        "排名: 3",
        "报价有效期至2026年8月31日。",
        "报价：0.82元/个，含税含运",
        "报价(元/个): 0.76",
        "报价：0.77元/个，含税含运",
        "运费口径: 另计",
        "650mL单淋膜纸餐盒报价0.68元/个，未税，运费另计。",
        "750mL产品最小起订量为10万件，制版费4万元。",
        "付款条件为30%预付、70%验收后30日支付。",
        "制版费(元): 40000",
        "准时交付率: 0.932",
        "准时交付率: 0.884",
        "准时交付率: 0.901",
        "平均交付周期: 15个工作日",
        "供应商B在2026年4月发生一次重大延期，原因是主线设备故障。",
        "质量缺陷率: 0.018",
        "质量缺陷率: 0.005",
        "试单数量(件): 50000",
        "供应商B投诉平均关闭时间为2.1天。",
        "评估期内没有食品安全事件。",
        "签发日期：2024年5月20日",
        "样品：650mL单淋膜纸餐盒",
        "关键方法编号：GB 4806.8-2022（扫描模糊，需人工核对范围）",
        "有效期至2025年12月31日",
        "证书持有人：博远供应链有限公司（关联公司）",
        "有效期至2027年10月31日",
        "合同主体为博远包装有限公司，注册资本为3000万元。",
        "供应商自述月产能为800万件，本次未获得独立审计或设备利用记录。",
        "供应商自述月产能为800万件，本次未获得独立审计或设备利用记录。",
        "供应商自述月产能为800万件，本次未获得独立审计或设备利用记录。",
        "现场看到两条备用产线。",
        "原材料供应商执行年度复审，抽查到2025年度复审记录。",
        "付款金额：18万元",
        "收款人：博远供应链有限公司（关联公司）",
        "利益冲突声明共6名评审人，已收5份，仍缺1名审批人签署。",
        "份额分配尚未完成断供压力测试或替代方案测算。",
        "价格、证书、合同主体和检测范围风险关闭前，不得无条件续签供应商B。",
        "评估期间为2025年7月1日至2026年6月30日，证据截止日为2026年6月30日。",
        "本轮共登记17组证据材料，均在截止日前收到。",
    ]
    evidence_locators = [
        "paragraph 3", "paragraph 4", "paragraph 4", "paragraph 3",
        "公式说明!C3", "TCO复核!K3", "page 1 · bbox", "page 1 · bbox",
        "TCO复核!D3", "page 1 · bbox", "TCO复核!G3", "page 1 · bbox",
        "page 1 · bbox", "page 1 · bbox", "TCO复核!H3", "履约KPI!B2",
        "履约KPI!B3", "履约KPI!B4", "履约KPI!C3", "paragraph 3",
        "履约KPI!D3", "履约KPI!D4", "试单验收!B2", "paragraph 4",
        "paragraph 5", "page 1 · bbox", "page 1 · bbox", "image",
        "image", "image", "page 1 · bbox", "paragraph 3",
        "paragraph 4", "paragraph 4", "paragraph 4", "paragraph 3",
        "paragraph 5", "image", "image", "paragraph 5",
        "paragraph 4", "paragraph 6", "paragraph 3", "paragraph 7",
    ]
    issue_codes = [
        ["direct_support"], ["direct_support"], ["direct_support"], ["numeric_mismatch"],
        ["derived_calculation", "numeric_mismatch"], ["numeric_mismatch"], ["weak_support"],
        ["direct_support"], ["numeric_mismatch"], ["direct_support"], ["scope_mismatch"],
        ["scope_mismatch"], ["direct_support"], ["direct_support"], ["direct_support"],
        ["direct_support"], ["source_conflict", "numeric_mismatch"], ["direct_support"],
        ["unit_mismatch"], ["direct_support"], ["numeric_mismatch"], ["source_conflict"],
        ["direct_support"], ["direct_support"], ["direct_support"], ["stale_evidence"],
        ["scope_mismatch"], ["ocr_uncertainty"], ["stale_evidence"], ["entity_mismatch"],
        ["direct_support"], ["direct_support"], ["source_conflict"], ["missing_evidence"],
        ["missing_evidence"], ["direct_support"], ["direct_support"], ["derived_calculation"],
        ["entity_mismatch"], ["qualifier_overreach"], ["weak_support"], ["qualifier_overreach"],
        ["direct_support"], ["direct_support"],
    ]
    labels = {
        "name": "claimledger-procurement-flagship-v1",
        "disclosure": DISCLOSURE,
        "cases": [
            {
                "name": "procurement-renewal",
                "report": "report.docx",
                "sources": "evidence",
                "profile": "lite",
                "rule_pack": "procurement-zh",
                "as_of_date": "2026-07-18",
                "claims": [
                    {
                        "text": claim,
                        "should_flag": status != "supported",
                        "expected_status": status,
                        "accepted_statuses": accepted_overrides.get(index, [status]),
                        "expected_issue_codes": codes,
                        "expected_severity": "low" if status == "supported" else "medium" if status in {"partial", "stale", "needs_review"} else "high",
                        "evidence_file": evidence_file,
                        "evidence_quote": evidence_quote,
                        "evidence_locator": evidence_locator,
                    }
                    for index, (claim, status, codes, evidence_file, evidence_quote, evidence_locator) in enumerate(
                        zip(report_claims, statuses, issue_codes, evidence_files, evidence_quotes, evidence_locators),
                        start=1,
                    )
                ],
            }
        ],
    }
    labels_path = output / "labels.json"
    labels_path.write_text(json.dumps(labels, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "case_id": "CL-PROC-001",
        "title": "食品包装供应商续签决策报告证据审计",
        "disclosure": labels["disclosure"],
        "claims": len(report_claims),
        "planted_risks": sum(status != "supported" for status in statuses),
        "evidence_groups": 17,
        "report": report_path.name,
        "sources": evidence.name,
        "labels": labels_path.name,
    }
    manifest_path = output / "case_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "scenario": "procurement",
        "report": str(report_path),
        "sources": str(evidence),
        "labels": str(labels_path),
        "manifest": str(manifest_path),
    }


def _operations_demo(output: Path) -> dict[str, str]:
    """Generate a reconstructed monthly operations evidence-assurance case."""

    evidence = output / "evidence"
    evidence.mkdir(parents=True, exist_ok=True)
    report_claims = [
        "报告统计期间为2026年7月1日至2026年7月31日。",
        "7月共创建订单15420单。",
        "7月完成订单14880单。",
        "7月支付订单14600单。",
        "7月订单完成率为96.5%。",
        "7月完成订单环比增长6.3%。",
        "7月完成订单同比增长18.0%。",
        "7月全国仓准时交付率为96.4%。",
        "7月华东仓准时交付率为96.4%。",
        "7月华南仓准时交付率为97.2%。",
        "7月全国平均交付周期为2.8个自然日。",
        "7月平均交付周期环比下降30%。",
        "7月所有区域仓平均交付周期均低于3个自然日。",
        "7月投诉率按完成订单口径计算为0.42%。",
        "7月退货订单率为1.8%。",
        "7月退货订单率环比下降0.4个百分点。",
        "7月缺货SKU共37个。",
        "7月缺货SKU环比下降25%。",
        "7月全国仓库存准确率为99.2%。",
        "7月所有仓库库存准确率均不低于99%。",
        "7月承运商S在24小时内提货达成率为99.1%。",
        "7月承运商S没有发生重大履约异常。",
        "7月运费支出为186万元。",
        "7月每完成订单运费为125元。",
        "7月运费支出环比下降8%。",
        "7月华东仓产能利用率为82%。",
        "7月全国仓产能利用率为82%。",
        "截至目前，承运商S货运保险仍然有效。",
    ]
    report = Document()
    _style_document(
        report,
        "全国电商履约运营月报",
        "经营例会送审稿｜AI 初稿待证据审计",
        metadata=[
            ("送审对象", "经营管理委员会"),
            ("案例编号", "CL-OPS-001"),
            ("证据截止日", "2026-08-05"),
            ("交付状态", "存在口径与时效风险，未经人工闭环不得交付"),
        ],
    )
    report.add_heading("一、订单与增长", level=1)
    for claim in report_claims[:7]:
        report.add_paragraph(claim, style="List Number")
    report.add_heading("二、履约与客户体验", level=1)
    for claim in report_claims[7:16]:
        report.add_paragraph(claim, style="List Number")
    report.add_heading("三、库存、成本与承运商", level=1)
    for claim in report_claims[16:]:
        report.add_paragraph(claim, style="List Number")
    report_path = output / "report.docx"
    report.save(report_path)

    _save_docx(
        evidence / "01_运营口径与证据截点.docx",
        "运营口径与证据截点",
        [
            "报告统计期间为2026年7月1日至2026年7月31日，证据截止日为2026年8月5日。",
            "订单生命周期分为创建、支付和完成，三个分母不得混用。",
            "订单完成率按完成订单数除以创建订单数计算，区域指标必须与全国加权指标分开。",
            "同比使用2025年7月，环比使用2026年6月，交付周期统一采用自然日。",
        ],
    )

    orders = Workbook()
    order_sheet = orders.active
    order_sheet.title = "订单概览"
    order_sheet.append([
        "月份", "创建订单(单)", "完成订单(单)", "支付订单(单)",
        "订单完成率", "完成订单环比", "完成订单同比",
    ])
    order_sheet.append(["2026-06", 14520, 14000, 13720, 0.9642, "基期", "基期另表"])
    order_sheet.append(["2026-07", 15420, 14880, 14600, 0.965, 0.063, 0.127])
    prior = orders.create_sheet("同比基期")
    prior.append(["月份", "完成订单(单)"])
    prior.append(["2025-07", 13200])
    _style_workbook(orders)
    orders.save(evidence / "02_订单生命周期导出.xlsx")

    fulfillment = Workbook()
    region_sheet = fulfillment.active
    region_sheet.title = "区域履约"
    region_sheet.append(["区域", "准时交付率", "平均交付周期(个自然日)"])
    region_sheet.append(["全国仓", 0.964, 2.8])
    region_sheet.append(["华东仓", 0.941, 2.6])
    region_sheet.append(["华南仓", 0.972, 2.4])
    region_sheet.append(["华北仓", 0.960, 2.9])
    region_sheet.append(["西部仓", 0.928, 3.6])
    trend_sheet = fulfillment.create_sheet("交付趋势")
    trend_sheet.append(["月份", "全国平均交付周期(个自然日)", "环比下降率"])
    trend_sheet.append(["2026-06", 3.4, "基期"])
    trend_sheet.append(["2026-07", 2.8, 0.176])
    _style_workbook(fulfillment)
    fulfillment.save(evidence / "03_区域履约KPI.xlsx")

    service = Workbook()
    service_sheet = service.active
    service_sheet.title = "投诉退货"
    service_sheet.append([
        "月份", "投诉件数", "投诉率", "投诉率分母",
        "退货订单数", "退货订单率", "退货率环比下降",
    ])
    service_sheet.append(["2026-06", 56, 0.0041, "支付订单口径", 280, 0.020, "基期"])
    service_sheet.append(["2026-07", 61, 0.0042, "支付订单14600单", 268, 0.018, "0.2个百分点"])
    _style_workbook(service)
    service.save(evidence / "04_投诉与退货统计.xlsx")

    inventory = Workbook()
    accuracy = inventory.active
    accuracy.title = "库存准确率"
    accuracy.append(["区域", "库存准确率"])
    accuracy.append(["全国仓", 0.992])
    accuracy.append(["华东仓", 0.995])
    accuracy.append(["华南仓", 0.994])
    accuracy.append(["华北仓", 0.991])
    accuracy.append(["西部仓", 0.978])
    stockout = inventory.create_sheet("缺货SKU")
    stockout.append(["月份", "缺货SKU(个)", "环比下降率"])
    stockout.append(["2026-06", 42, "基期"])
    stockout.append(["2026-07", 37, 0.119])
    capacity = inventory.create_sheet("产能利用")
    capacity.append(["区域", "产能利用率"])
    capacity.append(["全国仓", 0.74])
    capacity.append(["华东仓", 0.82])
    capacity.append(["华南仓", 0.76])
    capacity.append(["华北仓", 0.71])
    capacity.append(["西部仓", 0.63])
    _style_workbook(inventory)
    inventory.save(evidence / "05_库存与产能台账.xlsx")

    _save_pdf(
        evidence / "06_承运商SLA与提货记录.pdf",
        "承运商S月度SLA复核",
        [
            "统计月份：2026年7月",
            "承运商S在24小时内提货达成率为99.1%。",
            "统计口径：全部已分配运输任务，不剔除延迟任务。",
        ],
    )
    _save_docx(
        evidence / "07_重大履约异常RCA.docx",
        "重大履约异常RCA",
        [
            "2026年7月18日，承运商S发生1次重大履约异常，造成华东仓96票订单延迟。",
            "根因是干线调度系统故障，承运商于2026年7月21日完成修复。",
            "该事件已进入月度SLA扣罚和持续改进清单。",
        ],
    )

    freight = Workbook()
    freight_sheet = freight.active
    freight_sheet.title = "运费结算"
    freight_sheet.append(["月份", "运费支出(万元)", "每完成订单运费(元)", "环比变化率"])
    freight_sheet.append(["2026-06", 176, 125.71, "基期"])
    freight_sheet.append(["2026-07", 186, 125, 0.057])
    _style_workbook(freight)
    freight.save(evidence / "08_运费结算与单位成本.xlsx")
    _save_pdf(
        evidence / "09_承运商货运保险.pdf",
        "承运商S货运保险凭证",
        [
            "被保险人：承运商S物流有限公司",
            "保险期间：2025年7月1日至2026年6月30日",
            "有效期至2026年6月30日，期满后未提供续保凭证。",
        ],
    )

    statuses = [
        "supported", "supported", "supported", "supported", "supported", "supported", "conflict",
        "supported", "conflict", "supported", "supported", "conflict", "conflict", "conflict",
        "supported", "conflict", "supported", "conflict", "supported", "conflict", "supported",
        "conflict", "supported", "supported", "conflict", "supported", "conflict", "stale",
    ]
    issue_codes = [
        ["direct_support"], ["direct_support"], ["direct_support"], ["direct_support"],
        ["direct_support"], ["direct_support"], ["numeric_mismatch"], ["direct_support"],
        ["numeric_mismatch", "scope_mismatch"], ["direct_support"], ["direct_support"],
        ["numeric_mismatch"], ["qualifier_overreach", "numeric_mismatch"], ["scope_mismatch"],
        ["direct_support"], ["numeric_mismatch"], ["direct_support"], ["numeric_mismatch"],
        ["direct_support"], ["qualifier_overreach", "numeric_mismatch"], ["direct_support"],
        ["qualifier_overreach"], ["direct_support"], ["direct_support"], ["numeric_mismatch"],
        ["direct_support"], ["scope_mismatch", "numeric_mismatch"], ["stale_evidence"],
    ]
    evidence_files = [
        "01_运营口径与证据截点.docx",
        "02_订单生命周期导出.xlsx", "02_订单生命周期导出.xlsx", "02_订单生命周期导出.xlsx",
        "02_订单生命周期导出.xlsx", "02_订单生命周期导出.xlsx", "02_订单生命周期导出.xlsx",
        "03_区域履约KPI.xlsx", "03_区域履约KPI.xlsx", "03_区域履约KPI.xlsx",
        "03_区域履约KPI.xlsx", "03_区域履约KPI.xlsx", "03_区域履约KPI.xlsx",
        "04_投诉与退货统计.xlsx", "04_投诉与退货统计.xlsx", "04_投诉与退货统计.xlsx",
        "05_库存与产能台账.xlsx", "05_库存与产能台账.xlsx", "05_库存与产能台账.xlsx",
        "05_库存与产能台账.xlsx", "06_承运商SLA与提货记录.pdf",
        "07_重大履约异常RCA.docx", "08_运费结算与单位成本.xlsx",
        "08_运费结算与单位成本.xlsx", "08_运费结算与单位成本.xlsx",
        "05_库存与产能台账.xlsx", "05_库存与产能台账.xlsx", "09_承运商货运保险.pdf",
    ]
    evidence_quotes = [
        "报告统计期间为2026年7月1日至2026年7月31日，证据截止日为2026年8月5日。",
        "创建订单(单): 15420", "完成订单(单): 14880", "支付订单(单): 14600",
        "订单完成率: 0.965", "完成订单环比: 0.063", "完成订单同比: 0.127",
        "准时交付率: 0.964", "准时交付率: 0.941", "准时交付率: 0.972",
        "平均交付周期(个自然日): 2.8", "环比下降率: 0.176",
        "平均交付周期(个自然日): 3.6", "投诉率分母: 支付订单14600单",
        "退货订单率: 0.018", "退货率环比下降: 0.2个百分点",
        "缺货SKU(个): 37", "环比下降率: 0.119", "库存准确率: 0.992",
        "库存准确率: 0.978", "承运商S在24小时内提货达成率为99.1%。",
        "2026年7月18日，承运商S发生1次重大履约异常，造成华东仓96票订单延迟。",
        "运费支出(万元): 186", "每完成订单运费(元): 125", "环比变化率: 0.057",
        "产能利用率: 0.82", "产能利用率: 0.74",
        "有效期至2026年6月30日，期满后未提供续保凭证。",
    ]
    evidence_locators = [
        "paragraph 3",
        "订单概览!B3", "订单概览!C3", "订单概览!D3", "订单概览!E3",
        "订单概览!F3", "订单概览!G3",
        "区域履约!B2", "区域履约!B3", "区域履约!B4", "区域履约!C2",
        "交付趋势!C3", "区域履约!C6",
        "投诉退货!D3", "投诉退货!F3", "投诉退货!G3",
        "缺货SKU!B3", "缺货SKU!C3", "库存准确率!B2", "库存准确率!B6",
        "page 1 · bbox", "paragraph 3",
        "运费结算!B3", "运费结算!C3", "运费结算!D3",
        "产能利用!B3", "产能利用!B2", "page 1 · bbox",
    ]
    accepted_overrides = {
        13: ["conflict", "partial"],
        20: ["conflict", "partial"],
    }
    labels = {
        "name": "claimledger-operations-golden-v1",
        "disclosure": DISCLOSURE,
        "cases": [
            {
                "name": "operations-fulfillment-monthly",
                "report": "report.docx",
                "sources": "evidence",
                "profile": "deterministic",
                "rule_pack": "operations-zh",
                "as_of_date": "2026-08-05",
                "claims": [
                    {
                        "text": claim,
                        "should_flag": status != "supported",
                        "expected_status": status,
                        "accepted_statuses": accepted_overrides.get(index, [status]),
                        "expected_issue_codes": codes,
                        "expected_severity": "low" if status == "supported" else "medium" if status == "stale" else "high",
                        "evidence_file": evidence_file,
                        "evidence_quote": evidence_quote,
                        "evidence_locator": evidence_locator,
                    }
                    for index, (claim, status, codes, evidence_file, evidence_quote, evidence_locator) in enumerate(
                        zip(report_claims, statuses, issue_codes, evidence_files, evidence_quotes, evidence_locators),
                        start=1,
                    )
                ],
            }
        ],
    }
    labels_path = output / "labels.json"
    labels_path.write_text(json.dumps(labels, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "case_id": "CL-OPS-001",
        "title": "全国电商履约运营月报证据审计",
        "disclosure": DISCLOSURE,
        "claims": len(report_claims),
        "planted_risks": sum(status != "supported" for status in statuses),
        "evidence_groups": 9,
        "report": report_path.name,
        "sources": evidence.name,
        "labels": labels_path.name,
    }
    manifest_path = output / "case_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "scenario": "operations",
        "report": str(report_path),
        "sources": str(evidence),
        "labels": str(labels_path),
        "manifest": str(manifest_path),
    }


def _consulting_demo(output: Path) -> dict[str, str]:
    """Generate a reconstructed consulting market-entry delivery case."""

    evidence = output / "evidence"
    evidence.mkdir(parents=True, exist_ok=True)
    report_claims = [
        "研究期间为2026年1月1日至2026年6月30日。",
        "本报告市场口径包含视觉相机、边缘计算、软件和系统集成。",
        "2025年全国工业视觉质检市场规模达到85亿元。",
        "2025年华东市场规模达到68亿元。",
        "2025年硬件市场规模达到68亿元。",
        "2025年全国全栈市场规模为68亿元。",
        "2026年基准情景市场规模为78.2亿元。",
        "基准情景2025年至2028年复合增长率为15%。",
        "基准情景未来三年市场规模将保持25%的增长率。",
        "乐观情景2025年至2028年复合增长率为25%。",
        "核心收入预计采用乐观情景。",
        "本轮访谈共覆盖24家企业。",
        "访谈样本包含18家大型制造企业和6家中小企业。",
        "所有受访客户均将价格视为唯一采购因素。",
        "24家样本中有15家将价格列入前三项采购因素。",
        "24家样本中有18家将交付稳定性列入前三项采购因素。",
        "交付稳定性进入前三项采购因素的样本占75%。",
        "全部24家受访企业均计划在未来12个月采购。",
        "计划在未来12个月采购的样本占58.3%。",
        "本轮访谈可以代表全国市场需求。",
        "6家试点客户的平均漏检率下降22%。",
        "漏检率下降22%的结论适用于全部制造企业。",
        "6家试点客户的平均投资回收期为14个月。",
        "截至目前，竞品一体机成交价区间为35万元至80万元。",
        "竞争对手甲2025年相关业务收入为12亿元。",
        "竞争对手乙在全国市场的份额为40%。",
        "现行政策要求所有工厂在2027年前部署AI质检系统。",
        "项目可以无条件在全国全面商业化，无需补充验证。",
    ]
    report = Document()
    _style_document(
        report,
        "工业视觉质检市场进入建议",
        "客户交付稿｜AI 初稿待证据审计",
        metadata=[
            ("送审对象", "客户战略委员会"),
            ("案例编号", "CL-CONS-001"),
            ("证据截止日", "2026-07-18"),
            ("交付状态", "存在市场边界、样本外推和时效风险"),
        ],
    )
    report.add_heading("一、市场规模与预测", level=1)
    for claim in report_claims[:11]:
        report.add_paragraph(claim, style="List Number")
    report.add_heading("二、访谈与试点", level=1)
    for claim in report_claims[11:23]:
        report.add_paragraph(claim, style="List Number")
    report.add_heading("三、竞品、政策与建议", level=1)
    for claim in report_claims[23:]:
        report.add_paragraph(claim, style="List Number")
    report_path = output / "report.docx"
    report.save(report_path)

    _save_docx(
        evidence / "01_研究口径与情景选择.docx",
        "研究口径与情景选择",
        [
            "研究期间为2026年1月1日至2026年6月30日，证据截止日为2026年7月18日。",
            "本报告市场口径包含视觉相机、边缘计算、软件和系统集成，统称全国全栈市场。",
            "核心收入预计采用基准情景，乐观情景只用于压力测试。",
            "硬件市场仅包含视觉相机和边缘计算设备，不含软件与系统集成。",
        ],
    )
    market = Workbook()
    boundary = market.active
    boundary.title = "市场边界"
    boundary.append(["年份", "地域", "市场口径", "市场规模(亿元)"])
    boundary.append([2025, "全国市场", "全栈市场", 68])
    boundary.append([2025, "华东市场", "全栈市场", 24])
    boundary.append([2025, "全国市场", "硬件市场", 31])
    forecast = market.create_sheet("情景预测")
    forecast.append(["情景", "2025市场规模(亿元)", "2026市场规模(亿元)", "2028市场规模(亿元)", "复合增长率"])
    forecast.append(["基准情景", 68, 78.2, 103.5, 0.15])
    forecast.append(["乐观情景", 68, 85, 132.8, 0.25])
    forecast.append(["保守情景", 68, 72.5, 88.1, 0.09])
    _style_workbook(market)
    market.save(evidence / "02_市场规模与情景模型.xlsx")

    interviews = Workbook()
    sample = interviews.active
    sample.title = "样本概览"
    sample.append(["指标", "数量(家)"])
    sample.append(["总样本", 24])
    sample.append(["大型制造企业", 18])
    sample.append(["中小企业", 6])
    factors = interviews.create_sheet("采购因素")
    factors.append(["采购因素", "列入前三项数量(家)", "样本占比"])
    factors.append(["价格", 15, 0.625])
    factors.append(["交付稳定性", 18, 0.75])
    factors.append(["售后能力", 12, 0.50])
    purchase = interviews.create_sheet("采购计划")
    purchase.append(["计划窗口", "计划采购数量(家)", "样本数(家)", "样本占比"])
    purchase.append(["未来12个月", 14, 24, 0.583])
    _style_workbook(interviews)
    interviews.save(evidence / "03_访谈样本与编码结果.xlsx")
    _save_docx(
        evidence / "04_访谈方法与局限.docx",
        "访谈方法与局限",
        [
            "本轮访谈共覆盖24家企业，其中18家大型制造企业、6家中小企业。",
            "价格并非唯一采购因素，交付稳定性和售后能力也会影响采购决定。",
            "样本仅覆盖华东市场和华南市场，未做全国口径加权，不能代表全国市场需求。",
        ],
    )

    pilot = Workbook()
    pilot_sheet = pilot.active
    pilot_sheet.title = "试点结果"
    pilot_sheet.append(["范围", "试点客户(家)", "平均漏检率下降", "平均投资回收期(月)"])
    pilot_sheet.append(["6家试点样本内", 6, 0.22, 14])
    limitation = pilot.create_sheet("外推限制")
    limitation.append(["结论", "审计说明"])
    limitation.append(["漏检率下降22%", "仅能解释6家试点样本内结果，不能适用于全部制造企业。"])
    _style_workbook(pilot)
    pilot.save(evidence / "05_六家试点结果.xlsx")
    _save_pdf(
        evidence / "06_竞品价格清单_2024.pdf",
        "竞品一体机价格访谈清单",
        [
            "资料日期：2024年6月30日",
            "样本内成交价区间为35万元至80万元。",
            "该清单在2024年后未按同口径更新。",
        ],
    )
    _save_pdf(
        evidence / "07_竞争对手甲年报摘录.pdf",
        "竞争对手甲2025年年报摘录",
        [
            "2025年工业视觉质检相关业务收入为1.2亿元。",
            "该数字来自经审计分部附注，不等同于公司总收入。",
        ],
    )
    _save_docx(
        evidence / "08_竞争对手乙访谈摘要.docx",
        "竞争对手乙访谈摘要",
        [
            "竞争对手乙在华东市场受访样本中的提及率为40%。",
            "该比例不是全国市场份额，也没有经过收入或装机量验证。",
        ],
    )
    _save_pdf(
        evidence / "09_政策原文摘录.pdf",
        "智能制造试点政策摘录",
        [
            "发布日期：2026年3月15日",
            "政策鼓励符合条件的工厂开展AI质检试点。",
            "政策未要求所有工厂在2027年前完成部署，也未设置普遍部署义务。",
        ],
    )
    _save_docx(
        evidence / "10_项目评审纪要.docx",
        "项目评审纪要",
        [
            "委员会仅建议在华东市场和华南市场开展区域试点，补充样本验证后再议全国商业化。",
            "当前证据不足以支持无条件全面商业化，客户行业分层和渠道成本仍需验证。",
        ],
    )

    statuses = [
        "supported", "supported", "conflict", "conflict", "conflict", "supported", "supported",
        "supported", "conflict", "supported", "conflict", "supported", "supported", "conflict",
        "supported", "supported", "supported", "conflict", "supported", "conflict", "supported",
        "conflict", "supported", "stale", "conflict", "conflict", "conflict", "conflict",
    ]
    issue_codes = [
        ["direct_support"], ["direct_support"], ["numeric_mismatch"], ["scope_mismatch", "numeric_mismatch"],
        ["scope_mismatch", "numeric_mismatch"], ["direct_support"], ["direct_support"], ["direct_support"],
        ["scope_mismatch", "numeric_mismatch"], ["direct_support"], ["scope_mismatch"], ["direct_support"],
        ["direct_support"], ["qualifier_overreach"], ["direct_support"], ["direct_support"], ["direct_support"],
        ["numeric_mismatch", "qualifier_overreach"], ["direct_support"], ["scope_mismatch"],
        ["direct_support"], ["qualifier_overreach"], ["direct_support"], ["stale_evidence"],
        ["numeric_mismatch"], ["scope_mismatch"], ["qualifier_overreach"], ["scope_mismatch", "qualifier_overreach"],
    ]
    evidence_files = [
        "01_研究口径与情景选择.docx", "01_研究口径与情景选择.docx",
        "02_市场规模与情景模型.xlsx", "02_市场规模与情景模型.xlsx",
        "02_市场规模与情景模型.xlsx", "02_市场规模与情景模型.xlsx",
        "02_市场规模与情景模型.xlsx", "02_市场规模与情景模型.xlsx",
        "02_市场规模与情景模型.xlsx", "02_市场规模与情景模型.xlsx",
        "01_研究口径与情景选择.docx", "03_访谈样本与编码结果.xlsx",
        "04_访谈方法与局限.docx", "04_访谈方法与局限.docx",
        "03_访谈样本与编码结果.xlsx", "03_访谈样本与编码结果.xlsx",
        "03_访谈样本与编码结果.xlsx", "03_访谈样本与编码结果.xlsx",
        "03_访谈样本与编码结果.xlsx", "04_访谈方法与局限.docx",
        "05_六家试点结果.xlsx", "05_六家试点结果.xlsx", "05_六家试点结果.xlsx",
        "06_竞品价格清单_2024.pdf", "07_竞争对手甲年报摘录.pdf",
        "08_竞争对手乙访谈摘要.docx", "09_政策原文摘录.pdf", "10_项目评审纪要.docx",
    ]
    evidence_quotes = [
        "研究期间为2026年1月1日至2026年6月30日，证据截止日为2026年7月18日。",
        "本报告市场口径包含视觉相机、边缘计算、软件和系统集成，统称全国全栈市场。",
        "市场规模(亿元): 68", "市场规模(亿元): 24", "市场规模(亿元): 31",
        "市场规模(亿元): 68", "2026市场规模(亿元): 78.2", "复合增长率: 0.15",
        "复合增长率: 0.15", "复合增长率: 0.25",
        "核心收入预计采用基准情景，乐观情景只用于压力测试。",
        "数量(家): 24",
        "本轮访谈共覆盖24家企业，其中18家大型制造企业、6家中小企业。",
        "价格并非唯一采购因素，交付稳定性和售后能力也会影响采购决定。",
        "列入前三项数量(家): 15", "列入前三项数量(家): 18", "样本占比: 0.75",
        "计划采购数量(家): 14", "样本占比: 0.583",
        "样本仅覆盖华东市场和华南市场，未做全国口径加权，不能代表全国市场需求。",
        "平均漏检率下降: 0.22",
        "审计说明: 仅能解释6家试点样本内结果，不能适用于全部制造企业。",
        "平均投资回收期(月): 14",
        "样本内成交价区间为35万元至80万元。",
        "2025年工业视觉质检相关业务收入为1.2亿元。",
        "竞争对手乙在华东市场受访样本中的提及率为40%。",
        "政策未要求所有工厂在2027年前完成部署，也未设置普遍部署义务。",
        "当前证据不足以支持无条件全面商业化，客户行业分层和渠道成本仍需验证。",
    ]
    evidence_locators = [
        "paragraph 3", "paragraph 4",
        "市场边界!D2", "市场边界!D3", "市场边界!D4", "市场边界!D2",
        "情景预测!C2", "情景预测!E2", "情景预测!E2", "情景预测!E3",
        "paragraph 5", "样本概览!B2", "paragraph 3", "paragraph 4",
        "采购因素!B2", "采购因素!B3", "采购因素!C3",
        "采购计划!B2", "采购计划!D2", "paragraph 5",
        "试点结果!C2", "外推限制!B2", "试点结果!D2",
        "page 1 · bbox", "page 1 · bbox", "paragraph 3", "page 1 · bbox", "paragraph 4",
    ]
    labels = {
        "name": "claimledger-consulting-golden-v1",
        "disclosure": DISCLOSURE,
        "cases": [
            {
                "name": "consulting-market-entry",
                "report": "report.docx",
                "sources": "evidence",
                "profile": "deterministic",
                "rule_pack": "consulting-zh",
                "as_of_date": "2026-07-18",
                "claims": [
                    {
                        "text": claim,
                        "should_flag": status != "supported",
                        "expected_status": status,
                        "accepted_statuses": [status],
                        "expected_issue_codes": codes,
                        "expected_severity": "low" if status == "supported" else "medium" if status == "stale" else "high",
                        "evidence_file": evidence_file,
                        "evidence_quote": evidence_quote,
                        "evidence_locator": evidence_locator,
                    }
                    for claim, status, codes, evidence_file, evidence_quote, evidence_locator in zip(
                        report_claims, statuses, issue_codes, evidence_files, evidence_quotes, evidence_locators
                    )
                ],
            }
        ],
    }
    labels_path = output / "labels.json"
    labels_path.write_text(json.dumps(labels, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "case_id": "CL-CONS-001",
        "title": "工业视觉质检市场进入建议证据审计",
        "disclosure": DISCLOSURE,
        "claims": len(report_claims),
        "planted_risks": sum(status != "supported" for status in statuses),
        "evidence_groups": 10,
        "report": report_path.name,
        "sources": evidence.name,
        "labels": labels_path.name,
    }
    manifest_path = output / "case_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "scenario": "consulting",
        "report": str(report_path),
        "sources": str(evidence),
        "labels": str(labels_path),
        "manifest": str(manifest_path),
    }


def _golden_suite(output: Path) -> dict[str, str]:
    """Generate the complete 100-claim cross-industry golden dataset."""

    generated = {
        "procurement": _procurement_demo(output / "procurement"),
        "operations": _operations_demo(output / "operations"),
        "consulting": _consulting_demo(output / "consulting"),
    }
    cases = []
    total_claims = 0
    total_risks = 0
    total_evidence_groups = 0
    for scenario, result in generated.items():
        payload = json.loads(Path(result["labels"]).read_text(encoding="utf-8"))
        for case in payload["cases"]:
            case = dict(case)
            case["report"] = f"{scenario}/{case['report']}"
            case["sources"] = f"{scenario}/{case['sources']}"
            cases.append(case)
        case_manifest = json.loads(Path(result["manifest"]).read_text(encoding="utf-8"))
        total_claims += int(case_manifest["claims"])
        total_risks += int(case_manifest["planted_risks"])
        total_evidence_groups += int(case_manifest["evidence_groups"])
    labels = {
        "name": "claimledger-cross-industry-golden-v2",
        "disclosure": DISCLOSURE,
        "cases": cases,
    }
    labels_path = output / "labels.json"
    labels_path.write_text(json.dumps(labels, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "suite_id": "CL-GOLDEN-V2",
        "title": "ClaimLedger 跨行业报告证据审计黄金集",
        "disclosure": DISCLOSURE,
        "industries": ["采购决策", "运营管理", "咨询研究"],
        "claims": total_claims,
        "planted_risks": total_risks,
        "evidence_groups": total_evidence_groups,
        "case_ids": ["CL-PROC-001", "CL-OPS-001", "CL-CONS-001"],
        "labels": labels_path.name,
    }
    manifest_path = output / "suite_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "scenario": "golden-v2",
        "root": str(output),
        "labels": str(labels_path),
        "manifest": str(manifest_path),
        "procurement": str(output / "procurement"),
        "operations": str(output / "operations"),
        "consulting": str(output / "consulting"),
    }


def _small_demo(output: Path, scenario: str) -> dict[str, str]:
    evidence_dir = output / "evidence"
    evidence_dir.mkdir(exist_ok=True)
    report = Document()
    source = Document()
    workbook = Workbook()
    sheet = workbook.active
    if scenario == "lithium":
        report.add_heading("锂原料市场简报", 0)
        report.add_paragraph("截至目前，电池级碳酸锂价格为每吨120000元，较上月增长15%。")
        report.add_paragraph("SC6锂辉石与低品位锂云母可以直接按照吨价比较采购成本。")
        report.add_paragraph("全球航运指数下降表明澳洲至中国的锂矿进口路线运价下降20%。")
        report.add_paragraph("2026年7月，电池级碳酸锂现货参考价为每吨120000元。")
        source.add_heading("采购证据摘录", 0)
        source.add_paragraph("2026年7月，电池级碳酸锂现货参考价为每吨120000元，月度涨幅为5%。")
        source.add_paragraph("SC6为约6%品位锂辉石精矿；锂云母原矿品位和折算回收率不同，吨价不可直接比较。")
        source.add_paragraph("波罗的海全球航运指数不能代表澳洲至中国的具体进口航线运价。")
        source.add_paragraph("2026年7月，电池级碳酸锂现货参考价为每吨120000元。")
        sheet.title = "Market Data"
        sheet.append(["日期", "品类", "价格(元/吨)", "月度涨幅"])
        sheet.append(["2026-07", "电池级碳酸锂", 120000, "5%"])
    elif scenario == "operations":
        report.add_heading("运营月报", 0)
        report.add_paragraph("7月完成订单1280单，同比增长18%。")
        report.add_paragraph("本月平均交付周期下降至4.2天，环比下降30%。")
        report.add_paragraph("7月完成订单1280单。")
        source.add_heading("运营口径说明", 0)
        source.add_paragraph("7月完成订单1280单，去年7月完成1200单，同比增长6.7%。")
        source.add_paragraph("6月平均交付周期为5.0天，7月为4.2天，环比下降16%。")
        source.add_paragraph("7月完成订单1280单。")
        sheet.title = "Operations"
        sheet.append(["月份", "订单", "平均交付天数"])
        sheet.append(["2026-06", 1200, 5.0])
        sheet.append(["2026-07", 1280, 4.2])
    elif scenario == "consulting":
        report.add_heading("市场咨询报告", 0)
        report.add_paragraph("截至目前，目标市场规模达到85亿元，预计未来三年保持25%增长。")
        report.add_paragraph("主要客户均将价格视为唯一采购因素。")
        report.add_paragraph("2024年研究估算目标市场规模约为62亿元。")
        source.add_heading("访谈与公开材料摘录", 0)
        source.add_paragraph("2024年研究估算目标市场规模约为62亿元；后续缺少同口径更新。")
        source.add_paragraph("客户访谈显示价格、交付稳定性和售后能力共同影响采购决定。")
        source.add_paragraph("2024年研究估算目标市场规模约为62亿元。")
        sheet.title = "Research"
        sheet.append(["年份", "市场规模(亿元)", "来源口径"])
        sheet.append([2024, 62, "公开研究估算"])
    else:
        raise ValueError(f"unknown demo scenario: {scenario}")
    report_path = output / "report.docx"
    report.save(report_path)
    source_path = evidence_dir / "public-evidence.docx"
    source.save(source_path)
    workbook_path = evidence_dir / "market-data.xlsx"
    workbook.save(workbook_path)
    return {"scenario": scenario, "report": str(report_path), "sources": str(evidence_dir), "xlsx": str(workbook_path)}


def create_demo(output: Path, scenario: str = "lithium") -> dict[str, str]:
    output.mkdir(parents=True, exist_ok=True)
    if scenario == "procurement":
        return _procurement_demo(output)
    if scenario == "operations":
        return _operations_demo(output)
    if scenario == "consulting":
        return _consulting_demo(output)
    if scenario == "golden-v2":
        return _golden_suite(output)
    return _small_demo(output, scenario)
